"""SVF-RD-V4: V3 grid + per-song Total Revenue by Indian FY, plus Statement-V1 summary."""
import glob, os, re, zipfile, warnings, datetime as dt
from collections import defaultdict
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

IN_DIR = "/home/swatantra/RD/input/batch-1"
OUT_DIR = "/home/swatantra/RD/output/batch1-output"
SRC = os.path.join(OUT_DIR, "SVF-RD-V3.xlsx")
DST = os.path.join(OUT_DIR, "SVF-RD-V4.xlsx")
DST2 = os.path.join(OUT_DIR, "Statement-V1.xlsx")
TOL = 0.05
MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
MLBL = {v: k.capitalize() for k, v in MONTHS.items()}
SOCIETIES = {"008": ("APRA", "Australia"), "021": ("BMI", "USA"), "023": ("BUMA", "Netherlands"),
             "026": ("CASH", "Hong Kong"), "058": ("SACEM", "France"), "080": ("SUISA", "Switzerland"),
             "101": ("SOCAN", "Canada"), "104": ("MACP", "Malaysia"), "106": ("COMPASS", "Singapore"),
             "126": ("MCT", "Thailand"), "128": ("IMRO", "Ireland")}


def s(v):
    return "" if v is None else str(v).strip()


def num(v):
    try:
        f = float(v)
        return 0.0 if f != f else f
    except (TypeError, ValueError):
        return 0.0


def month_end(y, m):
    return dt.date(y + (m == 12), 1 if m == 12 else m + 1, 1) - dt.timedelta(days=1)


def fin_year(d):
    return d.year if d.month >= 4 else d.year - 1


def fy_label(y):
    return f"FY {y}-{str(y + 1)[-2:]}"


def stmt_date(path):
    with zipfile.ZipFile(path) as z:
        ts = max(i.date_time for i in z.infolist())
    return dt.date(ts[0], ts[1], ts[2])


def parse_periods(stem):
    fy = ""
    m = re.search(r"F\.?\s?Y\.?\s*(\d{4})\s*-\s*(\d{2,4})", stem, re.I)
    if m:
        fy = f"{m.group(1)}-{m.group(2)[-2:]}"
    hits = []
    for mm in re.finditer(r"([A-Za-z]{3,12})\.?\s+(\d{4})", stem):
        k = mm.group(1)[:3].lower()
        if k in MONTHS:
            hits.append((MONTHS[k], int(mm.group(2))))
    if hits:
        return dt.date(hits[0][1], hits[0][0], 1), month_end(hits[-1][1], hits[-1][0]), fy
    if fy:
        y = int(fy.split("-")[0])
        return dt.date(y, 4, 1), month_end(y + 1, 3), fy
    return None, None, fy


def parse_filename(path):
    fn = os.path.basename(path)
    stem = os.path.splitext(fn)[0]
    n = int(re.search(r"\(S-(\d+)\)", fn).group(1))
    low = stem.lower()
    sub = re.findall(r"\b([PM]\d{4}[A-Z]\d{3})\b", stem)
    runs = re.findall(r"\b([PM]\d{4})\b", stem)
    if sub:
        dist_no = sub[0]
    elif re.search(r"[PM]\d{4}\s+to\s+[PM]\d{4}", stem, re.I) and len(runs) > 1:
        dist_no = f"{runs[0]} to {runs[1]}"
    else:
        dist_no = runs[0] if runs else ""
    soc = country = ""
    ms = re.search(r"(\d{3})\s+([A-Z]{2,10})\s*$", stem.strip())
    if ms and ms.group(1) in SOCIETIES:
        soc, country = SOCIETIES[ms.group(1)]
    if "redistribution" in low:
        cat = "Redistribution"
    elif "mechanical" in low:
        cat = "Mechanical (MUSERK)"
    elif "apple music" in low:
        cat = "Apple Music"
    elif "spotify" in low:
        cat = "Spotify"
    elif "youtube" in low and "pre" in low:
        cat = "YouTube Pre-Claims"
    elif "youtube" in low and "post" in low:
        cat = "YouTube Post-Claims"
    elif "facebook" in low or "meta" in low:
        cat = "Facebook / Meta"
    elif "radio" in low:
        cat = "Radio"
    elif "zee television" in low:
        cat = "Zee TV Broadcast"
    elif "overseas" in low:
        cat = "Overseas"
    else:
        cat = "Other / Unclassified"
    st, en, fy = parse_periods(stem)
    if st:
        plabel = f"{MLBL[st.month]} {st.year} - {MLBL[en.month]} {en.year}"
        if fy:
            plabel += f" (FY {fy})"
    else:
        plabel = "Not specified"
    section = f"Overseas - {soc} ({country})" if (cat == "Overseas" and soc) else cat
    return dict(s_no=n, sheet=f"S-{n}", file=fn, dist_no=dist_no, section=section,
                p_start=st, p_end=en, period=plabel, date=stmt_date(path))


def fy_months(st, en):
    out = defaultdict(int)
    y, m = st.year, st.month
    while (y, m) <= (en.year, en.month):
        out[fin_year(dt.date(y, m, 1))] += 1
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def split_period(label):
    """period label -> {fy_start: weight fraction} ; {} when not stated"""
    lab = re.sub(r"\(.*?\)", "", s(label)).strip()
    hits = []
    for mm in re.finditer(r"([A-Za-z]{3,12})\.?\s+(\d{4})", lab):
        k = mm.group(1)[:3].lower()
        if k in MONTHS:
            hits.append((MONTHS[k], int(mm.group(2))))
    if not hits:
        return {}
    st = dt.date(hits[0][1], hits[0][0], 1)
    en = month_end(hits[-1][1], hits[-1][0])
    mm = fy_months(st, en)
    tot = sum(mm.values())
    return {y: k / tot for y, k in mm.items()}


def penny_fix(parts, target):
    """round dict of floats to 2dp so they sum exactly to target"""
    out = {k: round(v, 2) for k, v in parts.items()}
    gap = round(target - round(sum(out.values()), 2), 2)
    if gap and out:
        big = max(out, key=lambda k: abs(out[k]))
        out[big] = round(out[big] + gap, 2)
    return out


# ================================================================== sheet 1
wb = openpyxl.load_workbook(SRC)
ws = wb["Royalty by Song"]
ws.title = "SVF-RD-V4"

BOLD = Font(bold=True)
SECF = Font(bold=True, size=12, color="1F1F1F")
CTR = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
MED = Side(style="medium", color="808080")
THIN = Side(style="thin", color="BFBFBF")
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
MONEY = "#,##0.00"

GRID_LAST_COL = 184
TOTAL_ROW = next(r for r in range(3, ws.max_row + 1) if s(ws.cell(r, 2).value) == "TOTAL")
GRID_LAST_ROW = max(r for r in range(3, TOTAL_ROW) if ws.cell(r, 3).value is not None)

# ---- per-song FY split from the grid itself
rows = {}
fys = set()
for r in range(3, GRID_LAST_ROW + 1):
    raw = defaultdict(float)
    for c in range(5, GRID_LAST_COL + 1, 4):
        amt = ws.cell(r, c + 1).value
        if amt is None:
            continue
        w = split_period(ws.cell(r, c + 2).value)
        if not w:
            raw["NA"] += num(amt)
        else:
            for y, f in w.items():
                raw[y] += num(amt) * f
    tgt = round(num(ws.cell(r, 4).value), 2)
    rows[r] = penny_fix(raw, tgt) if raw else {}
    fys |= {k for k in raw if k != "NA"}

keys = sorted(fys) + (["NA"] if any("NA" in v for v in rows.values()) else [])
headers = [(fy_label(k) if k != "NA" else "Period Not Stated") for k in keys] + ["Total Revenue"]

START = GRID_LAST_COL + 4                       # leave 3 blank columns
FYFILL = ["E8F2E8", "E6EDF7", "F9E9F2", "F5F0E5", "EEEAF7", "E7F0F2"]
BANNER = PatternFill("solid", fgColor="C9B99B")
sc = ws.cell(1, START, value="TOTAL REVENUE  (by Indian Financial Year, 1 Apr - 31 Mar, "
                             "of the period the royalty was earned in)")
sc.font = SECF
sc.alignment = CTR
for c in range(START, START + len(headers)):
    ws.cell(1, c).fill = BANNER
ws.merge_cells(start_row=1, start_column=START, end_row=1, end_column=START + len(headers) - 1)
ws.cell(1, START).border = Border(left=MED, top=MED, bottom=THIN)
ws.cell(1, START + len(headers) - 1).border = Border(right=MED, top=MED, bottom=THIN)

for j, h in enumerate(headers):
    c = ws.cell(2, START + j, value=h)
    c.font = BOLD
    c.fill = PatternFill("solid", fgColor="D8CBB3" if j < len(keys) else "C9B99B")
    c.alignment = WRAP
    c.border = Border(top=THIN, bottom=MED, left=THIN, right=THIN)

col_tot = defaultdict(float)
for r in range(3, GRID_LAST_ROW + 1):
    d = rows[r]
    for j, k in enumerate(keys):
        c = ws.cell(r, START + j, value=round(d.get(k, 0.0), 2))
        c.number_format = MONEY
        c.fill = PatternFill("solid", fgColor=FYFILL[j % len(FYFILL)])
        c.border = Border(left=THIN, right=THIN)
        col_tot[k] += d.get(k, 0.0)
    tc = ws.cell(r, START + len(keys), value=round(sum(d.values()), 2))
    tc.number_format = MONEY
    tc.font = BOLD
    tc.fill = PatternFill("solid", fgColor="F0E9DC")
    ws.cell(r, START).border = Border(left=MED)
    tc.border = Border(right=MED, left=THIN)

for j, k in enumerate(keys):
    c = ws.cell(TOTAL_ROW, START + j, value=round(col_tot[k], 2))
    c.number_format = MONEY
    c.font = BOLD
    c.fill = BANNER
    c.border = Border(top=MED, bottom=MED)
gt = ws.cell(TOTAL_ROW, START + len(keys), value=round(sum(col_tot.values()), 2))
gt.number_format = MONEY
gt.font = BOLD
gt.fill = BANNER
gt.border = Border(top=MED, bottom=MED, right=MED)
ws.cell(TOTAL_ROW, START).border = Border(top=MED, bottom=MED, left=MED)
for j in range(len(headers)):
    ws.column_dimensions[get_column_letter(START + j)].width = 16
for j in range(1, 4):
    ws.column_dimensions[get_column_letter(GRID_LAST_COL + j)].width = 3
ws.auto_filter.ref = f"A2:{get_column_letter(START + len(headers) - 1)}{GRID_LAST_ROW}"

nc = ws.cell(TOTAL_ROW + 1, START,
             value="Basis: each distribution's amount is placed in the Indian FY (1 Apr - 31 Mar) "
                   "of the period it was earned in; a statement covering more than one FY is split "
                   "month-weighted. Row total = Total Amount (col D).")
nc.font = Font(italic=True, size=10, color="404040")
ws.merge_cells(start_row=TOTAL_ROW + 1, start_column=START, end_row=TOTAL_ROW + 1,
               end_column=START + len(headers) - 1)

grand_grid = round(sum(num(ws.cell(r, 4).value) for r in range(3, GRID_LAST_ROW + 1)), 2)
print(f"V4 songs={GRID_LAST_ROW - 2}  grand={grand_grid:,.2f}  fy_total={sum(col_tot.values()):,.2f}")
mismatch = [r for r in range(3, GRID_LAST_ROW + 1)
            if abs(round(sum(rows[r].values()), 2) - round(num(ws.cell(r, 4).value), 2)) > 0.005]
print(f"row mismatches={len(mismatch)}")
for k in keys:
    print(f"   {fy_label(k) if k != 'NA' else 'Not stated':<12} {round(col_tot[k], 2):>15,.2f}")

# ================================================================== sheet 2
paths = sorted(glob.glob(os.path.join(IN_DIR, "*.xlsx")),
               key=lambda p: int(re.search(r"\(S-(\d+)\)", os.path.basename(p)).group(1)))
stmts = []
for p in paths:
    b = openpyxl.load_workbook(p, data_only=True)
    g = [list(r) for r in b[b.sheetnames[0]].iter_rows(values_only=True)]
    b.close()
    if s(g[0][0]).lower() == "track name":
        print(f"skip catalog: {os.path.basename(p)}")
        continue
    md = parse_filename(p)
    hdr = [s(x).upper() for x in g[3]]
    schema = "Overseas" if ("RADIO" in hdr and "TV" in hdr) else "Standard"
    amt_col = 16 if schema == "Overseas" else 11
    tr = next(i for i, r in enumerate(g) if any(s(c) == "TOTAL ROYALTIES" for c in r))
    total = num(g[tr][amt_col])
    works, lines, cur = set(), 0, None
    for r in g[4:tr]:
        if s(r[0]) != "":
            cur = s(r[0])
        keep = s(r[10]) != "" if schema == "Standard" else (
            s(r[4]) != "" and s(r[5]) != "" and s(r[amt_col]) != "")
        if keep:
            lines += 1
            works.add(cur)
    md.update(total=total, works=len(works), lines=lines, schema=schema)
    stmts.append(md)
stmts.sort(key=lambda m: m["s_no"])
grand_stmt = round(sum(m["total"] for m in stmts), 2)

wb2 = openpyxl.Workbook()
st = wb2.active
st.title = "Statement-V1"
SH = ["S. No", "Statement File", "Distribution Number", "Source / Category", "Period Covered",
      "Financial Year(s) Covered", "Statement Date", "Works", "Royalty Lines", "Total Amount",
      "% of Grand Total"]
t = st.cell(1, 1, value="STATEMENT SUMMARY  -  source 2 (input/batch-1), one row per royalty "
                        "distribution statement")
t.font = Font(bold=True, size=13)
t.fill = HEAD_FILL
st.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SH))
for i, h in enumerate(SH, start=1):
    c = st.cell(2, i, value=h)
    c.font = BOLD
    c.fill = HEAD_FILL
    c.alignment = WRAP
    c.border = Border(top=THIN, bottom=MED, left=THIN, right=THIN)

r = 3
for m in stmts:
    w = split_period(m["period"])
    fyl = ", ".join(fy_label(y) for y in sorted(w)) if w else "Not specified"
    vals = [m["sheet"], m["file"], m["dist_no"], m["section"], m["period"], fyl, m["date"],
            m["works"], m["lines"], round(m["total"], 2),
            m["total"] / grand_stmt if grand_stmt else 0]
    band = PatternFill("solid", fgColor="F2F2F2") if r % 2 else None
    for j, v in enumerate(vals, start=1):
        c = st.cell(r, j, value=v)
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        if band:
            c.fill = band
        if j == 7:
            c.number_format = "DD-MMM-YYYY"
        if j in (8, 9):
            c.alignment = Alignment(horizontal="center")
        if j == 10:
            c.number_format = MONEY
        if j == 11:
            c.number_format = "0.00%"
    r += 1
tot = r
vals = [f"TOTAL  ({len(stmts)} statements)", "", "", "", "", "", "",
        "", sum(m["lines"] for m in stmts), grand_stmt, 1.0]
for j, v in enumerate(vals, start=1):
    c = st.cell(tot, j, value=v)
    c.font = BOLD
    c.fill = PatternFill("solid", fgColor="C9B99B")
    c.border = Border(top=MED, bottom=MED)
    if j == 9:
        c.alignment = Alignment(horizontal="center")
    if j == 10:
        c.number_format = MONEY
    if j == 11:
        c.number_format = "0.00%"
st.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=7)

# ---- roll-ups
def block(title, hdrs, data, row0, cnt_total=None, amt_total=None):
    c = st.cell(row0, 1, value=title)
    c.font = Font(bold=True, size=12)
    c.fill = HEAD_FILL
    st.merge_cells(start_row=row0, start_column=1, end_row=row0, end_column=len(hdrs))
    for i, h in enumerate(hdrs, start=1):
        x = st.cell(row0 + 1, i, value=h)
        x.font = BOLD
        x.fill = HEAD_FILL
        x.alignment = WRAP
        x.border = Border(top=THIN, bottom=MED, left=THIN, right=THIN)
    rr = row0 + 2
    for i, vals in enumerate(data):
        f = PatternFill("solid", fgColor=FYFILL[i % len(FYFILL)])
        for j, v in enumerate(vals, start=1):
            x = st.cell(rr, j, value=v)
            x.fill = f
            x.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if isinstance(v, float) and j == len(vals) - 1:
                x.number_format = MONEY
            if j == len(vals):
                x.number_format = "0.00%"
            if j == 2:
                x.alignment = Alignment(horizontal="center")
        rr += 1
    for j, v in enumerate(["TOTAL",
                           len(stmts) if cnt_total is None else cnt_total,
                           grand_stmt if amt_total is None else amt_total, 1.0], start=1):
        x = st.cell(rr, j, value=v)
        x.font = BOLD
        x.fill = PatternFill("solid", fgColor="C9B99B")
        x.border = Border(top=MED, bottom=MED)
        if j == 2:
            x.alignment = Alignment(horizontal="center")
        if j == 3:
            x.number_format = MONEY
        if j == 4:
            x.number_format = "0.00%"
    return rr + 3


by_dist = defaultdict(lambda: [0, 0.0])
for m in stmts:
    d = by_dist[m["dist_no"] or "(none)"]
    d[0] += 1
    d[1] += m["total"]
data_d = [[k, v[0], round(v[1], 2), v[1] / grand_stmt] for k, v in
          sorted(by_dist.items(), key=lambda kv: -kv[1][1])]
nxt = block("BY DISTRIBUTION NUMBER", ["Distribution Number", "Statements", "Total Amount",
                                       "% of Grand Total"], data_d, tot + 3)

by_sec = defaultdict(lambda: [0, 0.0])
for m in stmts:
    d = by_sec[m["section"]]
    d[0] += 1
    d[1] += m["total"]
data_s = [[k, v[0], round(v[1], 2), v[1] / grand_stmt] for k, v in
          sorted(by_sec.items(), key=lambda kv: -kv[1][1])]
nxt = block("BY SOURCE / CATEGORY", ["Source / Category", "Statements", "Total Amount",
                                     "% of Grand Total"], data_s, nxt)

fy_amt, fy_cnt = defaultdict(float), defaultdict(set)
for m in stmts:
    w = split_period(m["period"])
    if not w:
        fy_amt["NA"] += m["total"]
        fy_cnt["NA"].add(m["s_no"])
    else:
        for y, f in w.items():
            fy_amt[y] += m["total"] * f
            fy_cnt[y].add(m["s_no"])
kk = sorted([k for k in fy_amt if k != "NA"]) + (["NA"] if "NA" in fy_amt else [])
fy_amt = penny_fix({k: fy_amt[k] for k in kk}, grand_stmt)
data_f = [[fy_label(k) if k != "NA" else "Period not stated", len(fy_cnt[k]), fy_amt[k],
           fy_amt[k] / grand_stmt] for k in kk]
nxt = block("BY EARNING PERIOD  (Indian FY, month-weighted split of statements spanning "
            "more than one year)", ["Financial Year", "Statements Covering It", "Total Amount",
                                    "% of Grand Total"], data_f, nxt, cnt_total=len(stmts))

recs = [("Grand total of the 45 statements (source 2)", grand_stmt),
        ("Grand total of the song grid on sheet SVF-RD-V4", grand_grid),
        ("Total of the Total Revenue (FY) columns on sheet SVF-RD-V4",
         round(sum(col_tot.values()), 2)),
        ("Difference (statements - song grid)", round(grand_stmt - grand_grid, 2))]
for i, (lbl, v) in enumerate(recs):
    c = st.cell(nxt + i, 1, value=lbl)
    c.font = BOLD if i == 3 else Font()
    st.merge_cells(start_row=nxt + i, start_column=1, end_row=nxt + i, end_column=2)
    x = st.cell(nxt + i, 3, value=v)
    x.number_format = MONEY
    x.font = BOLD
    if i == 3:
        x.fill = PatternFill("solid", fgColor="E8F2E8" if abs(v) <= TOL else "F7E7E7")

for w, cl in ((8, "A"), (74, "B"), (20, "C"), (28, "D"), (34, "E"), (28, "F"), (14, "G"),
              (9, "H"), (13, "I"), (16, "J"), (14, "K")):
    st.column_dimensions[cl].width = w
st.row_dimensions[2].height = 30
st.freeze_panes = "C3"
st.auto_filter.ref = f"A2:K{tot - 1}"

wb.save(DST)
wb2.save(DST2)
print(f"\nstatements={len(stmts)}  grand={grand_stmt:,.2f}  diff={grand_stmt - grand_grid:+.2f}")
print(f"saved -> {DST}\nsaved -> {DST2}")
