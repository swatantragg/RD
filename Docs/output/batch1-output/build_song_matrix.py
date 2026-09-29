"""SVF royalty -> single-sheet song x payment matrix (sections per revenue source)."""
import glob, os, re, zipfile, warnings, datetime as dt
warnings.filterwarnings("ignore")
import pandas as pd
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

IN_DIR = "/home/swatantra/RD/input/batch-1"
OUT_DIR = "/home/swatantra/RD/output"
OUT_XLSX = os.path.join(OUT_DIR, "SVF_Royalty_Song_Matrix.xlsx")
TOL = 0.05

SOCIETIES = {
    "008": ("APRA", "Australia"), "021": ("BMI", "USA"), "023": ("BUMA", "Netherlands"),
    "026": ("CASH", "Hong Kong"), "058": ("SACEM", "France"), "080": ("SUISA", "Switzerland"),
    "101": ("SOCAN", "Canada"), "104": ("MACP", "Malaysia"), "106": ("COMPASS", "Singapore"),
    "126": ("MCT", "Thailand"), "128": ("IMRO", "Ireland"),
}
MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
MLBL = {v: k.capitalize() for k, v in MONTHS.items()}


def s(v):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    return str(v).strip()


def num(v):
    try:
        f = float(v)
        return 0.0 if np.isnan(f) else f
    except (TypeError, ValueError):
        return 0.0


def month_end(y, m):
    return dt.date(y + (m == 12), 1 if m == 12 else m + 1, 1) - dt.timedelta(days=1)


def stmt_date(path):
    """Statement generation date = timestamp stamped inside the .xlsx package."""
    with zipfile.ZipFile(path) as z:
        ts = max(i.date_time for i in z.infolist())
    return dt.date(ts[0], ts[1], ts[2])


def parse_periods(stem):
    fy = ""
    m = re.search(r"F\.?\s?Y\.?\s*(\d{4})\s*-\s*(\d{2,4})", stem, re.I)
    if m:
        fy = f"{m.group(1)}-{m.group(2)[-2:]}"
    hits = []
    for mm in re.finditer(r"([A-Za-z]{3,12})\.?\s+(\d{4})", stem):
        k = mm.group(1)[:3].lower()
        if k in MONTHS:
            hits.append((MONTHS[k], int(mm.group(2))))
    if hits:
        return dt.date(hits[0][1], hits[0][0], 1), month_end(hits[-1][1], hits[-1][0]), fy
    if fy:
        y = int(fy.split("-")[0])
        return dt.date(y, 4, 1), month_end(y + 1, 3), fy
    return None, None, fy


def parse_filename(path):
    fn = os.path.basename(path)
    stem = os.path.splitext(fn)[0]
    n = int(re.search(r"\(S-(\d+)\)", fn).group(1))
    low = stem.lower()
    redis = "redistribution" in low
    runs = re.findall(r"\b([PM]\d{4})\b", stem)
    run = f"{runs[0]}-{runs[1]}" if (re.search(r"P\d{4}\s+to\s+P\d{4}", stem, re.I) and len(runs) > 1) \
        else (runs[0] if runs else "")
    soc_code = soc = country = ""
    ms = re.search(r"(\d{3})\s+([A-Z]{2,10})\s*$", stem.strip())
    if ms and ms.group(1) in SOCIETIES:
        soc_code = ms.group(1)
        soc, country = SOCIETIES[soc_code]
    if redis:
        cat = "Redistribution"
    elif "mechanical" in low:
        cat = "Mechanical (MUSERK)"
    elif "apple music" in low:
        cat = "Apple Music"
    elif "spotify" in low:
        cat = "Spotify"
    elif "youtube" in low and "pre" in low:
        cat = "YouTube Pre-Claims"
    elif "youtube" in low and "post" in low:
        cat = "YouTube Post-Claims"
    elif "facebook" in low or "meta" in low:
        cat = "Facebook / Meta"
    elif "radio" in low:
        cat = "Radio"
    elif "zee television" in low:
        cat = "Zee TV Broadcast"
    elif "overseas" in low:
        cat = "Overseas"
    else:
        cat = "Other / Unclassified"
    st, en, fy = parse_periods(stem)
    if st:
        plabel = f"{MLBL[st.month]} {st.year} - {MLBL[en.month]} {en.year}"
        if fy:
            plabel += f" (FY {fy})"
    else:
        plabel = "Not specified"
    section = f"Overseas - {soc} ({country})" if (cat == "Overseas" and soc) else cat
    return dict(s_no=n, sheet=f"S-{n}", file=fn, run=run, category=cat, section=section,
                p_start=st, p_end=en, fy=fy, period=plabel, soc=soc, country=country,
                date=stmt_date(path), redis=redis)


# ------------------------------------------------------------------ discover
paths = sorted(glob.glob(os.path.join(IN_DIR, "*.xlsx")),
               key=lambda p: int(re.search(r"\(S-(\d+)\)", os.path.basename(p)).group(1)))
catalog_path, money_paths = None, []
for p in paths:
    if s(pd.read_excel(p, header=None, nrows=1).iloc[0, 0]).lower() == "track name":
        catalog_path = p
    else:
        money_paths.append(p)
print(f"files={len(paths)}  catalog={os.path.basename(catalog_path)}  money={len(money_paths)}")

# ------------------------------------------------------------------ extract
rows, validation, groups = [], [], {}
for p in money_paths:
    md = parse_filename(p)
    df = pd.read_excel(p, sheet_name=0, header=None)
    hdr = [s(x).upper() for x in df.iloc[3].tolist()]
    schema = "Overseas" if ("RADIO" in hdr and "TV" in hdr) else "Standard"
    amt_col = 16 if schema == "Overseas" else 11
    member_no = s(df.iloc[0, 1])
    member_no = int(float(member_no)) if re.fullmatch(r"\d+(\.0+)?", member_no) else member_no
    member = s(df.iloc[1, 1]).rstrip(",")
    mask = df.apply(lambda r: r.astype(str).str.strip().eq("TOTAL ROYALTIES").any(), axis=1)
    tr = int(df.index[mask][0])
    file_total = num(df.iloc[tr, amt_col])

    det = df.iloc[4:tr].copy()
    det.columns = range(df.shape[1])
    for c in (0, 1, 2, 3):
        det[c] = det[c].ffill()
    keep = det[10].map(lambda v: s(v) != "") if schema == "Standard" else (
        det[4].map(lambda v: s(v) != "") & det[5].map(lambda v: s(v) != "")
        & det[amt_col].map(lambda v: s(v) != ""))
    money = det[keep]

    got = 0.0
    for _, r in money.iterrows():
        wid = s(r[0])
        try:
            wid = int(float(wid)) if wid else None
        except ValueError:
            wid = None
        amt = num(r[amt_col])
        got += amt
        rows.append(dict(gid=md["s_no"], section=md["section"], sheet=md["sheet"],
                         work=wid, title=s(r[1]), lang=s(r[3]), member=member,
                         member_no=member_no, source=s(r[10]) if schema == "Standard" else md["soc"],
                         amount=amt))
    groups[md["s_no"]] = dict(md, schema=schema, total=file_total, member=member,
                              member_no=member_no, lines=len(money))
    diff = got - file_total
    validation.append(dict(Sheet=md["sheet"], Section=md["section"], Schema=schema,
                           Lines=len(money), Extracted=got, FileTotal=file_total,
                           Diff=diff, Result="PASS" if abs(diff) <= TOL else "FAIL"))
    print(f"  {md['sheet']:<5} {schema:<8} n={len(money):<6} sum={got:>13.4f} file={file_total:>13.4f} "
          f"{'PASS' if abs(diff) <= TOL else 'FAIL'}  {md['section']}")

tx = pd.DataFrame(rows)
val = pd.DataFrame(validation)

# ------------------------------------------------------------------ catalog
cat = pd.read_excel(catalog_path, header=None, skiprows=1, names=["Track Name", "Internal NO", "ISRC"])
cat["Internal NO"] = pd.to_numeric(cat["Internal NO"], errors="coerce").astype("Int64")
cat = cat[cat["Internal NO"].notna()]
cat["Track Name"] = cat["Track Name"].map(s)
cat["ISRC"] = cat["ISRC"].map(s)


def isrc_union(series):
    out = []
    for v in series:
        for part in str(v).split("|"):
            part = part.strip()
            if part and part not in out:
                out.append(part)
    return " | ".join(out)


lu = (cat.groupby("Internal NO", sort=False)
      .agg(**{"Track Name": ("Track Name", "first"), "ISRC": ("ISRC", isrc_union)}).reset_index())

tx["work"] = tx["work"].astype("Int64")
tx = tx.merge(lu, left_on="work", right_on="Internal NO", how="left")
tx["In Catalog"] = np.where(tx["Track Name"].notna(), "Yes", "No")
tx["Song"] = tx["Track Name"].fillna("").where(lambda c: c != "", tx["title"])
tx["ISRC"] = tx["ISRC"].fillna("")

# ------------------------------------------------------------------ pivot grid
cell = tx.groupby(["work", "gid"], as_index=False)["amount"].sum()
songs = (tx.groupby("work", as_index=False)
         .agg(Song=("Song", "first"), ISRC=("ISRC", "first"), InCat=("In Catalog", "first"),
              Total=("amount", "sum")))
songs = songs.sort_values(["Total", "Song"], ascending=[False, True]).reset_index(drop=True)

SECTION_ORDER = ["YouTube Pre-Claims", "YouTube Post-Claims", "Facebook / Meta", "Spotify",
                 "Apple Music", "Radio", "Zee TV Broadcast", "Mechanical (MUSERK)",
                 "Other / Unclassified", "Redistribution"]
present = sorted({g["section"] for g in groups.values()})
ordered_sections = [x for x in SECTION_ORDER if x in present] + \
                   sorted([x for x in present if x not in SECTION_ORDER])

layout = []                                     # (section, [gid,...])
for sec in ordered_sections:
    gids = [g for g, m in groups.items() if m["section"] == sec]
    gids.sort(key=lambda g: (groups[g]["date"], g))
    layout.append((sec, gids))

# ------------------------------------------------------------------ write
wb = Workbook()
ws = wb.active
ws.title = "Royalty by Song"
BOLD = Font(bold=True)
SEC = Font(bold=True, size=14)
CTR = Alignment(horizontal="center")

FIXED = ["ISRC", "Song Name", "Internal No", "Total Amount"]
for i, h in enumerate(FIXED, start=1):
    c = ws.cell(row=2, column=i, value=h)
    c.font = BOLD

col = len(FIXED) + 1
spans, gid_col = [], {}
for sec, gids in layout:
    start = col
    for g in gids:
        gid_col[g] = col
        for j, h in enumerate(("Date", "Amount", "Period")):
            c = ws.cell(row=2, column=col + j, value=h)
            c.font = BOLD
        col += 3
    sc = ws.cell(row=1, column=start, value=sec.upper())
    sc.font = SEC
    sc.alignment = CTR
    if col - 1 > start:
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
    spans.append((sec, start, col - 1, len(gids)))
LAST_COL = col - 1

amt_by_work = {w: {} for w in songs["work"]}
for w, g, a in cell.itertuples(index=False):
    amt_by_work[w][g] = a

r = 3
for row in songs.itertuples(index=False):
    ws.cell(row=r, column=1, value=row.ISRC)
    ws.cell(row=r, column=2, value=row.Song)
    ws.cell(row=r, column=3, value=int(row.work))
    c = ws.cell(row=r, column=4, value=round(float(row.Total), 2))
    c.number_format = "#,##0.00"
    for g, amt in amt_by_work[row.work].items():
        base = gid_col[g]
        md = groups[g]
        d = ws.cell(row=r, column=base, value=md["date"])
        d.number_format = "DD-MMM-YYYY"
        a = ws.cell(row=r, column=base + 1, value=round(float(amt), 2))
        a.number_format = "#,##0.00"
        ws.cell(row=r, column=base + 2, value=md["period"])
    r += 1
LAST_ROW = r - 1

# totals row
tr_ = LAST_ROW + 2
tc = ws.cell(row=tr_, column=2, value="TOTAL")
tc.font = BOLD
c = ws.cell(row=tr_, column=4, value=round(float(songs["Total"].sum()), 2))
c.number_format = "#,##0.00"
c.font = BOLD
for g, base in gid_col.items():
    c = ws.cell(row=tr_, column=base + 1, value=round(float(
        cell[cell["gid"] == g]["amount"].sum()), 2))
    c.number_format = "#,##0.00"
    c.font = BOLD

ws.freeze_panes = "D3"
ws.auto_filter.ref = f"A2:{get_column_letter(LAST_COL)}{LAST_ROW}"
ws.column_dimensions["A"].width = 30
ws.column_dimensions["B"].width = 42
ws.column_dimensions["C"].width = 13
ws.column_dimensions["D"].width = 14
for _, start, end, _n in spans:
    for cc in range(start, end + 1, 3):
        ws.column_dimensions[get_column_letter(cc)].width = 13
        ws.column_dimensions[get_column_letter(cc + 1)].width = 12
        ws.column_dimensions[get_column_letter(cc + 2)].width = 26
wb.save(OUT_XLSX)

val.to_csv(os.path.join(OUT_DIR, "validation_song_matrix.csv"), index=False)

# ------------------------------------------------------------------ analysis
print("\n=== LAYOUT ===")
for sec, start, end, n in spans:
    print(f"  {sec:<34} cols {get_column_letter(start)}..{get_column_letter(end)}  "
          f"payments={n}  amount={tx[tx['section']==sec]['amount'].sum():>14,.2f}")
print(f"\nGrid: {LAST_ROW-2} song rows x {LAST_COL} columns  -> {OUT_XLSX}")
print(f"Grand total: {tx['amount'].sum():,.2f}   lines={len(tx):,}   works={tx['work'].nunique():,}")
print(f"Validation: {val['Result'].value_counts().to_dict()}  max|diff|={val['Diff'].abs().max():.6f}")
print(f"Sum of file TOTAL ROYALTIES = {val['FileTotal'].sum():,.2f}")
print("\nBy member:")
print(tx.groupby(["member_no", "member"])["amount"].sum().round(2).to_string())
print("\nCatalog coverage:")
w = tx.groupby("work").agg(inc=("In Catalog", "first"), amt=("amount", "sum"))
print(w.groupby("inc").agg(works=("amt", "size"), amount=("amt", "sum")).round(2).to_string())
print("\nTop 15 songs:")
print(songs.head(15)[["Song", "work", "Total", "InCat"]].to_string(index=False,
      formatters={"Total": lambda v: f"{v:,.2f}"}))
print("\nSongs by number of payments received:")
npay = cell.groupby("work").size()
print(npay.value_counts().sort_index().to_string())
