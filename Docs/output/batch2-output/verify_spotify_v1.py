"""Reconcile Spotify-RD-SKV1 / Spotify-Statement-V1 against the raw batch-2 statement
and against the batch-2 column inside SVF-RD-SKV5."""
import glob, os, warnings
warnings.filterwarnings("ignore")
import openpyxl

OUT = "/home/swatantra/RD/output/batch2-output"
s = lambda v: "" if v is None else str(v).strip()
n2 = lambda v: round(float(v or 0), 2)
ok = True


def chk(label, a, b, tol=0.0):
    global ok
    good = abs(round(a - b, 4)) <= tol
    ok &= good
    print(f"{'PASS' if good else 'FAIL'}  {label:<62} {a:>14,.2f} vs {b:>14,.2f}")


def flag(label, good, extra=""):
    global ok
    ok &= bool(good)
    print(f"{'PASS' if good else 'FAIL'}  {label}{('  ' + extra) if extra else ''}")


# ---- raw source
p = glob.glob("/home/swatantra/RD/input/batch-2/*.xlsx")[0]
wb = openpyxl.load_workbook(p, data_only=True)
g = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
wb.close()
tr = next(i for i, r in enumerate(g) if any(s(c) == "TOTAL ROYALTIES" for c in r))
SRC_TOTAL = round(float(g[tr][11]), 2)
raw, carry = {}, [None] * 4
lines = 0
for r in g[4:tr]:
    for c in (0, 1, 2, 3):
        if s(r[c]) != "":
            carry[c] = r[c]
    if s(r[10]) == "":
        continue
    lines += 1
    raw[int(float(s(carry[0])))] = raw.get(int(float(s(carry[0]))), 0.0) + float(r[11] or 0)
print(f"raw batch-2: total={SRC_TOTAL:,.2f}  works={len(raw)}  lines={lines}\n")

# ---- Spotify-RD-SKV1
G = openpyxl.load_workbook(os.path.join(OUT, "Spotify-RD-SKV1.xlsx"))["Spotify-RD-SKV1"]
TR = next(r for r in range(3, G.max_row + 1) if s(G.cell(r, 2).value) == "TOTAL")
LR = max(r for r in range(3, TR) if G.cell(r, 3).value is not None)
FY0 = next(c for c in range(5, G.max_column + 1)
           if str(G.cell(1, c).value or "").startswith("TOTAL REVENUE"))
FYH = [s(G.cell(2, c).value) for c in range(FY0, G.max_column + 1)]
grid = {G.cell(r, 3).value: n2(G.cell(r, 6).value) for r in range(3, LR + 1)}
flag(f"grid song rows == raw distinct works", len(grid) == len(raw), f"{len(grid)} / {len(raw)}")
flag("grid work ids == raw work ids", set(grid) == set(raw))
naive = round(sum(round(v, 2) for v in raw.values()), 2)
gap_p = int(round(abs(SRC_TOTAL - naive) / 0.01))
adj = [w for w in raw if grid.get(w) != round(raw[w], 2)]
off = [w for w in adj if abs(round(grid.get(w, 0) - round(raw[w], 2), 2)) > 0.01]
flag(f"every song amount is its raw sum rounded to 2dp, +/- at most one paisa  "
     f"({len(raw) - len(adj)} exact, {len(adj)} carrying the rounding remainder)", not off)
flag(f"paisa carriers == rounding gap of the naive sum "
     f"({naive:,.2f} -> {SRC_TOTAL:,.2f} = {gap_p}p)", len(adj) == gap_p)
for w in off[:5]:
    print("      ", w, grid.get(w), round(raw[w], 2))
chk("grid TOTAL row == printed TOTAL ROYALTIES", n2(G.cell(TR, 6).value), SRC_TOTAL)
chk("grid TOTAL row == sum of Amount column",
    n2(G.cell(TR, 6).value), round(sum(grid.values()), 2))
chk("col D grand == Amount column total", n2(G.cell(TR, 4).value), n2(G.cell(TR, 6).value))
badd = [r for r in range(3, LR + 1) if n2(G.cell(r, 4).value) != n2(G.cell(r, 6).value)]
flag(f"every row: Total Amount == its Spotify Amount  ({LR - 2 - len(badd)}/{LR - 2})", not badd)
print(f"      grid columns: {G.max_column}  FY headers: {FYH}")
badfy = 0
for r in range(3, LR + 1):
    parts = [n2(G.cell(r, FY0 + i).value) for i in range(len(FYH) - 1)]
    if abs(round(sum(parts), 2) - n2(G.cell(r, 4).value)) > 0.004 or \
       n2(G.cell(r, FY0 + len(FYH) - 1).value) != n2(G.cell(r, 4).value):
        badfy += 1
flag(f"every row's FY split == its Total Amount  ({LR - 2 - badfy}/{LR - 2})", not badfy)
for i, h in enumerate(FYH):
    chk(f"FY block '{h}' total == its column sum", n2(G.cell(TR, FY0 + i).value),
        round(sum(n2(G.cell(r, FY0 + i).value) for r in range(3, LR + 1)), 2))
chk("FY block Total Revenue == printed TOTAL ROYALTIES",
    n2(G.cell(TR, FY0 + len(FYH) - 1).value), SRC_TOTAL)

# ---- Spotify-Statement-V1
S = openpyxl.load_workbook(os.path.join(OUT, "Spotify-Statement-V1.xlsx"))["Spotify-Statement-V1"]
chk("statement row Total Amount == printed TOTAL ROYALTIES", n2(S.cell(3, 10).value), SRC_TOTAL)
chk("statement TOTAL row == printed TOTAL ROYALTIES", n2(S.cell(4, 10).value), SRC_TOTAL)
flag("statement Works == raw distinct works", S.cell(3, 8).value == len(raw),
     f"{S.cell(3, 8).value} / {len(raw)}")
flag("statement Royalty Lines == raw kept lines", S.cell(3, 9).value == lines,
     f"{S.cell(3, 9).value} / {lines}")


def rollup(title):
    r0 = next(r for r in range(1, S.max_row + 1) if s(S.cell(r, 1).value).startswith(title))
    r, tot = r0 + 2, 0.0
    while s(S.cell(r, 1).value) != "TOTAL":
        tot += n2(S.cell(r, 3).value)
        r += 1
    return round(tot, 2), n2(S.cell(r, 3).value)


for tl in ("BY DISTRIBUTION NUMBER", "BY SOURCE / CATEGORY", "BY EARNING PERIOD"):
    t, pt = rollup(tl)
    chk(f"roll-up '{tl}' rows sum == its TOTAL", t, pt)
    chk(f"roll-up '{tl}' TOTAL == printed statement total", pt, SRC_TOTAL)
t, pt = rollup("TOP 20 SONGS")
chk("roll-up 'TOP 20 SONGS' TOTAL == printed statement total", pt, SRC_TOTAL)
top20 = sorted(grid.values(), reverse=True)[:20]
chk("TOP 20 rows sum == top-20 amounts in the grid", t, round(sum(top20), 2))

# ---- consistency with the consolidated SVF-RD-SKV5 batch-2 column
C = openpyxl.load_workbook(os.path.join(OUT, "SVF-RD-SKV5.xlsx"))["SVF-RD-SKV5"]
CTR_ = next(r for r in range(3, C.max_row + 1) if s(C.cell(r, 2).value) == "TOTAL")
CLR = max(r for r in range(3, CTR_) if C.cell(r, 3).value is not None)
base = next(c for c in range(5, C.max_column + 1)
            if s(C.cell(CTR_, c + 3).value) == "Not specified"
            and s(C.cell(2, c).value) == "Date")
cons = {C.cell(r, 3).value: n2(C.cell(r, base + 1).value)
        for r in range(3, CLR + 1) if C.cell(r, base + 1).value is not None}
flag("SVF-RD-SKV5 batch-2 column covers the same works", set(cons) == set(grid),
     f"{len(cons)} / {len(grid)}")
diff = [w for w in grid if cons.get(w) != grid[w]]
flag(f"every amount identical in both workbooks  ({len(grid) - len(diff)}/{len(grid)})", not diff)
for w in diff[:5]:
    print("      ", w, grid[w], cons.get(w))
chk("SVF-RD-SKV5 batch-2 column total == this workbook's total",
    n2(C.cell(CTR_, base + 1).value), SRC_TOTAL)

print("\n" + ("ALL CHECKS PASSED" if ok else "*** FAILURES PRESENT ***"))
