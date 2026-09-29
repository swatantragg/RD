"""Reconcile SVF-RD-SKV5 / Statement-SKV2 against the raw source statements."""
import glob, os, re, warnings
warnings.filterwarnings("ignore")
import openpyxl
from collections import defaultdict

OUT = "/home/swatantra/RD/output/batch2-output"
G = openpyxl.load_workbook(os.path.join(OUT, "SVF-RD-SKV5.xlsx"))["SVF-RD-SKV5"]
S = openpyxl.load_workbook(os.path.join(OUT, "Statement-SKV2.xlsx"))["Statement-SKV2"]
s = lambda v: "" if v is None else str(v).strip()
n2 = lambda v: round(float(v or 0), 2)

TR = next(r for r in range(3, G.max_row + 1) if s(G.cell(r, 2).value) == "TOTAL")
FY0 = next(c for c in range(5, G.max_column + 1)
           if str(G.cell(1, c).value or "").startswith("TOTAL REVENUE"))
GLAST, LR = FY0 - 4, max(r for r in range(3, TR) if G.cell(r, 3).value is not None)
FYH = [s(G.cell(2, c).value) for c in range(FY0, G.max_column + 1)]
ok = True


def chk(label, a, b, tol=0.0):
    global ok
    good = abs(round(a - b, 4)) <= tol
    ok &= good
    print(f"{'PASS' if good else 'FAIL'}  {label:<62} {a:>15,.2f} vs {b:>15,.2f}")


# ---- 1. raw source totals per statement
src = {}
for p in glob.glob("/home/swatantra/RD/input/batch-1/*.xlsx") + \
         glob.glob("/home/swatantra/RD/input/batch-2/*.xlsx"):
    wb = openpyxl.load_workbook(p, data_only=True)
    g = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
    wb.close()
    if s(g[0][0]).lower() == "track name":
        continue
    hdr = [s(x).upper() for x in g[3]]
    ac = 16 if ("RADIO" in hdr and "TV" in hdr) else 11
    tr = next(i for i, r in enumerate(g) if any(s(c) == "TOTAL ROYALTIES" for c in r))
    src[os.path.basename(p)] = round(float(g[tr][ac]), 2)
print(f"source statements = {len(src)}   printed grand = {round(sum(src.values()), 2):,.2f}\n")

# ---- 2. Statement-SKV2 rows vs raw
rows = {}
r = 3
while s(S.cell(r, 1).value).startswith("S-"):
    rows[s(S.cell(r, 2).value)] = (n2(S.cell(r, 10).value), r)
    r += 1
TOTR = r
print(f"Statement rows = {len(rows)}")
miss = set(src) ^ set(rows)
print(f"{'PASS' if not miss else 'FAIL'}  file coverage identical            {miss if miss else ''}")
ok &= not miss
badr = [f for f in src if rows.get(f, (None,))[0] != src[f]]
print(f"{'PASS' if not badr else 'FAIL'}  every statement amount == its printed TOTAL ROYALTIES"
      f"  ({len(src) - len(badr)}/{len(src)})")
ok &= not badr
for f in badr:
    print("      ", f, rows.get(f), src[f])

chk("Statement TOTAL row == sum of its own rows",
    n2(S.cell(TOTR, 10).value), round(sum(v[0] for v in rows.values()), 2))
chk("Statement TOTAL row == sum of raw printed totals",
    n2(S.cell(TOTR, 10).value), round(sum(src.values()), 2))
chk("Statement TOTAL lines == sum of its own rows",
    S.cell(TOTR, 9).value, sum(S.cell(v[1], 9).value for v in rows.values()))

# ---- 3. grid: per-distribution column totals vs raw printed totals
blocks = []
ban = [(c, G.cell(1, c).value) for c in range(5, GLAST + 1) if G.cell(1, c).value]
for i, (c, v) in enumerate(ban):
    end = ban[i + 1][0] - 1 if i + 1 < len(ban) else GLAST
    for b in range(c, end + 1, 4):
        blocks.append((str(v), b))
print(f"\ngrid distribution blocks = {len(blocks)}  (statements = {len(src)})")
ok &= len(blocks) == len(src)
colsum, bad = [], 0
for sec, b in blocks:
    t = n2(G.cell(TR, b + 1).value)
    live = round(sum(n2(G.cell(r, b + 1).value) for r in range(3, LR + 1)
                     if G.cell(r, b + 1).value is not None), 2)
    if abs(t - live) > 0.004:
        bad += 1
        print("   FAIL col total != column sum", sec, G.cell(TR, b + 3).value, t, live)
    colsum.append(t)
print(f"{'PASS' if not bad else 'FAIL'}  every TOTAL-row figure == its own column sum")
ok &= not bad
want = sorted(src.values())
got = sorted(colsum)
same = want == got
print(f"{'PASS' if same else 'FAIL'}  multiset of 46 column totals == 46 printed statement totals")
ok &= same
if not same:
    print("   only in grid:", [x for x in got if x not in want][:10])
    print("   only in src :", [x for x in want if x not in got][:10])

# ---- 4. grid internal consistency
chk("col D grand == sum of song rows",
    n2(G.cell(TR, 4).value), round(sum(n2(G.cell(r, 4).value) for r in range(3, LR + 1)), 2))
chk("col D grand == sum of all distribution column totals",
    n2(G.cell(TR, 4).value), round(sum(colsum), 2))
badrow = 0
for r in range(3, LR + 1):
    live = round(sum(n2(G.cell(r, b + 1).value) for _, b in blocks
                     if G.cell(r, b + 1).value is not None), 2)
    if abs(live - n2(G.cell(r, 4).value)) > 0.004:
        badrow += 1
print(f"{'PASS' if not badrow else 'FAIL'}  every song's Total Amount == sum of its distribution "
      f"cells  ({LR - 2 - badrow}/{LR - 2})")
ok &= not badrow

# ---- 5. FY revenue block
badfy = 0
for r in range(3, LR + 1):
    parts = [n2(G.cell(r, FY0 + i).value) for i in range(len(FYH) - 1)]
    if abs(round(sum(parts), 2) - n2(G.cell(r, FY0 + len(FYH) - 1).value)) > 0.004 or \
       abs(round(sum(parts), 2) - n2(G.cell(r, 4).value)) > 0.004:
        badfy += 1
print(f"{'PASS' if not badfy else 'FAIL'}  every song's FY split == its Total Amount "
      f"({LR - 2 - badfy}/{LR - 2})")
ok &= not badfy
for i, h in enumerate(FYH):
    t = n2(G.cell(TR, FY0 + i).value)
    live = round(sum(n2(G.cell(r, FY0 + i).value) for r in range(3, LR + 1)), 2)
    chk(f"FY block '{h}' total == its column sum", t, live)
chk("FY block Total Revenue == col D grand",
    n2(G.cell(TR, FY0 + len(FYH) - 1).value), n2(G.cell(TR, 4).value))
chk("FY year columns sum == FY Total Revenue",
    round(sum(n2(G.cell(TR, FY0 + i).value) for i in range(len(FYH) - 1)), 2),
    n2(G.cell(TR, FY0 + len(FYH) - 1).value))

# ---- 6. Statement roll-up blocks
def rollup(title):
    r0 = next(r for r in range(1, S.max_row + 1) if s(S.cell(r, 1).value).startswith(title))
    r = r0 + 2
    tot = 0.0
    n = 0
    while s(S.cell(r, 1).value) != "TOTAL":
        tot += n2(S.cell(r, 3).value)
        n += S.cell(r, 2).value or 0
        r += 1
    return round(tot, 2), n, n2(S.cell(r, 3).value), S.cell(r, 2).value


for tl in ("BY DISTRIBUTION NUMBER", "BY SOURCE / CATEGORY", "BY INPUT BATCH",
           "BY EARNING PERIOD"):
    t, n, pt, pn = rollup(tl)
    chk(f"roll-up '{tl}' rows sum == its TOTAL", t, pt)
    chk(f"roll-up '{tl}' == statement grand", pt, n2(S.cell(TOTR, 10).value))

# ---- 7. cross-file reconciliation footer
rec = {}
for r in range(TOTR, S.max_row + 1):
    if s(S.cell(r, 1).value).startswith(("Grand total", "Total of the", "Difference")):
        rec[s(S.cell(r, 1).value)] = n2(S.cell(r, 3).value)
for k, v in rec.items():
    print(f"      {k:<62} {v:>15,.2f}")
chk("footer difference (statements - song grid) is zero", rec.get(
    "Difference (statements - song grid)", 9e9), 0.0)

print("\n" + ("ALL CHECKS PASSED" if ok else "*** FAILURES PRESENT ***"))
