"""Spotify-RD-SKV1 + Spotify-Statement-V1: same layout as SVF-RD-SKV5 / Statement-SKV2,
built from input/batch-2 only (Spotify, Oct 2025 - Mar 2026)."""
import glob, os, re, zipfile, warnings, datetime as dt
from collections import defaultdict, OrderedDict
warnings.filterwarnings("ignore")
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

IN2 = "/home/swatantra/RD/input/batch-2"
CATALOG = "/home/swatantra/RD/input/batch-1/(S-46)SVF list of Song -April 2026.xlsx"
OUT_DIR = "/home/swatantra/RD/output/batch2-output"
OUT1 = os.path.join(OUT_DIR, "Spotify-RD-SKV1.xlsx")
OUT2 = os.path.join(OUT_DIR, "Spotify-Statement-V1.xlsx")
SH1, SH2 = "Spotify-RD-SKV1", "Spotify-Statement-V1"
TOL = 0.05

SOCIETIES = {
    "008": ("APRA", "Australia"), "021": ("BMI", "USA"), "023": ("BUMA", "Netherlands"),
    "026": ("CASH", "Hong Kong"), "058": ("SACEM", "France"), "080": ("SUISA", "Switzerland"),
    "101": ("SOCAN", "Canada"), "104": ("MACP", "Malaysia"), "106": ("COMPASS", "Singapore"),
    "126": ("MCT", "Thailand"), "128": ("IMRO", "Ireland"),
}
MONTHS = {m: i + 1 for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
MLBL = {v: k.capitalize() for k, v in MONTHS.items()}


def s(v):
    return "" if v is None else str(v).strip()


def num(v):
    try:
        f = float(v)
        return 0.0 if f != f else f
    except (TypeError, ValueError):
        return 0.0


def month_end(y, m):
    return dt.date(y + (m == 12), 1 if m == 12 else m + 1, 1) - dt.timedelta(days=1)


def stmt_date(path):
    with zipfile.ZipFile(path) as z:
        ts = max(i.date_time for i in z.infolist())
    return dt.date(ts[0], ts[1], ts[2])


def fin_year(d):
    return d.year if d.month >= 4 else d.year - 1


def fy_label(y):
    return f"FY {y}-{str(y + 1)[-2:]}"


def parse_periods(stem):
    fy = ""
    m = re.search(r"F\.?\s?Y\.?\s*(\d{4})\s*-\s*(\d{2,4})", stem, re.I)
    if m:
        fy = f"{m.group(1)}-{m.group(2)[-2:]}"
    hits = []
    for mm in re.finditer(r"([A-Za-z]{3,12})\.?\s+(\d{4})", stem):
        k = mm.group(1)[:3].lower()
        if k in MONTHS:
            hits.append((MONTHS[k], int(mm.group(2))))
    if hits:
        return dt.date(hits[0][1], hits[0][0], 1), month_end(hits[-1][1], hits[-1][0]), fy
    if fy:
        y = int(fy.split("-")[0])
        return dt.date(y, 4, 1), month_end(y + 1, 3), fy
    return None, None, fy


def parse_filename(path, s_no, batch):
    fn = os.path.basename(path)
    stem = os.path.splitext(fn)[0]
    low = stem.lower()
    redis = "redistribution" in low
    sub = re.findall(r"\b([PM]\d{4}[A-Z]\d{3})\b", stem)
    runs = re.findall(r"\b([PM]\d{4})\b", stem)
    if sub:
        dist_no = sub[0]
    elif re.search(r"[PM]\d{4}\s+to\s+[PM]\d{4}", stem, re.I) and len(runs) > 1:
        dist_no = f"{runs[0]} to {runs[1]}"
    else:
        dist_no = runs[0] if runs else "Not specified"
    soc = country = ""
    ms = re.search(r"(\d{3})\s+([A-Z]{2,10})\s*$", stem.strip())
    if ms and ms.group(1) in SOCIETIES:
        soc, country = SOCIETIES[ms.group(1)]
    if redis:
        cat = "Redistribution"
    elif "mechanical" in low:
        cat = "Mechanical (MUSERK)"
    elif "apple music" in low:
        cat = "Apple Music"
    elif "spotify" in low:
        cat = "Spotify"
    elif "youtube" in low and "pre" in low:
        cat = "YouTube Pre-Claims"
    elif "youtube" in low and "post" in low:
        cat = "YouTube Post-Claims"
    elif "facebook" in low or "meta" in low:
        cat = "Facebook / Meta"
    elif "radio" in low:
        cat = "Radio"
    elif "zee television" in low:
        cat = "Zee TV Broadcast"
    elif "overseas" in low:
        cat = "Overseas"
    else:
        cat = "Other / Unclassified"
    st, en, fy = parse_periods(stem)
    if st:
        plabel = f"{MLBL[st.month]} {st.year} - {MLBL[en.month]} {en.year}"
        if fy:
            plabel += f" (FY {fy})"
    else:
        plabel = "Not specified"
    section = f"Overseas - {soc} ({country})" if (cat == "Overseas" and soc) else cat
    return dict(s_no=s_no, sheet=f"S-{s_no}", file=fn, batch=batch, dist_no=dist_no,
                category=cat, section=section, p_start=st, p_end=en, fy=fy, period=plabel,
                soc=soc, country=country, date=stmt_date(path), redis=redis)


def grid(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def fy_months(st, en):
    out = defaultdict(int)
    y, m = st.year, st.month
    while (y, m) <= (en.year, en.month):
        out[fin_year(dt.date(y, m, 1))] += 1
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return out


def split_period(label):
    lab = re.sub(r"\(.*?\)", "", s(label)).strip()
    hits = []
    for mm in re.finditer(r"([A-Za-z]{3,12})\.?\s+(\d{4})", lab):
        k = mm.group(1)[:3].lower()
        if k in MONTHS:
            hits.append((MONTHS[k], int(mm.group(2))))
    if not hits:
        return {}
    st = dt.date(hits[0][1], hits[0][0], 1)
    en = month_end(hits[-1][1], hits[-1][0])
    mm = fy_months(st, en)
    tot = sum(mm.values())
    return {y: k / tot for y, k in mm.items()}


def penny_fix(parts, target):
    out = {k: round(v, 2) for k, v in parts.items()}
    gap = round(target - round(sum(out.values()), 2), 2)
    if gap and out:
        big = max(out, key=lambda k: abs(out[k]))
        out[big] = round(out[big] + gap, 2)
    return out


# ------------------------------------------------------------------ discover
b2 = sorted(glob.glob(os.path.join(IN2, "*.xlsx")), key=lambda p: os.path.basename(p).lower())
money = []
for i, p in enumerate(b2):
    g = grid(p)
    if s(g[0][0]).lower() == "track name":
        continue
    money.append((p, i + 1, "batch-2", g))
print(f"batch-2 statements = {len(money)}")

# ------------------------------------------------------------------ extract
tx, validation, groups = [], [], {}
for p, n, b, g in money:
    md = parse_filename(p, n, b)
    hdr = [s(x).upper() for x in g[3]]
    schema = "Overseas" if ("RADIO" in hdr and "TV" in hdr) else "Standard"
    amt_col = 16 if schema == "Overseas" else 11
    tr = next(i for i, r in enumerate(g) if any(s(c) == "TOTAL ROYALTIES" for c in r))
    file_total = num(g[tr][amt_col])

    got, lines, works, carry = 0.0, 0, set(), [None, None, None, None]
    for r in g[4:tr]:
        for c in (0, 1, 2, 3):
            if s(r[c]) != "":
                carry[c] = r[c]
        keep = s(r[10]) != "" if schema == "Standard" else (
            s(r[4]) != "" and s(r[5]) != "" and s(r[amt_col]) != "")
        if not keep:
            continue
        raw = s(carry[0])
        try:
            wid = int(float(raw)) if raw else None
        except ValueError:
            wid = None
        if wid is None:
            wid = f"?{md['sheet']}:{raw or 'BLANK'}"
        amt = num(r[amt_col])
        got += amt
        lines += 1
        works.add(raw)
        tx.append((wid, md["s_no"], s(carry[1]), amt))
    groups[md["s_no"]] = dict(md, schema=schema, total=file_total, lines=lines,
                              works=len(works), extracted=got)
    diff = got - file_total
    validation.append((md["sheet"], md["batch"], md["section"], md["dist_no"], schema, lines,
                       got, file_total, diff, "PASS" if abs(diff) <= TOL else "FAIL"))
fails = [v for v in validation if v[9] == "FAIL"]
print(f"per-file extraction: {len(validation) - len(fails)} PASS / {len(fails)} FAIL  "
      f"max|diff|={max(abs(v[8]) for v in validation):.9f}")

# ------------------------------------------------------------------ catalog
catalog_rows = grid(CATALOG)
cat_name, cat_isrc = {}, {}
for t, nn, i in ((r + [None, None, None])[:3] for r in catalog_rows[1:]):
    try:
        k = int(float(s(nn)))
    except ValueError:
        continue
    if s(t) == "":
        continue
    cat_name.setdefault(k, s(t))
    seen = cat_isrc.setdefault(k, [])
    for part in s(i).split("|"):
        part = part.strip()
        if part and part not in seen:
            seen.append(part)

# ------------------------------------------------------------------ pivot
cell = defaultdict(float)
work_gids = defaultdict(set)
song = OrderedDict()
for wid, gid, title, amt in tx:
    cell[(wid, gid)] += amt
    work_gids[wid].add(gid)
    d = song.setdefault(wid, dict(title=title, total=0.0))
    d["total"] += amt
    if not d["title"]:
        d["title"] = title

by_gid = defaultdict(list)
for (w, g) in cell:
    by_gid[g].append(w)
cell_r = {}
for g, works_ in by_gid.items():
    target = round(sum(cell[(w, g)] for w in works_), 2)
    for w in works_:
        cell_r[(w, g)] = round(cell[(w, g)], 2)
    gap = round(target - round(sum(cell_r[(w, g)] for w in works_), 2), 2)
    if gap:
        step = 0.01 if gap > 0 else -0.01
        pool = [w for w in works_ if cell[(w, g)] != 0] or works_
        pool.sort(key=lambda w: (cell[(w, g)] - cell_r[(w, g)]), reverse=gap > 0)
        for i in range(int(round(abs(gap) / 0.01))):
            w = pool[i % len(pool)]
            cell_r[(w, g)] = round(cell_r[(w, g)] + step, 2)
_src_target = round(sum(round(g["total"], 2) for g in groups.values()), 2)
_gap0 = round(_src_target - round(sum(cell_r.values()), 2), 2)
if _gap0:
    _k = max(cell_r, key=lambda k: cell_r[k])
    cell_r[_k] = round(cell_r[_k] + _gap0, 2)
cell = cell_r

for wid in song:
    song[wid]["total"] = round(sum(cell.get((wid, g), 0.0) for g in work_gids[wid]), 2)
for wid, d in song.items():
    k = wid if isinstance(wid, int) else None
    d["name"] = cat_name.get(k) or d["title"]
    d["isrc"] = " | ".join(cat_isrc.get(k, []))
    d["incat"] = k in cat_name
order = sorted(song, key=lambda w: (-song[w]["total"], song[w]["name"]))

SECTION_ORDER = ["YouTube Pre-Claims", "YouTube Post-Claims", "Facebook / Meta", "Spotify",
                 "Apple Music", "Radio", "Zee TV Broadcast", "Mechanical (MUSERK)",
                 "Other / Unclassified", "Redistribution"]
present = sorted({g["section"] for g in groups.values()})
ordered_sections = [x for x in SECTION_ORDER if x in present] + \
                   sorted([x for x in present if x not in SECTION_ORDER])
layout = [(sec, sorted([g for g, m in groups.items() if m["section"] == sec],
                       key=lambda g: (groups[g]["date"], g))) for sec in ordered_sections]

PALETTE = {
    "Spotify":              ("BFD9BF", "E8F2E8"),
    "Apple Music":          ("E0B4B4", "F7E7E7"),
    "YouTube Pre-Claims":   ("E7C3D8", "F9E9F2"),
    "YouTube Post-Claims":  ("CBC1E0", "EEEAF7"),
    "Facebook / Meta":      ("B7CBE4", "E6EDF7"),
    "Radio":                ("E6CBA6", "F9F0E2"),
    "Zee TV Broadcast":     ("ADCECA", "E6F1EF"),
    "Mechanical (MUSERK)":  ("DAD7A0", "F3F2DF"),
    "Other / Unclassified": ("CFCFCF", "EFEFEF"),
    "Redistribution":       ("D5BFA7", "F2E9DE"),
}
OVERSEAS_CYCLE = [("B9C4D6", "E9EDF4"), ("C2CFB8", "ECF1E8"), ("DCC0C0", "F5EAEA"),
                  ("C9C2D9", "EFECF5"), ("DBCEB1", "F5F0E5"), ("B3CDD1", "E7F0F2"),
                  ("D2BFCE", "F2E9F0"), ("C9CDA8", "F0F1E3")]
sec_fill, k = {}, 0
for sec, _ in layout:
    if sec in PALETTE:
        sec_fill[sec] = PALETTE[sec]
    else:
        sec_fill[sec] = OVERSEAS_CYCLE[k % len(OVERSEAS_CYCLE)]
        k += 1

# ------------------------------------------------------------------ sheet 1
wb = Workbook()
ws = wb.active
ws.title = SH1
BOLD = Font(bold=True)
SECF = Font(bold=True, size=12, color="1F1F1F")
CTR = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
MED = Side(style="medium", color="808080")
THIN = Side(style="thin", color="BFBFBF")
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
BAND = PatternFill("solid", fgColor="F2F2F2")
MONEY = "#,##0.00"
DATEF = "DD-MMM-YYYY"

FIXED = ["ISRC", "Song Name", "Internal No", "Total Amount"]
BLOCK = ("Date", "Amount", "Period", "Distribution Number")
for i, h in enumerate(FIXED, start=1):
    c = ws.cell(row=2, column=i, value=h)
    c.font = BOLD
    c.fill = HEAD_FILL
    c.alignment = WRAP
    c.border = Border(top=THIN, bottom=MED, left=THIN, right=THIN)
ws.cell(row=1, column=1).border = Border(bottom=THIN)

col = len(FIXED) + 1
spans, gid_col = [], {}
for sec, gids in layout:
    start = col
    bfill = PatternFill("solid", fgColor=sec_fill[sec][0])
    for g in gids:
        gid_col[g] = col
        for j, h in enumerate(BLOCK):
            c = ws.cell(row=2, column=col + j, value=h)
            c.font = BOLD
            c.fill = bfill
            c.alignment = WRAP
            c.border = Border(top=THIN, bottom=MED, left=THIN, right=THIN)
        ws.cell(row=2, column=col + 3).border = Border(
            top=THIN, bottom=MED, left=THIN, right=Side(style="thin", color="808080"))
        col += 4
    end = col - 1
    sc = ws.cell(row=1, column=start, value=f"{sec.upper()}  ({len(gids)} distribution"
                                            f"{'s' if len(gids) > 1 else ''})")
    sc.font = SECF
    sc.alignment = CTR
    for cc in range(start, end + 1):
        ws.cell(row=1, column=cc).fill = bfill
    ws.cell(row=1, column=start).border = Border(left=MED, top=MED, bottom=THIN)
    ws.cell(row=1, column=end).border = Border(right=MED, top=MED, bottom=THIN)
    if end > start:
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
    spans.append((sec, start, end, len(gids)))
GRID_LAST_COL = col - 1

r = 3
for wid in order:
    d = song[wid]
    ws.cell(row=r, column=1, value=d["isrc"])
    ws.cell(row=r, column=2, value=d["name"])
    ws.cell(row=r, column=3, value=wid if isinstance(wid, int) else str(wid))
    c = ws.cell(row=r, column=4, value=round(d["total"], 2))
    c.number_format = MONEY
    c.font = BOLD
    if r % 2 == 1:
        for cc in range(1, 5):
            ws.cell(row=r, column=cc).fill = BAND
    for sec, start, end, _n in spans:
        dfill = PatternFill("solid", fgColor=sec_fill[sec][1])
        for cc in range(start, end + 1):
            ws.cell(row=r, column=cc).fill = dfill
        ws.cell(row=r, column=start).border = Border(left=MED)
        ws.cell(row=r, column=end).border = Border(right=MED)
    for g, base in gid_col.items():
        amt = cell.get((wid, g))
        if amt is None:
            continue
        md = groups[g]
        dc = ws.cell(row=r, column=base, value=md["date"])
        dc.number_format = DATEF
        ac = ws.cell(row=r, column=base + 1, value=amt)
        ac.number_format = MONEY
        ws.cell(row=r, column=base + 2, value=md["period"])
        ws.cell(row=r, column=base + 3, value=md["dist_no"])
    r += 1
GRID_LAST_ROW = r - 1

TOTAL_ROW = GRID_LAST_ROW + 2
tc = ws.cell(row=TOTAL_ROW, column=2, value="TOTAL")
tc.font = BOLD
grand = round(sum(d["total"] for d in song.values()), 2)
c = ws.cell(row=TOTAL_ROW, column=4, value=grand)
c.number_format = MONEY
c.font = BOLD
for cc in range(1, 5):
    ws.cell(row=TOTAL_ROW, column=cc).fill = HEAD_FILL
    ws.cell(row=TOTAL_ROW, column=cc).border = Border(top=MED, bottom=MED)
gid_total = defaultdict(float)
for (w, g), a in cell.items():
    gid_total[g] += a
for sec, start, end, _n in spans:
    bfill = PatternFill("solid", fgColor=sec_fill[sec][0])
    for cc in range(start, end + 1):
        ws.cell(row=TOTAL_ROW, column=cc).fill = bfill
        ws.cell(row=TOTAL_ROW, column=cc).border = Border(top=MED, bottom=MED)
    ws.cell(row=TOTAL_ROW, column=start).border = Border(top=MED, bottom=MED, left=MED)
    ws.cell(row=TOTAL_ROW, column=end).border = Border(top=MED, bottom=MED, right=MED)
for g, base in gid_col.items():
    c = ws.cell(row=TOTAL_ROW, column=base + 1, value=round(gid_total[g], 2))
    c.number_format = MONEY
    c.font = BOLD
    c = ws.cell(row=TOTAL_ROW, column=base + 3, value=groups[g]["dist_no"])
    c.font = BOLD

# ---- TOTAL REVENUE by Indian FY
rows_fy, fys = {}, set()
for rr in range(3, GRID_LAST_ROW + 1):
    raw = defaultdict(float)
    for cc in range(5, GRID_LAST_COL + 1, 4):
        amt = ws.cell(rr, cc + 1).value
        if amt is None:
            continue
        w = split_period(ws.cell(rr, cc + 2).value)
        if not w:
            raw["NA"] += num(amt)
        else:
            for y, f in w.items():
                raw[y] += num(amt) * f
    tgt = round(num(ws.cell(rr, 4).value), 2)
    rows_fy[rr] = penny_fix(raw, tgt) if raw else {}
    fys |= {x for x in raw if x != "NA"}

keys = sorted(fys) + (["NA"] if any("NA" in v for v in rows_fy.values()) else [])
headers = [(fy_label(x) if x != "NA" else "Period Not Stated") for x in keys] + ["Total Revenue"]

START = GRID_LAST_COL + 4
FYFILL = ["E8F2E8", "E6EDF7", "F9E9F2", "F5F0E5", "EEEAF7", "E7F0F2"]
BANNER = PatternFill("solid", fgColor="C9B99B")
sc = ws.cell(1, START, value="TOTAL REVENUE  (by Indian Financial Year, 1 Apr - 31 Mar, "
                             "of the period the royalty was earned in)")
sc.font = SECF
sc.alignment = CTR
for cc in range(START, START + len(headers)):
    ws.cell(1, cc).fill = BANNER
ws.merge_cells(start_row=1, start_column=START, end_row=1, end_column=START + len(headers) - 1)
ws.cell(1, START).border = Border(left=MED, top=MED, bottom=THIN)
ws.cell(1, START + len(headers) - 1).border = Border(right=MED, top=MED, bottom=THIN)
for j, h in enumerate(headers):
    c = ws.cell(2, START + j, value=h)
    c.font = BOLD
    c.fill = PatternFill("solid", fgColor="D8CBB3" if j < len(keys) else "C9B99B")
    c.alignment = WRAP
    c.border = Border(top=THIN, bottom=MED, left=THIN, right=THIN)

col_tot = defaultdict(float)
for rr in range(3, GRID_LAST_ROW + 1):
    d = rows_fy[rr]
    for j, x in enumerate(keys):
        c = ws.cell(rr, START + j, value=round(d.get(x, 0.0), 2))
        c.number_format = MONEY
        c.fill = PatternFill("solid", fgColor=FYFILL[j % len(FYFILL)])
        c.border = Border(left=THIN, right=THIN)
        col_tot[x] += d.get(x, 0.0)
    tcell = ws.cell(rr, START + len(keys), value=round(sum(d.values()), 2))
    tcell.number_format = MONEY
    tcell.font = BOLD
    tcell.fill = PatternFill("solid", fgColor="F0E9DC")
    ws.cell(rr, START).border = Border(left=MED)
    tcell.border = Border(right=MED, left=THIN)
for j, x in enumerate(keys):
    c = ws.cell(TOTAL_ROW, START + j, value=round(col_tot[x], 2))
    c.number_format = MONEY
    c.font = BOLD
    c.fill = BANNER
    c.border = Border(top=MED, bottom=MED)
gt = ws.cell(TOTAL_ROW, START + len(keys), value=round(sum(col_tot.values()), 2))
gt.number_format = MONEY
gt.font = BOLD
gt.fill = BANNER
gt.border = Border(top=MED, bottom=MED, right=MED)
ws.cell(TOTAL_ROW, START).border = Border(top=MED, bottom=MED, left=MED)
nc = ws.cell(TOTAL_ROW + 1, START,
             value="Basis: each distribution's amount is placed in the Indian FY (1 Apr - 31 Mar) "
                   "of the period it was earned in; a statement covering more than one FY is split "
                   "month-weighted. Row total = Total Amount (col D).")
nc.font = Font(italic=True, size=10, color="404040")
ws.merge_cells(start_row=TOTAL_ROW + 1, start_column=START, end_row=TOTAL_ROW + 1,
               end_column=START + len(headers) - 1)

ws.freeze_panes = "E3"
ws.auto_filter.ref = f"A2:{get_column_letter(START + len(headers) - 1)}{GRID_LAST_ROW}"
ws.row_dimensions[1].height = 22
ws.row_dimensions[2].height = 30
for w_, cl in ((34, "A"), (42, "B"), (13, "C"), (15, "D")):
    ws.column_dimensions[cl].width = w_
for _, start, end, _n in spans:
    for cc in range(start, end + 1, 4):
        ws.column_dimensions[get_column_letter(cc)].width = 13
        ws.column_dimensions[get_column_letter(cc + 1)].width = 12
        ws.column_dimensions[get_column_letter(cc + 2)].width = 26
        ws.column_dimensions[get_column_letter(cc + 3)].width = 15
for j in range(len(headers)):
    ws.column_dimensions[get_column_letter(START + j)].width = 16
for j in range(1, 4):
    ws.column_dimensions[get_column_letter(GRID_LAST_COL + j)].width = 3

# ------------------------------------------------------------------ sheet 2
stmts = sorted(groups.values(), key=lambda m: m["s_no"])
grand_stmt = round(sum(round(m["total"], 2) for m in stmts), 2)

wb2 = Workbook()
st = wb2.active
st.title = SH2
SH = ["S. No", "Statement File", "Distribution Number", "Source / Category", "Period Covered",
      "Financial Year(s) Covered", "Statement Date", "Works", "Royalty Lines", "Total Amount",
      "% of Grand Total"]
t = st.cell(1, 1, value=f"STATEMENT SUMMARY  -  input/batch-2, one row per royalty "
                        f"distribution statement ({len(stmts)} statement"
                        f"{'s' if len(stmts) > 1 else ''})")
t.font = Font(bold=True, size=13)
t.fill = HEAD_FILL
st.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(SH))
for i, h in enumerate(SH, start=1):
    c = st.cell(2, i, value=h)
    c.font = BOLD
    c.fill = HEAD_FILL
    c.alignment = WRAP
    c.border = Border(top=THIN, bottom=MED, left=THIN, right=THIN)

rr = 3
for m in stmts:
    w = split_period(m["period"])
    fyl = ", ".join(fy_label(y) for y in sorted(w)) if w else "Not specified"
    vals = [m["sheet"], m["file"], m["dist_no"], m["section"], m["period"], fyl, m["date"],
            m["works"], m["lines"], round(m["total"], 2),
            round(m["total"], 2) / grand_stmt if grand_stmt else 0]
    band = PatternFill("solid", fgColor="F2F2F2") if rr % 2 else None
    for j, v in enumerate(vals, start=1):
        c = st.cell(rr, j, value=v)
        c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        if band:
            c.fill = band
        if j == 7:
            c.number_format = "DD-MMM-YYYY"
        if j in (8, 9):
            c.alignment = Alignment(horizontal="center")
        if j == 10:
            c.number_format = MONEY
        if j == 11:
            c.number_format = "0.00%"
    rr += 1
tot = rr
vals = [f"TOTAL  ({len(stmts)} statement{'s' if len(stmts) > 1 else ''})", "", "", "", "", "", "",
        "", sum(m["lines"] for m in stmts), grand_stmt, 1.0]
for j, v in enumerate(vals, start=1):
    c = st.cell(tot, j, value=v)
    c.font = BOLD
    c.fill = PatternFill("solid", fgColor="C9B99B")
    c.border = Border(top=MED, bottom=MED)
    if j == 9:
        c.alignment = Alignment(horizontal="center")
    if j == 10:
        c.number_format = MONEY
    if j == 11:
        c.number_format = "0.00%"
st.merge_cells(start_row=tot, start_column=1, end_row=tot, end_column=7)


def block(title, hdrs, data, row0, cnt_total=None, amt_total=None):
    c = st.cell(row0, 1, value=title)
    c.font = Font(bold=True, size=12)
    c.fill = HEAD_FILL
    st.merge_cells(start_row=row0, start_column=1, end_row=row0, end_column=len(hdrs))
    for i, h in enumerate(hdrs, start=1):
        x = st.cell(row0 + 1, i, value=h)
        x.font = BOLD
        x.fill = HEAD_FILL
        x.alignment = WRAP
        x.border = Border(top=THIN, bottom=MED, left=THIN, right=THIN)
    r2 = row0 + 2
    for i, vals_ in enumerate(data):
        f = PatternFill("solid", fgColor=FYFILL[i % len(FYFILL)])
        for j, v in enumerate(vals_, start=1):
            x = st.cell(r2, j, value=v)
            x.fill = f
            x.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if isinstance(v, float) and j == len(vals_) - 1:
                x.number_format = MONEY
            if j == len(vals_):
                x.number_format = "0.00%"
            if j == 2:
                x.alignment = Alignment(horizontal="center")
        r2 += 1
    for j, v in enumerate(["TOTAL",
                           len(stmts) if cnt_total is None else cnt_total,
                           grand_stmt if amt_total is None else amt_total, 1.0], start=1):
        x = st.cell(r2, j, value=v)
        x.font = BOLD
        x.fill = PatternFill("solid", fgColor="C9B99B")
        x.border = Border(top=MED, bottom=MED)
        if j == 2:
            x.alignment = Alignment(horizontal="center")
        if j == 3:
            x.number_format = MONEY
        if j == 4:
            x.number_format = "0.00%"
    return r2 + 3


by_dist = defaultdict(lambda: [0, 0.0])
for m in stmts:
    d = by_dist[m["dist_no"] or "(none)"]
    d[0] += 1
    d[1] += round(m["total"], 2)
data_d = [[k_, v[0], round(v[1], 2), v[1] / grand_stmt] for k_, v in
          sorted(by_dist.items(), key=lambda kv: -kv[1][1])]
nxt = block("BY DISTRIBUTION NUMBER", ["Distribution Number", "Statements", "Total Amount",
                                       "% of Grand Total"], data_d, tot + 3)

by_sec = defaultdict(lambda: [0, 0.0])
for m in stmts:
    d = by_sec[m["section"]]
    d[0] += 1
    d[1] += round(m["total"], 2)
data_s = [[k_, v[0], round(v[1], 2), v[1] / grand_stmt] for k_, v in
          sorted(by_sec.items(), key=lambda kv: -kv[1][1])]
nxt = block("BY SOURCE / CATEGORY", ["Source / Category", "Statements", "Total Amount",
                                     "% of Grand Total"], data_s, nxt)

fy_amt, fy_cnt = defaultdict(float), defaultdict(set)
for m in stmts:
    w = split_period(m["period"])
    if not w:
        fy_amt["NA"] += round(m["total"], 2)
        fy_cnt["NA"].add(m["s_no"])
    else:
        for y, f in w.items():
            fy_amt[y] += round(m["total"], 2) * f
            fy_cnt[y].add(m["s_no"])
kk = sorted([x for x in fy_amt if x != "NA"]) + (["NA"] if "NA" in fy_amt else [])
fy_amt = penny_fix({x: fy_amt[x] for x in kk}, grand_stmt)
data_f = [[fy_label(x) if x != "NA" else "Period not stated", len(fy_cnt[x]), fy_amt[x],
           fy_amt[x] / grand_stmt] for x in kk]
nxt = block("BY EARNING PERIOD  (Indian FY, month-weighted split of statements spanning "
            "more than one year)", ["Financial Year", "Statements Covering It", "Total Amount",
                                    "% of Grand Total"], data_f, nxt, cnt_total=len(stmts))

# ---- top earners from this batch
top = order[:20]
data_t = [[song[w]["name"], w if isinstance(w, int) else str(w), song[w]["total"],
           song[w]["total"] / grand_stmt if grand_stmt else 0] for w in top]
nxt = block("TOP 20 SONGS IN THIS BATCH", ["Song Name", "Internal No", "Total Amount",
                                           "% of Grand Total"], data_t, nxt,
            cnt_total=len(song), amt_total=grand)

recs = [(f"Grand total of the {len(stmts)} batch-2 statement"
         f"{'s' if len(stmts) > 1 else ''}", grand_stmt),
        (f"Grand total of the song grid on sheet {SH1}", grand),
        (f"Total of the Total Revenue (FY) columns on sheet {SH1}",
         round(sum(col_tot.values()), 2)),
        ("Difference (statements - song grid)", round(grand_stmt - grand, 2))]
for i, (lbl, v) in enumerate(recs):
    c = st.cell(nxt + i, 1, value=lbl)
    c.font = BOLD if i == 3 else Font()
    st.merge_cells(start_row=nxt + i, start_column=1, end_row=nxt + i, end_column=2)
    x = st.cell(nxt + i, 3, value=v)
    x.number_format = MONEY
    x.font = BOLD
    if i == 3:
        x.fill = PatternFill("solid", fgColor="E8F2E8" if abs(v) <= TOL else "F7E7E7")

for w_, cl in ((8, "A"), (74, "B"), (20, "C"), (28, "D"), (34, "E"), (28, "F"), (14, "G"),
               (9, "H"), (13, "I"), (16, "J"), (14, "K")):
    st.column_dimensions[cl].width = w_
st.row_dimensions[2].height = 30
st.freeze_panes = "C3"
st.auto_filter.ref = f"A2:K{tot - 1}"

wb.save(OUT1)
wb2.save(OUT2)

with open(os.path.join(OUT_DIR, "validation_spotify_v1.csv"), "w") as f:
    f.write("Sheet,Batch,Section,DistributionNumber,Schema,Lines,Extracted,FileTotal,Diff,Result\n")
    for v in validation:
        f.write(",".join(str(x) if not isinstance(x, float) else f"{x:.6f}"
                         for x in (v[0], v[1], f'"{v[2]}"', v[3], v[4], v[5], v[6], v[7],
                                   v[8], v[9])) + "\n")

print(f"\ngrid: {GRID_LAST_ROW - 2} songs x {START + len(headers) - 1} cols")
print(f"song grid grand      = {grand:,.2f}")
print(f"statement grand      = {grand_stmt:,.2f}")
print(f"FY revenue columns   = {sum(col_tot.values()):,.2f}")
print(f"difference           = {grand_stmt - grand:+.2f}")
for x in keys:
    print(f"   {fy_label(x) if x != 'NA' else 'Not stated':<12} {round(col_tot[x], 2):>15,.2f}")
print(f"works not in catalog = {sum(1 for d in song.values() if not d['incat'])}")
print(f"-> {OUT1}\n-> {OUT2}")
