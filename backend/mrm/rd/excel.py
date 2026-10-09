"""Workbook writers: the master (SKV layout, measured from SVF-RD-SKV8.xlsx) and the two companion reports
(layout of Missing-Songs-Revenue-From-IPRS-And-Spotify.xlsx / Non-SVF-Works-Revenue-From-IPRS-And-Spotify.xlsx).

Measured SKV system: row 1 = merged section bands (bold 12 pt, text 1F1F1F, band fill), row 2 = headers (bold, centred,
wrapped; identity headers D9D9D9, block headers in the band fill; thin BFBFBF borders, medium 808080 bottom), data rows
tinted per band, identity columns striped F2F2F2, money '#,##0.00', dates 'DD-MMM-YYYY', a 3-column spacer before the
financial-year block, freeze F3, autofilter over the header row, TOTAL row of values, notes under it.

Exports are deterministic: no timestamps inside the workbook, fixed document dates, fixed ZIP entry times.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import zipfile
from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..normalize import title_key
from . import names
from .engine import Band, Result, Row
from .rules import FY_FILLS, STATUS_LABEL

ZERO = Decimal("0")
FIXED_TIME = dt.datetime(2000, 1, 1)
MONEY, DATEF = "#,##0.00", "DD-MMM-YYYY"
GREY_H, GREY_R = "D9D9D9", "F2F2F2"
FY_BAND, FY_HEAD, FY_TOTAL = "C9B99B", "D8CBB3", "F0E9DC"
AUD_BAND, AUD_TINT = "D8CBB7", "F0E9DC"
MED = Side(style="medium", color="808080")
THIN = Side(style="thin", color="BFBFBF")
SEP = Side(style="thin", color="808080")
BAND_F = Font(bold=True, size=12, color="1F1F1F")
BOLD = Font(bold=True)
NOTE_F = Font(italic=True, size=10, color="404040")
CTR = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
AUD_AL = Alignment(vertical="center", wrap_text=True)
_fills: dict[str, PatternFill] = {}


def fill(c: str) -> PatternFill:
    if c not in _fills:
        _fills[c] = PatternFill("solid", fgColor=c)
    return _fills[c]


def put(ws, r: int, c: int, v, fmt: str | None = None, font=None, fl: str | None = None, border=None, al=None):
    """Write one cell. Untrusted text that looks like a formula is stored as plain text (V3 §13.5)."""
    if isinstance(v, Decimal):
        v = names.q2(v)
    cell = ws.cell(r, c, v)
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        cell.data_type = "s"
    if fmt:
        cell.number_format = fmt
    if font is not None:
        cell.font = font
    if fl:
        cell.fill = fill(fl)
    if border is not None:
        cell.border = border
    if al is not None:
        cell.alignment = al
    return cell


def stable(wb: Workbook) -> bytes:
    wb.properties.created = wb.properties.modified = FIXED_TIME
    wb.properties.creator = "Royalty platform"
    wb.properties.lastModifiedBy = None
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in sorted(src.infolist(), key=lambda i: (i.filename != "[Content_Types].xml", i.filename)):
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                import re
                data = re.sub(rb"<dcterms:modified[^>]*>[^<]*</dcterms:modified>",
                              b'<dcterms:modified xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z</dcterms:modified>',
                              data)
            zi = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)
    return out.getvalue()


def _works_cell(works: list[str]):
    if not works:
        return None
    if len(works) == 1 and works[0].isdigit():
        return int(works[0])
    return " | ".join(works)


def _band_header(ws, a: int, z: int, text: str, fl: str) -> None:
    put(ws, 1, a, text, font=BAND_F, al=CTR)
    for k in range(a, z + 1):
        ws.cell(1, k).fill = fill(fl)
        ws.cell(1, k).border = Border(top=MED, bottom=THIN, left=MED if k == a else None,
                                      right=MED if k == z else None)
    if z > a:
        ws.merge_cells(start_row=1, start_column=a, end_row=1, end_column=z)


def _head(ws, c: int, text: str, fl: str, width: float | None = None, sep: bool = False) -> None:
    put(ws, 2, c, text, font=BOLD, fl=fl, al=WRAP,
        border=Border(left=THIN, right=SEP if sep else THIN, top=THIN, bottom=MED))
    if width:
        ws.column_dimensions[get_column_letter(c)].width = width


# ================================================================== master workbook
def master(res: Result, *, client_code: str, generation: str, state: str, build_ref: str, ruleset: dict,
           software: str, inputs: list[dict], left_out: list[dict], decisions: list[dict],
           people: dict | None = None) -> tuple[bytes, dict]:
    wb = Workbook()
    ws = wb.active
    ws.title = generation[:31]
    plat_bands = [b for b in res.bands if b.kind == "PLATFORM"]
    # ---- column plan
    ident = [("ISRC", 34), ("Song Name", 42), ("Internal No", 13), ("Total Amount", 15)] + \
            [(b.total_header, 21) for b in plat_bands]
    for c, (h, w) in enumerate(ident, 1):
        _head(ws, c, h, GREY_H, w)
    ws.cell(1, 1).border = Border(bottom=THIN)
    col = len(ident) + 1
    spans = []
    colpos: dict[str, int] = {}
    for b in res.bands:
        start = col
        widths = (13, 12, 26, 15) if b.kind == "ROYALTY" else (13, 15, 13, 26)
        for c in b.cols:
            colpos[c.key] = col
            for j, h in enumerate(("Date", "Amount", "Period", "Distribution Number")):
                _head(ws, col + j, h, b.fill, widths[j], sep=(j == 3))
            col += 4
        spans.append((b, start, col - 1))
        _band_header(ws, start, col - 1, b.label, b.fill)
    for j in range(3):
        ws.column_dimensions[get_column_letter(col + j)].width = 3
    fy_start = col + 3
    fy_heads = [names.fy_label(k) if k != "NA" else "Period Not Stated" for k in res.fy_keys] + ["Total Revenue"]
    for j, h in enumerate(fy_heads):
        _head(ws, fy_start + j, h, FY_HEAD if j < len(res.fy_keys) else FY_BAND, 16)
    fy_end = fy_start + len(fy_heads) - 1
    _band_header(ws, fy_start, fy_end, "TOTAL REVENUE  (by Indian Financial Year, 1 Apr - 31 Mar, of the period the "
                                       "royalty was earned in)", FY_BAND)
    aud_start = fy_end + 1
    aud = [("Catalogue Status", 40), ("Royalty Booked On", 58)] + \
          [(b.basis_header, 46) for b in plat_bands] + [("Held / Flags", 50)]
    for j, (h, w) in enumerate(aud):
        _head(ws, aud_start + j, h, AUD_BAND, w)
    aud_end = aud_start + len(aud) - 1
    _band_header(ws, aud_start, aud_end, "AUDIT / TRACEABILITY", AUD_BAND)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30

    # ---- rows
    values_hash = hashlib.sha256()
    first = 3
    for n, row in enumerate(res.rows):
        r = first + n
        stripe = GREY_R if r % 2 == 1 else None
        vals = [" | ".join(row.all_isrcs) or None, row.title, _works_cell(row.all_works), row.royalty] + \
               [row.platform.get(b.key, ZERO) for b in plat_bands]
        for c, v in enumerate(vals, 1):
            put(ws, r, c, v, fmt=MONEY if c >= 4 else None, font=BOLD if c >= 4 else None, fl=stripe)
        for b, a, z in spans:
            for k in range(a, z + 1):
                cell = ws.cell(r, k)
                cell.fill = fill(b.tint)
            for c in b.cols:
                p = colpos[c.key]
                ws.cell(r, p).number_format = DATEF
                ws.cell(r, p + 1).number_format = MONEY
                cl = row.cells.get(c.key)
                if cl is None or cl.sibling:
                    continue
                put(ws, r, p, c.date)
                put(ws, r, p + 1, cl.shown)
                put(ws, r, p + 2, c.period_label)
                put(ws, r, p + 3, c.dist_no)
            ws.cell(r, a).border = Border(left=MED)
            ws.cell(r, z).border = Border(right=MED)
        for j, k in enumerate(res.fy_keys):
            put(ws, r, fy_start + j, row.fy.get(k, Decimal("0.00")), fmt=MONEY, fl=FY_FILLS[j % len(FY_FILLS)],
                border=Border(left=MED) if j == 0 else Border(left=THIN, right=THIN))
        put(ws, r, fy_end, row.royalty, fmt=MONEY, font=BOLD, fl=FY_TOTAL, border=Border(left=THIN, right=MED))
        audit = [row.status, " ; ".join(row.booked_on) or None] + \
                [" ; ".join(dict.fromkeys(row.basis.get(b.key, []))) or None for b in plat_bands] + \
                [" ; ".join(dict.fromkeys(row.flags)) or None]
        for j, v in enumerate(audit):
            put(ws, r, aud_start + j, v, fl=AUD_TINT, al=AUD_AL,
                border=Border(left=MED) if j == 0 else (Border(right=MED) if j == len(audit) - 1 else None))
        values_hash.update(json.dumps([_jsonable(v) for v in vals + audit] +
                                      [str(row.cells[c.key].shown) if c.key in row.cells else None
                                       for b in res.bands for c in b.cols], ensure_ascii=False).encode())
    last = first + len(res.rows) - 1
    # ---- TOTAL row
    tr = last + 2
    put(ws, tr, 2, "TOTAL", font=BOLD)
    put(ws, tr, 4, sum((x.royalty for x in res.rows), ZERO), fmt=MONEY, font=BOLD)
    for j, b in enumerate(plat_bands):
        put(ws, tr, 5 + j, sum((x.platform.get(b.key, ZERO) for x in res.rows), ZERO), fmt=MONEY, font=BOLD)
    for c in range(1, len(ident) + 1):
        ws.cell(tr, c).fill = fill(GREY_H)
        ws.cell(tr, c).border = Border(top=MED, bottom=MED)
    for b, a, z in spans:
        for k in range(a, z + 1):
            ws.cell(tr, k).fill = fill(b.fill)
            ws.cell(tr, k).border = Border(top=MED, bottom=MED, left=MED if k == a else None,
                                           right=MED if k == z else None)
        for c in b.cols:
            p = colpos[c.key]
            put(ws, tr, p + 1, c.shown_total, fmt=MONEY, font=BOLD)
            put(ws, tr, p + 3, c.dist_no if b.kind == "ROYALTY" else c.period_label, font=BOLD)
    fy_tot = {k: sum((x.fy.get(k, ZERO) for x in res.rows), ZERO) for k in res.fy_keys}
    for j, k in enumerate(res.fy_keys):
        put(ws, tr, fy_start + j, fy_tot[k], fmt=MONEY, font=BOLD, fl=FY_BAND,
            border=Border(top=MED, bottom=MED, left=MED if j == 0 else None))
    put(ws, tr, fy_end, sum((x.royalty for x in res.rows), ZERO), fmt=MONEY, font=BOLD, fl=FY_BAND,
        border=Border(top=MED, bottom=MED, right=MED))
    for j in range(len(aud)):
        put(ws, tr, aud_start + j, f"{len(res.rows):,} song rows" if j == 0 else None, font=BOLD if j == 0 else None,
            fl=AUD_BAND, border=Border(top=MED, bottom=MED, left=MED if j == 0 else None,
                                       right=MED if j == len(aud) - 1 else None))
    put(ws, tr + 1, fy_start, "Basis: each distribution's amount is placed in the Indian FY (1 Apr - 31 Mar) of the "
                              "period it was earned in; a statement covering more than one FY is split month-weighted. "
                              "Row total = Total Amount (col D).", font=NOTE_F)
    note = provenance_note(res, generation, state, build_ref, ruleset, left_out, decisions)
    put(ws, tr + 3, 1, note, font=NOTE_F, al=Alignment(vertical="top", wrap_text=True))
    ws.merge_cells(start_row=tr + 3, start_column=1, end_row=tr + 3, end_column=aud_end)
    ws.row_dimensions[tr + 3].height = 132
    ws.freeze_panes = "F3" if plat_bands else "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(aud_end)}{max(last, 2)}"

    run_summary(wb, res, client_code=client_code, generation=generation, state=state, build_ref=build_ref,
                ruleset=ruleset, software=software, inputs=inputs, left_out=left_out, decisions=decisions,
                people=people or {})
    meta = rd_meta(wb, res, client_code=client_code, generation=generation, build_ref=build_ref, ruleset=ruleset,
                   software=software, values_sha=values_hash.hexdigest())
    return stable(wb), meta


def _jsonable(v):
    if isinstance(v, Decimal):
        return str(names.q2(v))
    if isinstance(v, (dt.date, dt.datetime)):
        return v.isoformat()
    return v


def provenance_note(res: Result, generation: str, state: str, build_ref: str, ruleset: dict, left_out: list[dict],
                    decisions: list[dict]) -> str:
    st = res.statements
    parts = [f"{generation} ({'DRAFT - not approved' if state != 'PUBLISHED' else 'published'}; build {build_ref}; "
             f"rule set v{ruleset.get('version')}).  Grain: {res.grain}."]
    parts.append(f"Built from {len(st)} royalty statement{'s' if len(st) != 1 else ''} "
                 f"({_m(sum((c.shown_total for b in res.bands if b.kind == 'ROYALTY' for c in b.cols if c.origin == 'source'), ZERO))} society royalty)"
                 + "".join(f", {p.platform} revenue report from {p.distributor} "
                           f"({_m(sum((c.shown_total for b in res.bands if b.kind == 'PLATFORM' for c in b.cols if c.file_id == p.file_id), ZERO))} "
                           f"{p.basis.lower()} revenue)" for p in res.platforms)
                 + (f", on the base: {res.base_label}" if res.base_label else ", with no song list (rows built from "
                    "the statements)") + ".")
    parts.append("Money is placed only through codes: royalty by IPRS work number, platform revenue by ISRC. A work "
                 "listed on several rows is booked once (column 'Royalty Booked On' names the row). A song title "
                 "never moves money without an approved decision. Society royalty (column D) and platform revenue "
                 "are different kinds of money and are never added together.")
    if res.totals.get("held_rows"):
        parts.append(f"{res.totals['held_rows']} held row(s) carry money that waits for a decision "
                     f"({_m(res.totals['held_royalty'])} royalty, {_m(res.totals['held_platform'])} platform revenue); "
                     "it is counted in every total.")
    if left_out:
        parts.append("Left out of this report: " + "; ".join(f"{x['label']} - {x['reason']}" for x in left_out[:12])
                     + ("; …" if len(left_out) > 12 else "") + ".")
    if decisions:
        parts.append(f"{len(decisions)} approved decision(s) applied (listed on the Run Summary sheet).")
    parts.append("Every figure is computed from the source files and re-checked before the file is written.")
    return "  ".join(parts)


def _m(v) -> str:
    from .detect import money
    return money(v)


# ================================================================== run summary sheet
def run_summary(wb: Workbook, res: Result, *, client_code, generation, state, build_ref, ruleset, software, inputs,
                left_out, decisions, people: dict | None = None) -> None:
    ws = wb.create_sheet("Run Summary")
    ws.column_dimensions["A"].width = 34
    for c, w in zip("BCDEFG", (60, 22, 18, 18, 16, 40)):
        ws.column_dimensions[c].width = w
    r = 1
    put(ws, r, 1, f"RUN SUMMARY  -  {generation}", font=Font(bold=True, size=13), fl=GREY_H)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
    r = 3
    kv = [("Client", client_code), ("Report", generation),
          ("State", "DRAFT - not approved for use" if state != "PUBLISHED" else "Published (approved)"),
          ("Build", build_ref)] + \
         [(label, (people or {}).get(key)) for key, label in (("prepared_by", "Prepared by"),
                                                              ("submitted_by", "Sent for approval by"),
                                                              ("approved_by", "Approved and published by"))
          if (people or {}).get(key)] + \
         [("Rule set", f"v{ruleset.get('version')}  ({rulesha(ruleset)[:12]})"),
          ("Software", software), ("Grain", res.grain), ("Base", res.base_label or "none"),
          ("Society royalty in this report", res.totals["royalty"])] + \
         [(f"{b.total_header}", res.totals["platform"].get(b.key, ZERO)) for b in res.bands if b.kind == "PLATFORM"] + \
         [("Song rows", res.totals["rows"]), ("Held rows (money waiting for a decision)", res.totals["held_rows"]),
          ("Paid songs not in the base (Missing Songs report)", res.totals["missing"])]
    for k, v in kv:
        put(ws, r, 1, k, font=BOLD)
        put(ws, r, 2, v, fmt=MONEY if isinstance(v, Decimal) else None)
        r += 1
    r += 1
    r = _table(ws, r, "FILES USED", ["File", "Type", "Detected as", "Total", "Check", "Included", "Fingerprint (sha256)"],
               [[x["filename"], x["kind"], x["label"], x.get("total"), x.get("check"), x.get("included"),
                 (x.get("sha256") or "")[:16]] for x in inputs], money_cols=(4,))
    if left_out:
        r = _table(ws, r, "LEFT OUT", ["File", "Reason", "", "Total"],
                   [[x["label"], x["reason"], None, x.get("total")] for x in left_out], money_cols=(4,))
    r = _table(ws, r, "CHECKS (every one must pass before a report is written)", ["Check", "Result", "Detail"],
               [[f"{c['id']}  {c['text']}", "PASS" if c["ok"] else "FAIL", c.get("detail") or ""] for c in res.checks])
    if decisions:
        r = _table(ws, r, "DECISIONS APPLIED", ["#", "Decision", "Made by", "Approved by", "Reason"],
                   [[d["id"], d["text"], d.get("made_by"), d.get("approved_by"), d.get("reason")] for d in decisions])
    if res.notices:
        r = _table(ws, r, "NOTICES", ["Notice"], [[n] for n in res.notices])


def _table(ws, r: int, title: str, heads: list[str], rows: list[list], money_cols=()) -> int:
    put(ws, r, 1, title, font=Font(bold=True, size=11))
    r += 1
    for j, h in enumerate(heads, 1):
        put(ws, r, j, h, font=BOLD, fl=GREY_H, al=WRAP, border=Border(top=THIN, bottom=MED, left=THIN, right=THIN))
    r += 1
    for i, row in enumerate(rows):
        for j, v in enumerate(row, 1):
            put(ws, r, j, v, fmt=MONEY if j in money_cols else None, fl=GREY_R if i % 2 else None,
                al=Alignment(wrap_text=True, vertical="top"))
        r += 1
    return r + 1


def rulesha(rs: dict) -> str:
    from .rules import ruleset_sha
    return ruleset_sha(rs)


def rd_meta(wb: Workbook, res: Result, *, client_code, generation, build_ref, ruleset, software, values_sha) -> dict:
    ws = wb.create_sheet("_rd_meta")
    ws.sheet_state = "hidden"
    meta = dict(workspace=client_code, build=build_ref, generation=generation, rule_set_version=ruleset.get("version"),
                rule_set_sha256=rulesha(ruleset), software_version=software, grain=res.grain,
                royalty_total=str(res.totals["royalty"]),
                platform_totals=json.dumps({k: str(v) for k, v in res.totals["platform"].items()}, sort_keys=True),
                rows=len(res.rows), values_sha256=values_sha)
    r = 1
    for k, v in meta.items():
        put(ws, r, 1, k)
        put(ws, r, 2, v)
        r += 1
    r += 1
    put(ws, r, 1, "row_key")
    put(ws, r, 2, "isrcs")
    put(ws, r, 3, "work_nos")
    r += 1
    for row in res.rows:
        put(ws, r, 1, row.key)
        put(ws, r, 2, "|".join(row.all_isrcs))
        put(ws, r, 3, "|".join(row.all_works))
        r += 1
    r += 1
    put(ws, r, 1, "column_key")
    put(ws, r, 2, "band")
    put(ws, r, 3, "distribution")
    put(ws, r, 4, "period")
    r += 1
    for b in res.bands:
        for c in b.cols:
            put(ws, r, 1, c.key)
            put(ws, r, 2, b.section)
            put(ws, r, 3, c.dist_no)
            put(ws, r, 4, c.period_label)
            r += 1
    return meta


# ================================================================== companion reports
def companion_groups(res: Result, only_status: str | None = None) -> list[dict]:
    """Paid-but-not-in-base rows, grouped so that one reviewed song (one row of the client's review list, which
    carries both the ISRC and the work number) is one line - exactly as the reference reports show it."""
    groups: dict = {}
    order = []
    for r in res.findings:
        st = r.review.status if r.review else "NEW"
        if only_status and st != only_status:
            continue
        key = ("mark", id(r.review)) if r.review else ("row", r.key)
        if key not in groups:
            groups[key] = dict(rows=[], review=r.review)
            order.append(key)
        groups[key]["rows"].append(r)
    out = []
    for k in order:
        g = groups[k]
        rows: list[Row] = g["rows"]
        g["isrcs"] = list(dict.fromkeys(i for x in rows for i in x.all_isrcs))
        g["works"] = list(dict.fromkeys(w for x in rows for w in x.all_works))
        g["title"] = rows[0].title
        for x in rows:
            if x.origin == "recording":
                g["title"] = x.title
                break
        g["royalty"] = sum((x.royalty for x in rows), ZERO)
        g["platform"] = {b.key: sum((x.platform.get(b.key, ZERO) for x in rows), ZERO)
                         for b in res.bands if b.kind == "PLATFORM"}
        g["status"] = g["review"].status if g["review"] else "NEW"
        out.append(g)
    from .rules import UNRESOLVED
    out.sort(key=lambda g: (g["status"] not in UNRESOLVED,
                            -(g["royalty"] + sum(g["platform"].values(), ZERO)), g["title"].upper()))
    return out


def companion(res: Result, groups: list[dict], *, kind: str, client_code: str, client_name: str, generation: str,
              state: str, combined_total: bool) -> bytes:
    """kind: MISSING | NON_CLIENT."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Missing Songs Revenue" if kind == "MISSING" else f"Non-{client_code} Works Revenue"[:31]
    rows_in = [x for g in groups for x in g["rows"]]
    used = {k for x in rows_in for k, c in x.cells.items() if not c.sibling}
    plat_bands = [b for b in res.bands if b.kind == "PLATFORM"]
    roy_bands = [Band(key=b.key, kind=b.kind, section=b.section, label=b.label, fill=b.fill, tint=b.tint,
                      cols=[c for c in b.cols if c.key in used]) for b in res.bands if b.kind == "ROYALTY"]
    roy_bands = [b for b in roy_bands if b.cols]
    letters = iter("ABCDEFGHIJKLMNOPQRSTUVWXYZ")
    plan = []
    for b in [*roy_bands, *plat_bands] if not roy_bands else _interleave(res.bands, roy_bands, plat_bands):
        letter = next(letters, "")
        if b.kind == "ROYALTY":
            pers = " and ".join(dict.fromkeys(c.period_label for c in b.cols))
            text = (f"SECTION {letter}  -  REVENUE FROM IPRS   ({b.section}: {len(b.cols)} royalty distribution"
                    f"{'s' if len(b.cols) != 1 else ''} - {pers})")
        else:
            months = [c.month for c in b.cols if c.month]
            per = names.period_label(min(months), names.month_end_of(max(months))) if months else ""
            text = (f"SECTION {letter}  -  REVENUE EARNED FROM {b.platform.upper()}   ({b.distributor} report - "
                    f"{'net' if b.basis == 'NET' else 'gross'} revenue {b.platform} reported against the ISRC, month "
                    f"by month, {per})")
        plan.append((b, text))
    head = [("ISRC", 34), ("Song Name", 44), ("Internal No", 13), ("Total Royalty From IPRS", 16)] + \
           [(f"Total Revenue From {b.platform} ({b.distributor})", 20) for b in plat_bands]
    for c, (h, w) in enumerate(head, 1):
        _head(ws, c, h, GREY_H, w)
    col = len(head) + 1
    colpos = {}
    spans = []
    for b, text in plan:
        a = col
        for c in b.cols:
            colpos[c.key] = col
            for j, (h, w) in enumerate(zip(("Date", "Amount", "Period", "Distribution Number"),
                                           (13, 13, 22 if b.kind == "ROYALTY" else 13, 18 if b.kind == "ROYALTY"
                                            else 24))):
                _head(ws, col + j, h, b.fill, w)
            col += 4
        spans.append((b, a, col - 1))
        _band_header(ws, a, col - 1, text, b.fill)
    # FY block
    fy_roy = sorted({k for g in groups for x in g["rows"] for k, v in x.fy.items()}, key=lambda k: (k == "NA", k))
    fy_plat: dict[str, list] = {}
    for b in plat_bands:
        fy_plat[b.key] = sorted({names.fin_year(c.month) for c in b.cols if c.month})
    tot_heads = [(f"{names.fy_label(k)} - IPRS Royalty" if k != "NA" else "Period Not Stated - IPRS Royalty", 17)
                 for k in fy_roy] + \
                [(f"{names.fy_label(y)} - {b.platform} Revenue", 18) for b in plat_bands for y in fy_plat[b.key]]
    if combined_total:
        tot_heads.append(("Total (Both Sources)", 17))
    fy_a = col
    for j, (h, w) in enumerate(tot_heads):
        _head(ws, col + j, h, FY_HEAD, w)
    col += len(tot_heads)
    fy_years = sorted({k for k in fy_roy if k != "NA"} | {y for v in fy_plat.values() for y in v})
    fy_text = ", ".join(names.fy_label(y) for y in fy_years)
    _band_header(ws, fy_a, col - 1, f"TOTAL REVENUE  (Indian Financial Year, 1 Apr - 31 Mar - {fy_text})"
                 if fy_years else "TOTAL REVENUE", FY_BAND)
    spans.append((Band(key="FY", kind="FY", section="", label="", fill=FY_BAND, tint="E8F2E8", cols=[]), fy_a,
                  col - 1))
    # audit
    with_note = any(g["review"] and g["review"].note for g in groups)
    aud = [("Status Given In Review", 22), ("Review Status (standard)", 22)] + \
          ([("Note From Review", 40)] if with_note else []) + \
          [("Present In The IPRS Statements", 46)] + [(f"{b.platform} ({b.distributor}) Match Basis", 34)
                                                       for b in plat_bands] + \
          [(f"In {client_code} Song List?", 16), ("Possible Match In Song List", 40)]
    au_a = col
    for j, (h, w) in enumerate(aud):
        _head(ws, col + j, h, AUD_BAND, w)
    col += len(aud)
    _band_header(ws, au_a, col - 1, "AUDIT / TRACEABILITY", AUD_BAND)
    spans.append((Band(key="AUD", kind="AUD", section="", label="", fill=AUD_BAND, tint=AUD_TINT, cols=[]), au_a,
                  col - 1))
    ncol = col - 1
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 32

    stmt_titles = _statement_titles(res)
    base_codes = {i for r in res.rows if r.in_base for i in r.all_isrcs + r.all_works}
    money_cols = [4] + [5 + j for j in range(len(plat_bands))]
    for n, g in enumerate(groups):
        r = 3 + n
        stripe = GREY_R if n % 2 == 0 else None
        vals = [" | ".join(g["isrcs"]) or None, g["title"], " | ".join(g["works"]) or None, g["royalty"]] + \
               [g["platform"][b.key] for b in plat_bands]
        for c, v in enumerate(vals, 1):
            put(ws, r, c, v, fmt=MONEY if c >= 4 else None, font=BOLD if c >= 4 else None, fl=stripe)
        for b, a, z in spans:
            for k in range(a, z + 1):
                ws.cell(r, k).fill = fill(b.tint)
            ws.cell(r, a).border = Border(left=MED)
        for b, _t in plan:
            for c in b.cols:
                p = colpos[c.key]
                ws.cell(r, p).number_format = DATEF
                ws.cell(r, p + 1).number_format = MONEY
                amt = sum((x.cells[c.key].shown for x in g["rows"] if c.key in x.cells and not x.cells[c.key].sibling),
                          ZERO)
                present = any(c.key in x.cells and not x.cells[c.key].sibling for x in g["rows"])
                if not present:
                    continue
                put(ws, r, p, c.date)
                put(ws, r, p + 1, amt)
                put(ws, r, p + 2, c.period_label)
                put(ws, r, p + 3, c.dist_no)
        tv = []
        for k in fy_roy:
            tv.append(sum((x.fy.get(k, ZERO) for x in g["rows"]), ZERO))
        for b in plat_bands:
            for y in fy_plat[b.key]:
                tv.append(sum((x.cells[c.key].shown for x in g["rows"] for c in b.cols
                               if c.key in x.cells and c.month and names.fin_year(c.month) == y), ZERO))
        if combined_total:
            tv.append(g["royalty"] + sum(g["platform"].values(), ZERO))
        for j, v in enumerate(tv):
            put(ws, r, fy_a + j, v, fmt=MONEY)
        rv = g["review"]
        av = [(rv.original or STATUS_LABEL[rv.status]) if rv else STATUS_LABEL["NEW"],
              STATUS_LABEL[rv.status] if rv else STATUS_LABEL["NEW"]] + \
             ([rv.note if rv else None] if with_note else []) + \
             [_present_in_iprs(g, stmt_titles)] + \
             [_platform_basis(g, b) for b in plat_bands] + \
             ["Yes" if any(code in base_codes for code in (rv.isrcs + rv.work_nos if rv else [])) else "No",
              _possible(g)]
        for j, v in enumerate(av):
            put(ws, r, au_a + j, v, al=Alignment(vertical="center", wrap_text=True))
    last = 2 + len(groups)
    tr = last + 1
    blurb = ("every paid song that is not in " + (res.base_label or "the base")) if kind == "MISSING" else \
            f"the songs reviewed as not belonging to {client_name}"
    put(ws, tr, 2, f"TOTAL  ({len(groups)} song{'s' if len(groups) != 1 else ''} - {blurb})")
    for b, a, z in [(None, 1, len(head))] + spans:
        for k in range(a, z + 1):
            cell = ws.cell(tr, k)
            cell.fill = fill(b.fill if b else GREY_H)
            cell.font = BOLD
            cell.border = Border(top=MED, left=MED if (b is not None and k == a) else None)
    sum_cols = money_cols + [colpos[c.key] + 1 for b, _ in plan for c in b.cols] + \
               list(range(fy_a, fy_a + len(tot_heads)))
    for c in sum_cols:
        L = get_column_letter(c)
        cell = put(ws, tr, c, f"=SUM({L}3:{L}{max(last, 3)})" if groups else 0, fmt=MONEY, font=BOLD)
        cell.data_type = "f" if groups else "n"
    note = companion_note(res, groups, kind, client_code, client_name, generation, state, combined_total)
    put(ws, tr + 2, 1, note, font=Font(size=10, color="404040"), al=Alignment(vertical="top", wrap_text=True))
    ws.merge_cells(start_row=tr + 2, start_column=1, end_row=tr + 2, end_column=ncol)
    ws.row_dimensions[tr + 2].height = 118
    ws.freeze_panes = "F3" if plat_bands else "E3"
    ws.auto_filter.ref = f"A2:{get_column_letter(ncol)}{max(last, 2)}"
    return stable(wb)


def _interleave(all_bands: list[Band], roy: list[Band], plat: list[Band]) -> list[Band]:
    keep = {b.key for b in roy}
    return [b for b in all_bands if (b.kind == "ROYALTY" and b.key in keep) or b.kind == "PLATFORM"]


def _statement_titles(res: Result) -> dict:
    out: dict = defaultdict(set)
    for st in res.statements:
        for ln in st.lines:
            if ln.work_no and ln.title:
                out["t:" + title_key(ln.title)].add(ln.work_no)
                out["b:" + title_key(_no_film(ln.title))].add(ln.work_no)
    return out


def _no_film(t: str) -> str:
    import re
    return re.sub(r"_?\(from[^)]*\)", "", (t or "").lower())


def _present_in_iprs(g: dict, idx: dict) -> str:
    paid = [x for x in g["rows"] if x.origin == "work"]
    if paid:
        return "Yes - work no " + ", ".join(g["works"]) + f", royalty {_m(g['royalty'])}"
    hits = sorted(idx.get("t:" + title_key(g["title"]), set()) | idx.get("b:" + title_key(_no_film(g["title"])),
                                                                           set()))
    if hits:
        return "Work no absent; song title matches IPRS work " + ", ".join(hits)
    return "No - work number, song title and title-without-film-name all absent"


def _platform_basis(g: dict, b: Band) -> str:
    n = sum(c.lines for x in g["rows"] for k, c in x.cells.items() if k in {cc.key for cc in b.cols})
    if n:
        return f"ISRC - exact match in the {b.distributor} report ({n} monthly line{'s' if n != 1 else ''})"
    return f"No ISRC match in the {b.distributor} report"


def _possible(g: dict) -> str | None:
    best = None
    for x in g["rows"]:
        for p in x.proposals:
            if p["band"] != "unlikely" and (best is None or Decimal(p["score"]) > Decimal(best["score"])):
                best = p
    if not best:
        return None
    codes = " | ".join(best["isrcs"][:2] + best["work_nos"][:1])
    return f"'{best['title']}' ({codes}) - {best['why']}"


def companion_note(res: Result, groups: list[dict], kind: str, client_code: str, client_name: str, generation: str,
                   state: str, combined_total: bool) -> str:
    n = len(groups)
    roy = sum((g["royalty"] for g in groups), ZERO)
    plat = sum((sum(g["platform"].values(), ZERO) for g in groups), ZERO)
    if kind == "MISSING":
        rule = (f"WHO IS ON THIS LIST: every song that earned money (or was named with 0.00) in a statement or platform "
                f"report of this run, but whose IPRS work number and ISRC are not in {res.base_label}. A song whose "
                "title matches a song-list entry is still listed, with that entry under 'Possible Match In Song List' "
                "- a title alone is not proof that two songs are the same. Usage-only songs (plays but no money) are "
                "not listed.  Review statuses come from the client's review and are carried by ISRC and work number, "
                "so a decision is never asked twice.  Unreviewed songs are listed first.")
    else:
        rule = (f"WHO IS ON THIS LIST: only songs that a reviewer marked as not belonging to {client_name} "
                "('Not ours'). No song is put here automatically: earning nothing, or being missing from the song "
                "list, does not make a song someone else's.")
    money_text = (f"Totals: {_m(roy)} society royalty from IPRS and {_m(plat)} platform revenue on {n} song"
                  f"{'s' if n != 1 else ''}. Every figure comes from the same checked data as {generation}; "
                  "a blank cell means no line exists in that source (it is not an inferred zero).")
    kinds = ("Royalty distributed by IPRS is the rights holder's share; platform revenue is what the platform "
             "reported. They are different kinds of money" + (" and are added together only in the column "
                                                               "'Total (Both Sources)'." if combined_total else "."))
    state_text = "" if state == "PUBLISHED" else "  DRAFT - not approved for use."
    return f"{rule}  {money_text}  {kinds}{state_text}"
