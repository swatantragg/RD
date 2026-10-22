"""Plain tabular Excel exports (catalogue, records, reconciliation report, error reports): a heading block that says
what was selected, one header row, frozen and filterable. Money keeps its exact value in the cell and is shown as
#,##0.00; a missing value stays an empty cell (never 0). Text from uploads cannot become a formula (excel.put)."""
from __future__ import annotations

import datetime as dt
import re
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .excel import put, stable

MONEY = "#,##0.00"
PCT = "0.00####"
DATE = "DD-MMM-YYYY"
HEAD_FILL = "1E5128"
HEAD_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13)
NOTE_FONT = Font(italic=True, size=10, color="404040")


def _cell(ws, r: int, c: int, v, kind: str | None):
    if isinstance(v, Decimal):
        cell = ws.cell(r, c, v)                 # exact value; the format rounds only what is shown
        cell.number_format = PCT if kind == "pct" else MONEY
        return cell
    if isinstance(v, (dt.date, dt.datetime)):
        cell = ws.cell(r, c, v)
        cell.number_format = DATE
        return cell
    if isinstance(v, (list, tuple)):
        v = " | ".join(str(x) for x in v if x not in (None, ""))
    return put(ws, r, c, v)


def sheet(wb: Workbook, title: str, columns: list[tuple], rows: list[list], heading: list[str] | None = None,
          first: bool = False) -> None:
    """columns: (header, width, kind) with kind None | 'money' | 'pct' | 'date' | 'text'."""
    ws = wb.active if first else wb.create_sheet()
    ws.title = re.sub(r"[\[\]:*?/\\]", "-", title)[:31]
    r = 1
    for i, line in enumerate(heading or []):
        put(ws, r, 1, line, font=TITLE_FONT if i == 0 else NOTE_FONT)
        r += 1
    if heading:
        r += 1
    head_row = r
    for c, (h, width, _k) in enumerate(columns, 1):
        cell = put(ws, r, c, h, font=HEAD_FONT)
        cell.fill = PatternFill("solid", fgColor=HEAD_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width or 14
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            _cell(ws, r, c, v, columns[c - 1][2] if c <= len(columns) else None)
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    if rows:
        ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(columns))}{r}"


def book(sheets: list[dict]) -> bytes:
    """sheets: [{title, columns, rows, heading}] -> deterministic workbook bytes."""
    wb = Workbook()
    for i, sh in enumerate(sheets):
        sheet(wb, sh["title"], sh["columns"], sh["rows"], sh.get("heading"), first=(i == 0))
    return stable(wb)


def safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_")
