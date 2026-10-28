"""What the screens do: clients, processing runs, files, checks, builds, review decisions, publishing, downloads.

Every function takes the signed-in person (actor), opens one transaction scoped to the clients that person may see,
checks the person's role, records an audit event and returns plain data for the API. Messages are plain language;
technical detail is kept for the diagnostics view.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
from collections import Counter
from decimal import Decimal

from .. import __version__, audit, config, db, storage
from ..security import Forbidden
from . import access, detect, names, rules
from .model import (KIND_LABEL, KINDS, PLATFORM_REVENUE, PREVIOUS_OUTPUT, REVIEW_LIST, SONG_LIST,
                    STATEMENT, UNKNOWN)

SOFTWARE = f"royalty-platform {__version__}"
ZERO = Decimal("0")
STAGES = ["Checking statements", "Extracting royalty entries", "Checking financial totals", "Matching song records",
          "Preparing reports"]


class RDError(Exception):
    """A plain-language refusal shown to the user (HTTP 400)."""


# ================================================================== clients
def _client(conn, client_id: int) -> dict:
    c = db.one(conn, "SELECT * FROM client WHERE id = %s", (client_id,))
    if not c:
        raise RDError("This client does not exist.")
    c["holders"] = db.all(conn, "SELECT id, society, member_no, name, ipi_name_no, ipi_base_no FROM client_holder "
                                "WHERE client_id = %s ORDER BY id", (client_id,))
    c["settings"] = c.get("settings") or {}
    return c


def client_cfg(conn, client_id: int) -> dict:
    c = _client(conn, client_id)
    s = c["settings"]
    return dict(code=c["code"], name=c["name"], holders=c["holders"], settings=s,
                tolerance=rules.RULESET_V1["recon_tolerance"])


def _client_out(c: dict, counts: dict | None = None) -> dict:
    s = c.get("settings") or {}
    return dict(id=c["id"], code=c["code"], name=c["name"], active=c.get("active", True), holders=c.get("holders", []),
                output_prefix=s.get("output_prefix") or f"{c['code']}-RD-SKV",
                next_generation=s.get("next_generation") or 1, distributors=s.get("distributors") or [],
                song_list_name=s.get("song_list_name") or f"{c['code']} list", **(counts or {}))


def list_clients(actor: dict) -> list[dict]:
    with access.tx(actor) as conn:
        ids = access.client_ids(conn, actor)
        out = []
        for cid in ids:
            c = _client(conn, cid)
            if not c.get("active", True) and not access.can(actor, "manage"):
                continue
            last = db.one(conn, "SELECT generation, published_at FROM rd_build WHERE client_id = %s AND state IN "
                                "('PUBLISHED','SUPERSEDED') ORDER BY generation_no DESC LIMIT 1", (cid,))
            out.append(_client_out(c, dict(last_published=last["generation"] if last else None,
                                           last_published_at=last["published_at"] if last else None)))
        return out


def save_client(actor: dict, body: dict, client_id: int | None = None) -> dict:
    """Preparers and administrators add clients (the person who adds one can work on it at once) and change their
    settings; switching a client off is for administrators."""
    access.require(actor, "clients")
    if "active" in body:
        access.require(actor, "manage")
    code = (body.get("code") or "").strip().upper()
    name = (body.get("name") or "").strip()
    with access.tx(actor) as conn:
        if client_id is None:
            if not code or not name:
                raise RDError("A client needs a short code (for example SVF) and a name.")
            if not code.replace("-", "").isalnum() or len(code) > 12:
                raise RDError("The client code may use letters, digits and '-', up to 12 characters.")
            if db.scalar(conn, "SELECT id FROM client WHERE code = %s", (code,)):
                raise RDError(f"A client with the code {code} already exists. Choose another code.")
            settings = _settings({}, body)
            settings.setdefault("output_prefix", f"{code}-RD-SKV")
            settings.setdefault("next_generation", 1)
            settings.setdefault("song_list_name", f"{code} list")
            client_id = db.insert(conn, "client", dict(code=code, name=name, kind="PUBLISHER_AND_LABEL", society="IPRS",
                                                       territories=["IN"], settings=settings))
            conn.execute("INSERT INTO client_grant (user_id, client_id, granted_by) VALUES (%s, %s, %s) ON CONFLICT "
                         "DO NOTHING", (actor["id"], client_id, actor["id"]))
            access.set_scope(conn, access.client_ids(conn, actor))
            audit.record(conn, "client.created", actor, "client", client_id, dict(code=code, name=name))
        else:
            access.require_client(conn, actor, client_id)
            c = _client(conn, client_id)
            vals = dict(settings=_settings(c["settings"], body))
            if name:
                vals["name"] = name
            if "active" in body:
                if not isinstance(body["active"], bool):
                    raise RDError("Say whether the client is switched on or off.")
                vals["active"] = body["active"]
            db.update(conn, "client", client_id, vals)
            audit.record(conn, "client.updated", actor, "client", client_id, dict(fields=sorted(body)))
        for h in body.get("holders") or []:
            if h.get("member_no") and h.get("name"):
                conn.execute("INSERT INTO client_holder (client_id, society, member_no, name, ipi_name_no, ipi_base_no) "
                             "VALUES (%s, %s, %s, %s, %s, %s) ON CONFLICT (client_id, society, member_no) DO UPDATE SET "
                             "name = EXCLUDED.name", (client_id, h.get("society") or "IPRS", str(h["member_no"]).strip(),
                                                      h["name"].strip(), h.get("ipi_name_no"), h.get("ipi_base_no")))
        for hid in body.get("remove_holders") or []:
            conn.execute("DELETE FROM client_holder WHERE id = %s AND client_id = %s", (hid, client_id))
        return _client_out(_client(conn, client_id))


def _settings(cur: dict, body: dict) -> dict:
    s = dict(cur or {})
    for k in ("output_prefix", "song_list_name"):
        if body.get(k):
            s[k] = str(body[k]).strip()
    if body.get("next_generation"):
        s["next_generation"] = int(body["next_generation"])
    if isinstance(body.get("distributors"), list):
        s["distributors"] = [dict(name=str(d["name"]).strip(), currency=(d.get("currency") or "INR").upper(),
                                  basis=(d.get("basis") or "GROSS").upper(), platforms=d.get("platforms") or [])
                             for d in body["distributors"] if d.get("name")]
    return s


# ================================================================== runs and files
def _run(conn, run_id: int) -> dict:
    r = db.one(conn, "SELECT * FROM rd_run WHERE id = %s", (run_id,))
    if not r:
        raise RDError("This processing run does not exist, or you do not have access to it.")
    if r["state"] == "DELETED":
        raise RDError("This processing run was deleted.")
    return r


def create_run(actor: dict, client_id: int, title: str | None = None) -> dict:
    access.require(actor, "prepare")
    with access.tx(actor, client_id) as conn:
        c = _client(conn, client_id)
        title = (title or "").strip() or f"{c['code']} - {dt.date.today().strftime('%d %b %Y')}"
        taken = {r["title"] for r in db.all(conn, "SELECT title FROM rd_run WHERE client_id = %s AND state <> 'DELETED' "
                                                    "AND (title = %s OR title LIKE %s)", (client_id, title, title + " (%"))}
        if title in taken:                            # runs must be told apart on the Home screen
            n = 2
            while f"{title} ({n})" in taken:
                n += 1
            title = f"{title} ({n})"
        rid = db.insert(conn, "rd_run", dict(client_id=client_id, title=title, created_by=actor["id"]))
        audit.record(conn, "run.created", actor, "rd_run", rid, dict(client=c["code"], title=title))
    return run_view(actor, rid)


def close_run(actor: dict, run_id: int) -> None:
    """Removes a run from the working lists. Its files, results and published reports are kept."""
    access.require(actor, "prepare")
    with access.tx(actor) as conn:
        _run(conn, run_id)
        busy = db.scalar(conn, "SELECT count(*) FROM rd_build WHERE run_id = %s AND state IN ('QUEUED','RUNNING') OR "
                               "(run_id = %s AND state = 'SUBMITTED')", (run_id, run_id))
        if busy:
            raise RDError("This run is being processed or published, so it cannot be closed now.")
        conn.execute("UPDATE rd_run SET state = 'ARCHIVED', updated_at = now() WHERE id = %s", (run_id,))
        audit.record(conn, "run.closed", actor, "rd_run", run_id)


def _delete_block(conn, run_id: int) -> str | None:
    """Why a run cannot be deleted, or None when it can."""
    bs = db.all(conn, "SELECT state, generation FROM rd_build WHERE run_id = %s ORDER BY version FOR UPDATE", (run_id,))
    published = [b["generation"] for b in bs if b["state"] in ("PUBLISHED", "SUPERSEDED")]
    if published:
        return (f"It holds the published report {', '.join(published)}. Published reports are kept for good; you can "
                f"close the run instead.")
    if any(b["state"] in ("QUEUED", "RUNNING") for b in bs):
        return "It is being processed. Please wait until that finishes."
    if any(b["state"] == "SUBMITTED" for b in bs):
        return "It is being published. Please wait until that finishes."
    return None


def delete_run(actor: dict, run_id: int) -> dict:
    """Deletes a run that was never published: the run and its draft reports disappear from every list. The uploaded
    files stay stored (added to a new run, the same copy is used) and review decisions stay with the client. Rows are
    marked, not erased, so the audit trail stays complete."""
    access.require(actor, "prepare")
    with access.tx(actor) as conn:
        conn.execute("SELECT id FROM rd_run WHERE id = %s FOR UPDATE", (run_id,))     # one delete at a time
        run = _run(conn, run_id)
        why = _delete_block(conn, run_id)
        if why:
            raise RDError(f"This run cannot be deleted. {why}")
        drafts = db.scalar(conn, "SELECT count(*) FROM rd_build WHERE run_id = %s", (run_id,))
        conn.execute("UPDATE rd_run SET state = 'DELETED', deleted_at = now(), deleted_by = %s, updated_at = now() "
                     "WHERE id = %s", (actor["id"], run_id))
        conn.execute("UPDATE rd_report SET available = FALSE, reason = 'This report was deleted with its run.' WHERE "
                     "NOT final AND build_id IN (SELECT id FROM rd_build WHERE run_id = %s)", (run_id,))
        audit.record(conn, "run.deleted", actor, "rd_run", run_id, dict(title=run["title"], drafts=drafts))
    return dict(ok=True, message=f"'{run['title']}' and its draft reports were deleted.")


def add_file(actor: dict, run_id: int, filename: str, data: bytes) -> dict:
    access.require(actor, "prepare")
    sha = hashlib.sha256(data).hexdigest()
    with access.tx(actor) as conn:
        run = _run(conn, run_id)
        cid = run["client_id"]
        existing = db.one(conn, "SELECT * FROM rd_file WHERE client_id = %s AND sha256 = %s", (cid, sha))
        c = _client(conn, cid)
    if existing:
        with access.tx(actor) as conn:
            had = db.one(conn, "SELECT removed FROM rd_run_file WHERE run_id = %s AND file_id = %s",
                         (run_id, existing["id"]))
            if had and not had["removed"]:
                return dict(file=_file_out(existing), note="This file is already in this run.")
            if had:
                conn.execute("UPDATE rd_run_file SET removed = FALSE WHERE run_id = %s AND file_id = %s",
                             (run_id, existing["id"]))
            else:
                db.insert(conn, "rd_run_file", dict(run_id=run_id, file_id=existing["id"], client_id=cid,
                                                    added_by=actor["id"]), returning=None)
            audit.record(conn, "file.reused", actor, "rd_file", existing["id"], dict(run=run_id, filename=filename))
            _touch(conn, run_id)
            return dict(file=_file_out(existing), note=f"This file was uploaded before (on "
                        f"{existing['uploaded_at']:%d %b %Y}); the same copy is used, so nothing is counted twice.")
    safe_name = "".join(ch for ch in filename if ch not in "/\\\0")[:200] or "upload.xlsx"
    parsed = detect.identify(safe_name, data)                     # outside the transaction: CPU only
    key = f"{c['code']}/{sha}{'.csv' if safe_name.lower().endswith('.csv') else '.xlsx'}"
    storage.put("originals", key, data)
    with access.tx(actor) as conn:
        cfg = client_cfg(conn, cid)
        facts = parsed.summary.get("facts", {})
        ev = detect.evaluate(parsed.kind, facts, safe_name, cfg, None, parsed.summary.get("file_date"))
        if parsed.layout in ("REJECTED", "DAMAGED"):
            ev["status"] = detect.UNSUPPORTED if parsed.layout == "REJECTED" else detect.ERROR
            ev["problems"] = [p.as_dict() for p in parsed.problems]
            ev["label"] = "Not a readable workbook"
        elif parsed.kind == UNKNOWN:
            ev["problems"] = [p.as_dict() for p in parsed.problems]
            ev["status"] = detect.NEEDS_REVIEW
        fid = db.insert(conn, "rd_file", dict(
            client_id=cid, sha256=sha, filename=safe_name, bytes=len(data), object_key=key, kind=parsed.kind,
            layout=parsed.layout, why=parsed.why, facts=db.J(facts), status=ev["status"], label=ev["label"],
            meta=db.J(ev["meta"]), meta_source=db.J(ev["meta_source"]), problems=db.J(ev["problems"]),
            summary=db.J(ev.get("summary") or {}),
            file_date=parsed.summary.get("file_date"), uploaded_by=actor["id"]))
        if parsed.data is not None:
            from .store import save_data
            save_data(conn, fid, cid, parsed.kind, parsed.data)
        db.insert(conn, "rd_run_file", dict(run_id=run_id, file_id=fid, client_id=cid, added_by=actor["id"]),
                  returning=None)
        audit.record(conn, "file.uploaded", actor, "rd_file", fid, dict(run=run_id, filename=safe_name, sha256=sha,
                                                                        kind=parsed.kind, status=ev["status"]))
        if parsed.kind == REVIEW_LIST and ev["status"] == detect.READY:
            _import_review_statuses(conn, actor, cid, run_id, fid, safe_name, parsed.data)
        _touch(conn, run_id)
        f = db.one(conn, "SELECT * FROM rd_file WHERE id = %s", (fid,))
    return dict(file=_file_out(f), note=None)


def _touch(conn, run_id: int) -> None:
    conn.execute("UPDATE rd_run SET updated_at = now() WHERE id = %s", (run_id,))


def _import_review_statuses(conn, actor: dict, client_id: int, run_id: int, file_id: int, filename: str, data) -> int:
    """A client's reviewed missing-song list becomes review decisions, keyed by ISRC and work number, so the same
    question is never asked again in later runs. The reviewer's own wording is kept."""
    n = 0
    label = names.short_label(filename)
    for e in data.entries:
        st = rules.map_status(e.status_text)
        if not st or not (e.isrcs or e.work_nos):
            continue
        _song_status(conn, actor, client_id, dict(isrcs=e.isrcs, work_nos=e.work_nos, title=e.title), st,
                     e.status_text, e.note, source=f"review list {label}, row {e.row_no}", run_id=run_id)
        n += 1
    audit.record(conn, "review.imported", actor, "rd_file", file_id, dict(statuses=n))
    return n


def _song_status(conn, actor, client_id, subject: dict, status: str, original: str, note: str, source: str,
                 run_id=None, build_id=None) -> int:
    codes = sorted(set(subject.get("isrcs") or []) | set(subject.get("work_nos") or []))
    key = "|".join(codes)
    conn.execute("UPDATE rd_decision SET state = 'SUPERSEDED' WHERE client_id = %s AND kind = 'SONG_STATUS' AND "
                 "subject_key = %s AND state = 'ACTIVE'", (client_id, key))
    return db.insert(conn, "rd_decision", dict(
        client_id=client_id, kind="SONG_STATUS", subject=subject, subject_key=key,
        value=dict(status=status, original=original or rules.STATUS_LABEL[status], note=note or ""),
        source=source, moves_money=False, state="ACTIVE", made_by=actor["id"], run_id=run_id, build_id=build_id))


def _file_out(f: dict, choice: dict | None = None) -> dict:
    probs = f.get("problems") or []
    questions = [p for p in probs if p.get("field")]
    return dict(id=f["id"], filename=f["filename"], kind=f["kind"], kind_label=KIND_LABEL.get(f["kind"], f["kind"]),
                label=f.get("label"), status=f["status"], status_text=detect.STATUS_TEXT.get(f["status"], f["status"]),
                problems=[dict(code=p["code"], severity=p["severity"], message=p["message"], detail=p.get("detail"))
                          for p in probs],
                questions=[dict(field=q["field"], message=q["message"], options=_options(q["field"]),
                                required=q["severity"] in ("REVIEW", "ERROR")) for q in questions],
                summary=f.get("summary") or {}, meta={k: v for k, v in (f.get("meta") or {}).items()
                                                      if k in ("section", "dist_no", "period_label", "holder_member",
                                                               "holder_name", "platform", "distributor", "currency",
                                                               "basis", "generation", "label", "date")},
                meta_source=f.get("meta_source") or {}, uploaded_at=f.get("uploaded_at"), bytes=f.get("bytes"),
                choice=choice or {})


def _options(field: str) -> list:
    return {"kind": [dict(value=k, label=KIND_LABEL[k]) for k in KINDS if k != UNKNOWN],
            "category": [dict(value=c, label=c) for c in detect.CATEGORIES],
            "currency": [dict(value=c, label=c) for c in ("INR", "USD", "EUR", "GBP")],
            "basis": [dict(value="GROSS", label="Gross (before the distributor's fee)"),
                      dict(value="NET", label="Net (after the distributor's fee)")],
            "accept_edited": [dict(value="yes", label="Yes, use it"), dict(value="", label="No")]}.get(field, [])


def answer(actor: dict, run_id: int, file_id: int, answers: dict) -> dict:
    """The user's answer to a question about one file (metadata, or the file type)."""
    access.require(actor, "prepare")
    with access.tx(actor) as conn:
        run = _run(conn, run_id)
        f = db.one(conn, "SELECT * FROM rd_file WHERE id = %s", (file_id,))
        if not f or f["client_id"] != run["client_id"]:
            raise RDError("This file is not part of this run.")
        if "kind" in answers and f["kind"] == UNKNOWN:
            kind = answers["kind"]
            if kind not in KINDS or kind == UNKNOWN:
                raise RDError("Please choose one of the listed file types.")
    if "kind" in answers and f["kind"] == UNKNOWN:
        data = storage.get("originals", f["object_key"])
        parsed = detect.identify(f["filename"], data, forced_kind=answers["kind"])
        with access.tx(actor) as conn:
            cfg = client_cfg(conn, run["client_id"])
            if parsed.kind == UNKNOWN:
                probs = [p.as_dict() for p in parsed.problems] or [dict(
                    code="KIND_NOT_READABLE", severity="ERROR", field=None, detail=parsed.why,
                    message=f"This file could not be read as a {KIND_LABEL[answers['kind']].lower()}. Please check "
                            "that it has a heading row with the song name and ISRC or work number columns.")]
                db.update(conn, "rd_file", file_id, dict(status=detect.UNSUPPORTED, problems=db.J(probs),
                                                         overrides=db.J(dict(kind=answers["kind"]))))
            else:
                facts = parsed.summary.get("facts", {})
                ev = detect.evaluate(parsed.kind, facts, f["filename"], cfg, None, parsed.summary.get("file_date"))
                db.update(conn, "rd_file", file_id, dict(
                    kind=parsed.kind, kind_source="you chose", layout=parsed.layout, why=parsed.why, facts=db.J(facts),
                    status=ev["status"], label=ev["label"], meta=db.J(ev["meta"]), meta_source=db.J(ev["meta_source"]),
                    problems=db.J(ev["problems"]), summary=db.J(ev.get("summary") or {})))
                from .store import save_data
                save_data(conn, file_id, run["client_id"], parsed.kind, parsed.data)
                if parsed.kind == REVIEW_LIST and ev["status"] == detect.READY:
                    _import_review_statuses(conn, actor, run["client_id"], run_id, file_id, f["filename"], parsed.data)
            audit.record(conn, "file.kind_chosen", actor, "rd_file", file_id, dict(kind=answers["kind"]))
            return _file_out(db.one(conn, "SELECT * FROM rd_file WHERE id = %s", (file_id,)))
    with access.tx(actor) as conn:
        cfg = client_cfg(conn, run["client_id"])
        ov = dict(f["overrides"] or {})
        for k, v in answers.items():
            if k in ("category", "dist_no", "overseas_society", "period", "date", "platform", "distributor",
                     "currency", "basis", "accept_edited"):
                ov[k] = str(v).strip() if v is not None else ""
        ev = detect.evaluate(f["kind"], f["facts"], f["filename"], cfg, ov,
                             f["file_date"].isoformat() if f["file_date"] else None)
        db.update(conn, "rd_file", file_id, dict(overrides=db.J(ov), status=ev["status"], label=ev["label"],
                                                 meta=db.J(ev["meta"]), meta_source=db.J(ev["meta_source"]),
                                                 problems=db.J(ev["problems"]), summary=db.J(ev.get("summary") or {})))
        audit.record(conn, "file.answered", actor, "rd_file", file_id, dict(answers=answers))
        _touch(conn, run_id)
        return _file_out(db.one(conn, "SELECT * FROM rd_file WHERE id = %s", (file_id,)))


def remove_file(actor: dict, run_id: int, file_id: int) -> None:
    access.require(actor, "prepare")
    with access.tx(actor) as conn:
        _run(conn, run_id)
        conn.execute("UPDATE rd_run_file SET removed = TRUE WHERE run_id = %s AND file_id = %s", (run_id, file_id))
        audit.record(conn, "file.removed_from_run", actor, "rd_file", file_id, dict(run=run_id))
        _touch(conn, run_id)


def choose(actor: dict, run_id: int, file_id: int, choice: str) -> None:
    """Answers to run-level questions: 'use both' (SUPPLEMENTARY), or leave one file out."""
    access.require(actor, "prepare")
    if choice not in ("SUPPLEMENTARY", "LEAVE_OUT", "USE"):
        raise RDError("Unknown choice.")
    with access.tx(actor) as conn:
        _run(conn, run_id)
        cur = db.one(conn, "SELECT choice FROM rd_run_file WHERE run_id = %s AND file_id = %s", (run_id, file_id))
        if cur is None:
            raise RDError("This file is not part of this run.")
        ch = dict(cur["choice"] or {})
        ch["relation"] = None if choice == "USE" else choice
        ch["by"] = actor["id"]
        conn.execute("UPDATE rd_run_file SET choice = %s WHERE run_id = %s AND file_id = %s",
                     (db.J(ch), run_id, file_id))
        audit.record(conn, "file.choice", actor, "rd_file", file_id, dict(run=run_id, choice=choice))
        _touch(conn, run_id)


def _run_files(conn, run_id: int) -> list[dict]:
    return db.all(conn, "SELECT f.*, rf.choice, rf.removed, rf.added_at FROM rd_run_file rf JOIN rd_file f ON "
                        "f.id = rf.file_id WHERE rf.run_id = %s ORDER BY rf.added_at, f.id", (run_id,))


def run_checks(conn, run: dict, files: list[dict]) -> dict:
    """Checks across the files of one run (duplicates, overlaps, which base to use). Returns the files that will be
    used, the ones left out (with the reason) and the questions still open."""
    live = [f for f in files if not f["removed"]]
    out_files, used, left, issues = [], [], [], []
    for f in live:
        ch = f["choice"] or {}
        if ch.get("relation") == "LEAVE_OUT":
            left.append(dict(file_id=f["id"], label=f["filename"], reason="you chose to leave it out",
                             total=(f.get("summary") or {}).get("total")))
            continue
        if f["status"] in (detect.ERROR, detect.UNSUPPORTED, detect.LEFT_OUT):
            reason = next((p["message"] for p in f["problems"] if p["severity"] in ("ERROR", "EXCLUDE")),
                          detect.STATUS_TEXT[f["status"]])
            left.append(dict(file_id=f["id"], label=f["filename"], reason=reason,
                             total=(f.get("summary") or {}).get("total")))
            continue
        used.append(f)
    # same distribution twice
    by_id: dict[str, list[dict]] = {}
    for f in used:
        if f["kind"] == STATEMENT and f["status"] == detect.READY:
            by_id.setdefault(f["meta"].get("identity_key", ""), []).append(f)
    for ident, fs in by_id.items():
        if len(fs) < 2:
            continue
        fs = sorted(fs, key=lambda x: x["id"])
        first = fs[0]
        for other in fs[1:]:
            if other["facts"].get("content_sha256") == first["facts"].get("content_sha256"):
                used.remove(other)
                left.append(dict(file_id=other["id"], label=other["filename"],
                                 reason=f"same distribution and the same entries as {first['filename']} - counted once",
                                 total=other["summary"].get("total")))
            elif (other["choice"] or {}).get("relation") != "SUPPLEMENTARY":
                issues.append(dict(kind="SAME_DISTRIBUTION", blocking=True, file_ids=[first["id"], other["id"]],
                                   title="Two files may contain the same distribution",
                                   message=f"{first['filename']} and {other['filename']} are both "
                                           f"{first['meta'].get('section')} {first['meta'].get('dist_no') or ''} "
                                           f"({first['meta'].get('period_label')}) for member "
                                           f"{first['meta'].get('holder_member')}, but their entries differ "
                                           f"({detect.money(Decimal(first['summary']['total'] or 0))} and "
                                           f"{detect.money(Decimal(other['summary']['total'] or 0))}). Are they "
                                           "separate payments, or is one a corrected copy of the other?",
                                   actions=[dict(action="SUPPLEMENTARY", file_id=other["id"],
                                                 label="Use both - they are separate payments"),
                                            dict(action="LEAVE_OUT", file_id=first["id"],
                                                 label=f"Use only {other['filename']}"),
                                            dict(action="LEAVE_OUT", file_id=other["id"],
                                                 label=f"Use only {first['filename']}")]))
    for kind, what in ((SONG_LIST, "song lists"), (PREVIOUS_OUTPUT, "previous master reports")):
        fs = [f for f in used if f["kind"] == kind and f["status"] == detect.READY]
        if len(fs) > 1:
            issues.append(dict(kind="CHOOSE_BASE", blocking=True, file_ids=[f["id"] for f in fs],
                               title=f"Choose one of the {what}",
                               message=f"{len(fs)} {what} were uploaded. Only one can be used in a run.",
                               actions=[dict(action="LEAVE_OUT", file_id=o["id"], label=f"Use {f['filename']}")
                                        for f in fs for o in fs if o["id"] != f["id"]][:6]))
    # already in the previous report
    prior = next((f for f in used if f["kind"] == PREVIOUS_OUTPUT and f["status"] == detect.READY), None)
    if prior:
        cols = prior["facts"].get("columns") or []
        for f in used:
            if f["kind"] != STATEMENT or f["status"] != detect.READY:
                continue
            m = f["meta"]
            hit = [c for c in cols if c["kind"] == "ROYALTY" and c["section"] == m.get("section")
                   and (c["dist_no"] or "").upper().replace("NOT SPECIFIED", "") == (m.get("dist_no") or "").upper()
                   and _same_period(c["period_label"], m.get("period_label"))]
            if not hit:
                continue
            tot = names.q2(Decimal(f["facts"].get("lines_total") or "0"))
            if any(names.q2(Decimal(c["total"] or "0")) == tot for c in hit):
                used.remove(f)
                left.append(dict(file_id=f["id"], label=f["filename"], total=f["summary"].get("total"),
                                 reason=f"already included in {prior['meta'].get('generation')} with the same total - "
                                        "not added again"))
            elif (f["choice"] or {}).get("relation") != "SUPPLEMENTARY":
                issues.append(dict(kind="ALREADY_IN_BASE", blocking=True, file_ids=[f["id"]],
                                   title=f"{f['filename']} may already be in {prior['meta'].get('generation')}",
                                   message=f"{prior['meta'].get('generation')} already contains {m.get('section')} "
                                           f"{m.get('dist_no') or ''} for {m.get('period_label')}, but with a "
                                           f"different total ({detect.money(Decimal(hit[0]['total'] or 0))} there, "
                                           f"{detect.money(tot)} in this file). It may be a re-issue.",
                                   actions=[dict(action="SUPPLEMENTARY", file_id=f["id"],
                                                 label="Add it - it is a separate payment"),
                                            dict(action="LEAVE_OUT", file_id=f["id"], label="Leave it out")]))
    # one rights holder per run
    holders = sorted({f["meta"].get("holder_member") for f in used if f["kind"] == STATEMENT and
                      f["status"] == detect.READY and f["meta"].get("holder_member")})
    if len(holders) > 1:
        issues.append(dict(kind="MIXED_HOLDERS", blocking=True, file_ids=[],
                           title="Statements for more than one rights holder",
                           message=f"This run contains statements for members {', '.join(holders)}. Process each "
                                   "rights holder in its own run so their money is never added together.",
                           actions=[]))
    needs = [f for f in used if f["status"] == detect.NEEDS_REVIEW]
    for f in out_files:
        pass
    money = [f for f in used if f["kind"] in (STATEMENT, PLATFORM_REVENUE) and f["status"] == detect.READY]
    blocked = []
    if needs:
        blocked.append(f"{len(needs)} file{'s need' if len(needs) != 1 else ' needs'} your answer before processing.")
    if any(i["blocking"] for i in issues):
        blocked.append("Resolve the questions about the files above.")
    if not money and not prior:
        blocked.append("Add at least one royalty statement or platform revenue report.")
    return dict(used=used, left_out=left, issues=issues, needs_answer=needs, blocked=blocked,
                can_process=not blocked)


def _same_period(a: str | None, b: str | None) -> bool:
    return names.parse_period_label(a or "") == names.parse_period_label(b or "")


def run_view(actor: dict, run_id: int) -> dict:
    with access.tx(actor) as conn:
        run = _run(conn, run_id)
        c = _client(conn, run["client_id"])
        files = _run_files(conn, run_id)
        chk = run_checks(conn, run, files)
        left_ids = {x["file_id"]: x["reason"] for x in chk["left_out"]}
        builds = db.all(conn, "SELECT id, version, state, created_at, finished_at, generation, error FROM rd_build "
                              "WHERE run_id = %s ORDER BY version DESC", (run_id,))
        latest = builds[0] if builds else None
        changed = latest is not None and _inputs_sha(conn, run, chk["used"]) != db.scalar(
            conn, "SELECT inputs_sha256 FROM rd_build WHERE id = %s", (latest["id"],))
        maker = db.one(conn, "SELECT name FROM app_user WHERE id = %s", (run["created_by"],))
        fo = []
        for f in files:
            if f["removed"]:
                continue
            o = _file_out(f, f["choice"])
            if f["id"] in left_ids and f["status"] == detect.READY:
                o["status"], o["status_text"] = "LEFT_OUT", "Not included"
                o["problems"] = [dict(code="LEFT_OUT", severity="INFO", message=left_ids[f["id"]], detail="")] + \
                    o["problems"]
            fo.append(o)
        counts = Counter(o["status"] for o in fo)
        return dict(id=run["id"], title=run["title"], client=_client_out(c), created_by=maker["name"] if maker else None,
                    created_at=run["created_at"], updated_at=run["updated_at"], files=fo,
                    counts=dict(total=len(fo), ready=counts.get("READY", 0), needs_review=counts.get("NEEDS_REVIEW", 0),
                                left_out=counts.get("LEFT_OUT", 0), error=counts.get("ERROR", 0) +
                                counts.get("UNSUPPORTED", 0)),
                    issues=chk["issues"], blocked=chk["blocked"], can_process=chk["can_process"],
                    builds=builds, latest_build=latest["id"] if latest else None, changed_since_build=changed,
                    state=_run_state(conn, run, builds, chk),
                    closed=run["state"] == "ARCHIVED", **_deletable(conn, actor, run_id))


def _deletable(conn, actor: dict, run_id: int) -> dict:
    if not access.can(actor, "prepare"):
        return dict(can_delete=False, delete_note=None)
    why = _delete_block(conn, run_id)
    return dict(can_delete=why is None, delete_note=why)


def _run_state(conn, run: dict, builds: list[dict], chk: dict) -> dict:
    if not builds:
        if chk["needs_answer"] or chk["issues"]:
            return dict(code="FILES_NEED_ATTENTION", text="Files need attention")
        return dict(code="ADDING_FILES", text="Adding files")
    b = builds[0]
    if b["state"] in ("QUEUED", "RUNNING"):
        return dict(code="PROCESSING", text="Processing")
    if b["state"] == "FAILED":
        return dict(code="FAILED", text="Processing failed")
    if b["state"] == "SUBMITTED":
        return dict(code="PUBLISHING", text="Publishing")
    if b["state"] in ("PUBLISHED", "SUPERSEDED"):
        return dict(code="PUBLISHED", text=f"Published as {b['generation']}")
    n = db.scalar(conn, "SELECT count(*) FROM rd_issue WHERE build_id = %s AND severity IN ('REVIEW','BLOCKING') AND "
                        "NOT later", (b["id"],))
    return dict(code="NEEDS_REVIEW", text="Needs your review") if n else dict(code="READY", text="Ready to publish")


def list_runs(actor: dict, client_id: int | None = None, limit: int = 20) -> list[dict]:
    with access.tx(actor) as conn:
        rows = db.all(conn, "SELECT r.*, c.code AS client_code, c.name AS client_name, u.name AS created_by_name "
                            "FROM rd_run r JOIN client c ON c.id = r.client_id JOIN app_user u ON u.id = r.created_by "
                            "WHERE (%(c)s::BIGINT IS NULL OR r.client_id = %(c)s) AND r.state = 'OPEN' "
                            "ORDER BY r.updated_at DESC LIMIT %(n)s", dict(c=client_id, n=limit))
        out = []
        for r in rows:
            builds = db.all(conn, "SELECT id, version, state, generation, created_at FROM rd_build WHERE run_id = %s "
                                  "ORDER BY version DESC LIMIT 1", (r["id"],))
            files = _run_files(conn, r["id"])
            chk = run_checks(conn, r, files)
            out.append(dict(id=r["id"], title=r["title"], client_id=r["client_id"], client_code=r["client_code"],
                            client_name=r["client_name"], created_by=r["created_by_name"], updated_at=r["updated_at"],
                            files=sum(1 for f in files if not f["removed"]), state=_run_state(conn, r, builds, chk),
                            latest_build=builds[0]["id"] if builds else None))
        return out


# ================================================================== builds
def _inputs_sha(conn, run: dict, used: list[dict]) -> str:
    h = hashlib.sha256()
    for f in sorted(used, key=lambda x: x["sha256"]):
        h.update(f"{f['sha256']}|{f['kind']}|{json.dumps(f['overrides'], sort_keys=True)}|"
                 f"{(f.get('choice') or {}).get('relation')}\n".encode())
    for d in db.all(conn, "SELECT id FROM rd_decision WHERE client_id = %s AND state = 'ACTIVE' ORDER BY id",
                    (run["client_id"],)):
        h.update(f"d{d['id']}\n".encode())
    h.update(rules.ruleset_sha(rules.RULESET_V1).encode())
    return h.hexdigest()


def start_build(actor: dict, run_id: int) -> dict:
    access.require(actor, "prepare")
    with access.tx(actor) as conn:
        run = _run(conn, run_id)
        conn.execute("SELECT pg_advisory_xact_lock(84726100, %s)", (run["client_id"],))
        active = db.one(conn, "SELECT id FROM rd_build WHERE client_id = %s AND state IN ('QUEUED','RUNNING')",
                        (run["client_id"],))
        if active:
            raise RDError("Files for this client are already being processed. Please wait until that finishes.")
        files = _run_files(conn, run_id)
        chk = run_checks(conn, run, files)
        if not chk["can_process"]:
            raise RDError(" ".join(chk["blocked"]))
        version = (db.scalar(conn, "SELECT max(version) FROM rd_build WHERE run_id = %s", (run_id,)) or 0) + 1
        bid = db.insert(conn, "rd_build", dict(
            client_id=run["client_id"], run_id=run_id, version=version, state="QUEUED", created_by=actor["id"],
            stages=db.J([dict(name=s, state="WAITING") for s in STAGES]),
            inputs_sha256=_inputs_sha(conn, run, chk["used"]), ruleset_version=rules.RULESET_V1["version"],
            ruleset_sha256=rules.ruleset_sha(rules.RULESET_V1), software_version=SOFTWARE,
            left_out=db.J(chk["left_out"])))
        db.insert(conn, "rd_job", dict(client_id=run["client_id"], build_id=bid, state="QUEUED"), returning=None)
        audit.record(conn, "build.queued", actor, "rd_build", bid, dict(run=run_id, version=version))
        _touch(conn, run_id)
    from . import jobs
    jobs.wake()
    return dict(build_id=bid, version=version)


def build_view(actor: dict, build_id: int) -> dict:
    with access.tx(actor) as conn:
        b = db.one(conn, "SELECT * FROM rd_build WHERE id = %s", (build_id,))
        if not b:
            raise RDError("This result does not exist, or you do not have access to it.")
        run = _run(conn, b["run_id"])
        c = _client(conn, b["client_id"])
        issues = db.all(conn, "SELECT * FROM rd_issue WHERE build_id = %s ORDER BY issue_no", (build_id,))
        reports = db.all(conn, "SELECT id, kind, final, filename, rows, available, reason, totals, bytes, sha256 "
                               "FROM rd_report WHERE build_id = %s ORDER BY final DESC, kind", (build_id,))
        people = {r["id"]: r["name"] for r in db.all(conn, "SELECT id, name FROM app_user")}
        findings_count = db.scalar(conn, "SELECT count(*) FROM rd_finding WHERE build_id = %s", (build_id,))
        decided = _decided_issue_keys(conn, b["client_id"])
        for i in issues:
            i["decision"] = decided.get(_issue_key(i))
        final = [r for r in reports if r["final"]]
        shown = final or [r for r in reports if not r["final"]]
        is_final = bool(final)
        for r in shown:
            r["state"] = _report_state(b, r, is_final)
        blocking = [i for i in issues if i["severity"] == "BLOCKING" and not i["decision"]]
        can_publish = b["state"] == "CHECKED" and not blocking and access.can(actor, "prepare")
        return dict(id=b["id"], run_id=b["run_id"], run_title=run["title"], version=b["version"], state=b["state"],
                    state_text=_build_state_text(b), client=_client_out(c), stages=b["stages"], error=b["error"],
                    error_detail=b["error_detail"] if access.can(actor, "diagnostics") else None,
                    summary=b["summary"], checks=b["checks"], notices=b["notices"], left_out=b["left_out"],
                    generation=b["generation"], created_by=people.get(b["created_by"]), created_at=b["created_at"],
                    finished_at=b["finished_at"], published_by=people.get(b["approved_by"]),
                    published_at=b["published_at"], returned_reason=b["returned_reason"],
                    issues=[_issue_out(i) for i in issues], findings=findings_count, reports=shown,
                    can_publish=can_publish, blocking=len(blocking), inputs=b["inputs"],
                    publishing=b["state"] == "SUBMITTED")


def _build_state_text(b: dict) -> str:
    return {"QUEUED": "Waiting to start", "RUNNING": "Processing", "CHECKED": "Validated - ready to publish",
            "FAILED": "Failed - no reports were produced", "SUBMITTED": "Publishing",
            "PUBLISHED": f"Published as {b['generation']}", "SUPERSEDED": f"Published as {b['generation']} "
            "(a newer version exists)"}[b["state"]]


def _report_state(b: dict, r: dict, final: bool) -> dict:
    if not r["available"]:
        return dict(code="HELD", text="Not available", reason=r["reason"])
    if final:
        return dict(code="PUBLISHED", text="Published")
    if b["state"] == "SUBMITTED":
        return dict(code="PUBLISHING", text="Publishing (draft)")
    return dict(code="VALIDATED", text="Validated (draft - not yet published)")


def _issue_key(i: dict) -> str:
    s = i.get("subject") or {}
    return f"{i['kind']}|{s.get('code_kind')}|{s.get('code')}|{s.get('file_id')}"


def _decided_issue_keys(conn, client_id: int) -> dict:
    out = {}
    for d in db.all(conn, "SELECT id, kind, subject, state, value FROM rd_decision WHERE client_id = %s AND kind IN "
                          "('LINK','NOT_SAME','ASSIGN_HELD') AND state = 'ACTIVE' ORDER BY id",
                    (client_id,)):
        s = d["subject"]
        kind = {"LINK": "PROPOSAL", "NOT_SAME": "PROPOSAL", "ASSIGN_HELD": "SHARED_RECORDING"}[d["kind"]]
        out[f"{kind}|{s.get('code_kind')}|{s.get('code')}|None"] = dict(
            id=d["id"], kind=d["kind"], state=d["state"],
            text=_decision_text(d))
    return out


def _decision_text(d: dict) -> str:
    v, s = d.get("value") or {}, d.get("subject") or {}
    t = (v.get("target") or {}).get("title")
    if d["kind"] == "LINK":
        return f"Confirm '{s.get('title')}' is the same song as '{t}'"
    if d["kind"] == "NOT_SAME":
        return f"Keep '{s.get('title')}' separate from '{t}'"
    if d["kind"] == "ASSIGN_HELD":
        return f"Place the money of recording {s.get('code')} on '{t}'"
    return f"Status '{rules.STATUS_LABEL.get(v.get('status'), v.get('status'))}' for '{s.get('title')}'"


def _decision_out(d: dict, people: dict) -> dict:
    return dict(id=d["id"], kind=d["kind"], text=_decision_text(d), state=d["state"], reason=d.get("reason"),
                made_by=people.get(d["made_by"]), made_at=d["made_at"], made_by_id=d["made_by"],
                checked_by=people.get(d.get("checked_by")), subject=d["subject"], value=d["value"],
                original=d.get("original"), evidence=d.get("evidence"))


def _issue_out(i: dict) -> dict:
    return dict(no=i["issue_no"], kind=i["kind"], severity=i["severity"], category=i["category"], title=i["title"],
                message=i["message"], subject=i["subject"], actions=i["actions"], later=i["later"],
                decision=i.get("decision"))


# ================================================================== review
ACTION_TEXT = {"CONFIRM_MATCH": "Confirm match", "KEEP_SEPARATE": "Keep separate", "LATER": "Review later",
               "ASSIGN_HELD": "Place the money on this song", "SET_STATUS": "Set status"}


def act(actor: dict, build_id: int, issue_no: int, action: str, payload: dict) -> dict:
    """A reviewer's answer to one issue. It applies at once (the newest answer for a song replaces an older one) and
    reaches the reports when the files are processed again."""
    access.require(actor, "prepare")
    reason = (payload.get("reason") or "").strip() or None
    with access.tx(actor) as conn:
        b = db.one(conn, "SELECT * FROM rd_build WHERE id = %s", (build_id,))
        i = db.one(conn, "SELECT * FROM rd_issue WHERE build_id = %s AND issue_no = %s", (build_id, issue_no))
        if not b or not i:
            raise RDError("This item no longer exists. Reload the page.")
        if action not in (i["actions"] or []):
            raise RDError("That action is not available for this item.")
        s = i["subject"]
        if action == "LATER":
            conn.execute("UPDATE rd_issue SET later = TRUE WHERE build_id = %s AND issue_no = %s", (build_id, issue_no))
            audit.record(conn, "review.later", actor, "rd_build", build_id, dict(issue=issue_no))
            return dict(ok=True, message="Saved for later.")
        if i["kind"] == "PROPOSAL" and action in ("CONFIRM_MATCH", "KEEP_SEPARATE"):
            target = dict(s["target"])
            kind = "LINK" if action == "CONFIRM_MATCH" else "NOT_SAME"
            if kind == "LINK" and not reason:
                reason = "titles and evidence reviewed"
            _replace_active(conn, b["client_id"], ("LINK", "NOT_SAME"), f"{s['code_kind']}:{s['code']}")
            did = db.insert(conn, "rd_decision", dict(
                client_id=b["client_id"], kind=kind,
                subject=dict(code_kind=s["code_kind"], code=s["code"], title=s["title"], isrcs=s["isrcs"],
                             work_nos=s["work_nos"]),
                subject_key=f"{s['code_kind']}:{s['code']}",
                value=dict(target=dict(code_kind=target["code_kind"], code=target["code"], title=target["title"],
                                       row_key=target.get("row_key"))),
                original=dict(paid_title=s["title"], paid_codes=s["isrcs"] + s["work_nos"], royalty=s.get("royalty"),
                              platform=s.get("platform"), base_title=target["title"],
                              base_codes=target.get("isrcs", []) + target.get("work_nos", [])),
                evidence=dict(score=target.get("score"), why=target.get("why")), reason=reason,
                source=f"review of build {build_id}", moves_money=(kind == "LINK"),
                state="ACTIVE", made_by=actor["id"], run_id=b["run_id"],
                build_id=build_id))
            audit.record(conn, f"decision.{kind.lower()}", actor, "rd_decision", did, dict(issue=issue_no))
            msg = ("Saved. Process the files again to move the money."
                   if kind == "LINK" else "Saved. These songs will be kept separate from now on.")
            return dict(ok=True, decision_id=did, message=msg)
        if i["kind"] == "SHARED_RECORDING" and action == "ASSIGN_HELD":
            cands = {c["row_key"]: c for c in s["candidates"]}
            pick = cands.get(payload.get("row_key"))
            if not pick:
                raise RDError("Choose the song the money belongs to.")
            anchor_kind = "ISRC" if pick["isrcs"] else "WORK"
            anchor = (pick["isrcs"] or pick["work_nos"])[0]
            _replace_active(conn, b["client_id"], ("ASSIGN_HELD",), f"ISRC:{s['code']}")
            did = db.insert(conn, "rd_decision", dict(
                client_id=b["client_id"], kind="ASSIGN_HELD",
                subject=dict(code_kind="ISRC", code=s["code"], title=s["title"]),
                subject_key=f"ISRC:{s['code']}",
                value=dict(target=dict(code_kind=anchor_kind, code=anchor if anchor != s["code"] else
                                       (pick["work_nos"] or [anchor])[0], title=pick["title"],
                                       row_key=pick["row_key"])),
                original=dict(candidates=s["candidates"]), evidence={}, reason=reason or "assigned by reviewer",
                source=f"review of build {build_id}", moves_money=True, state="ACTIVE",
                made_by=actor["id"], run_id=b["run_id"], build_id=build_id))
            audit.record(conn, "decision.assign_held", actor, "rd_decision", did, dict(issue=issue_no))
            return dict(ok=True, decision_id=did, message="Saved. Process the files again to move the money.")
        if i["kind"] in ("ALREADY_IN_BASE",) and action in ("ADD_AS_SEPARATE", "LEAVE_OUT"):
            choose(actor, b["run_id"], s["file_id"], "SUPPLEMENTARY" if action == "ADD_AS_SEPARATE" else "LEAVE_OUT")
            return dict(ok=True, message="Saved. Process the files again to apply it.")
    raise RDError("That action is not available for this item.")


def _replace_active(conn, client_id: int, kinds: tuple, subject_key: str) -> None:
    conn.execute("UPDATE rd_decision SET state = 'SUPERSEDED' WHERE client_id = %s AND kind = ANY(%s) AND "
                 "subject_key = %s AND state = 'ACTIVE'", (client_id, list(kinds), subject_key))


def set_status(actor: dict, build_id: int, finding_nos: list[int], status: str, note: str | None = None) -> dict:
    """Review status of paid-but-not-in-list songs. It moves no money, so it applies at once (and is logged)."""
    access.require(actor, "prepare")
    if status not in rules.STATUS_LABEL or status == rules.NEW:
        raise RDError("Choose one of the listed statuses.")
    with access.tx(actor) as conn:
        b = db.one(conn, "SELECT * FROM rd_build WHERE id = %s", (build_id,))
        if not b:
            raise RDError("This result does not exist.")
        n = 0
        for no in finding_nos:
            f = db.one(conn, "SELECT * FROM rd_finding WHERE build_id = %s AND finding_no = %s", (build_id, no))
            if not f:
                continue
            _song_status(conn, actor, b["client_id"], dict(isrcs=list(f["isrcs"]), work_nos=list(f["work_nos"]),
                                                           title=f["title"]),
                         status, rules.STATUS_LABEL[status], note or "", source=f"review of build {build_id}",
                         run_id=b["run_id"], build_id=build_id)
            conn.execute("UPDATE rd_finding SET status = %s, status_original = %s WHERE build_id = %s AND "
                         "finding_no = %s", (status, rules.STATUS_LABEL[status], build_id, no))
            n += 1
        audit.record(conn, "review.status", actor, "rd_build", build_id, dict(findings=finding_nos, status=status))
    return dict(ok=True, updated=n, message=f"Saved for {n} song{'s' if n != 1 else ''}. Process again to update the "
                                            "reports.")


def findings(actor: dict, build_id: int, status: str | None = None, q: str | None = None, limit: int = 200,
             offset: int = 0) -> dict:
    with access.tx(actor) as conn:
        b = db.one(conn, "SELECT client_id FROM rd_build WHERE id = %s", (build_id,))
        if not b:
            raise RDError("This result does not exist.")
        where, args = ["build_id = %s"], [build_id]
        if status:
            where.append("status = %s")
            args.append(status)
        if q:
            where.append("(title ILIKE %s OR %s = ANY(isrcs) OR %s = ANY(work_nos))")
            args += [f"%{q}%", q.upper(), q]
        rows = db.all(conn, f"SELECT * FROM rd_finding WHERE {' AND '.join(where)} ORDER BY finding_no "
                            f"LIMIT %s OFFSET %s", (*args, min(limit, 500), offset))
        counts = {r["status"]: r["n"] for r in db.all(conn, "SELECT status, count(*) AS n FROM rd_finding WHERE "
                                                            "build_id = %s GROUP BY status", (build_id,))}
        return dict(items=[dict(no=r["finding_no"], title=r["title"], isrcs=r["isrcs"], work_nos=r["work_nos"],
                                royalty=r["royalty"], platform=r["platform"], status=r["status"],
                                status_text=rules.STATUS_LABEL.get(r["status"], r["status"]),
                                status_original=r["status_original"], sources=r["sources"], proposal=r["proposal"])
                           for r in rows], counts=counts,
                    statuses=[dict(value=k, label=v) for k, v in rules.STATUS_LABEL.items() if k != rules.NEW])


def publish(actor: dict, build_id: int) -> dict:
    """The preparer publishes a validated result: the reports are rebuilt from the same inputs, checked again and
    saved for good under the next report number (e.g. SVF-RD-SKV9)."""
    access.require(actor, "prepare")
    with access.tx(actor) as conn:
        b = db.one(conn, "SELECT * FROM rd_build WHERE id = %s FOR UPDATE", (build_id,))
        if not b or b["state"] != "CHECKED":
            raise RDError("Only a validated result can be published.")
        blocking = db.scalar(conn, "SELECT count(*) FROM rd_issue WHERE build_id = %s AND severity = 'BLOCKING'",
                             (build_id,))
        if blocking:
            raise RDError("Resolve the items marked 'must be resolved' first, then process again.")
        run = _run(conn, b["run_id"])
        if _inputs_sha(conn, run, run_checks(conn, run, _run_files(conn, run["id"]))["used"]) != b["inputs_sha256"]:
            raise RDError("Files, answers or review decisions changed after this version was processed. Process the "
                          "files again, then publish the new version.")
        conn.execute("UPDATE rd_build SET state = 'SUBMITTED', submitted_by = %(a)s, submitted_at = now(), "
                     "approved_by = %(a)s, approved_at = now(), returned_reason = NULL WHERE id = %(b)s",
                     dict(a=actor["id"], b=build_id))
        db.insert(conn, "rd_job", dict(client_id=b["client_id"], build_id=build_id, state="QUEUED"), returning=None)
        audit.record(conn, "build.publish_requested", actor, "rd_build", build_id)
    from . import jobs
    jobs.wake()
    return dict(ok=True, message="Publishing. The final reports will be ready in a moment.")

def _report_bytes(conn, actor: dict, report_id: int) -> tuple[dict, bytes]:
    """The stored workbook, after the same checks for a download and a preview."""
    r = db.one(conn, "SELECT * FROM rd_report WHERE id = %s", (report_id,))
    if not r:
        raise Forbidden("This report does not exist, or you do not have access to it.")
    b = db.one(conn, "SELECT state FROM rd_build WHERE id = %s", (r["build_id"],))
    if not r["available"]:
        raise RDError(r["reason"] or "This report is not available.")
    if not r["final"]:
        access.require(actor, "draft_download")
        if b["state"] in ("QUEUED", "RUNNING", "FAILED"):
            raise RDError("This report is not ready.")
    data = storage.get("reports", r["object_key"])
    if hashlib.sha256(data).hexdigest() != r["sha256"]:
        raise RDError("The stored report does not match its fingerprint. Please contact your administrator.")
    return r, data


def report_file(actor: dict, report_id: int) -> tuple[str, bytes]:
    with access.tx(actor) as conn:
        r, data = _report_bytes(conn, actor, report_id)
        name = with_person(r["filename"], actor)
        audit.record(conn, "report.downloaded", actor, "rd_report", report_id, dict(file=r["filename"], saved_as=name))
    return name, data


PREVIEW_ROWS, PREVIEW_COLS = 300, 60


def report_preview(actor: dict, report_id: int, sheet: int = 0) -> dict:
    """The first rows of one sheet as text, to look at a report before downloading it."""
    import io

    import openpyxl
    with access.tx(actor) as conn:
        r, data = _report_bytes(conn, actor, report_id)
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheets = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
        ws = sheets[max(0, min(sheet, len(sheets) - 1))]
        rows = []
        for row in ws.iter_rows(max_row=PREVIEW_ROWS, max_col=PREVIEW_COLS):
            rows.append([_preview_cell(c) for c in row])
        width = max((max((i + 1 for i, v in enumerate(row) if v), default=0) for row in rows), default=0)
        rows = [row[:width] + [""] * (width - len(row)) for row in rows]
        while rows and not any(rows[-1]):
            rows.pop()
        total = ws.max_row or len(rows)
        out = dict(filename=r["filename"], sheets=[x.title for x in sheets], sheet=sheets.index(ws), rows=rows,
                   total_rows=total, total_cols=ws.max_column or width, shown_rows=len(rows))
    finally:
        wb.close()
    return out


def _preview_cell(c) -> str:
    v = getattr(c, "value", None)
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if isinstance(v, (int, float, Decimal)):
        fmt = getattr(c, "number_format", "") or ""
        if "0.0" in fmt or (isinstance(v, float) and not v.is_integer()):
            return _inr(Decimal(str(v)))
        return str(int(v))
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d %b %Y")
    return str(v)


def _inr(v: Decimal) -> str:
    """12345678.5 -> '1,23,45,678.50' (Indian grouping, two decimals)."""
    whole, frac = f"{abs(v):.2f}".split(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        groups = []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        whole = ",".join([head] + groups if head else groups) + "," + tail
    return ("-" if v < 0 else "") + whole + "." + frac


def with_person(filename: str, actor: dict) -> str:
    """'SVF-RD-SKV9.xlsx' downloaded by Swatantra Kumar -> 'SVF-RD-SKV9_Swatantra-Kumar.xlsx'."""
    import re
    who = re.sub(r"[^A-Za-z0-9]+", "-", actor.get("name") or actor.get("email", "").split("@")[0]).strip("-")[:40]
    stem, dot, ext = filename.rpartition(".")
    return f"{stem}_{who}.{ext}" if who and dot else filename


def report_bundle(actor: dict, build_id: int) -> tuple[str, bytes]:
    import io
    import zipfile
    with access.tx(actor) as conn:
        reps = db.all(conn, "SELECT id, final FROM rd_report WHERE build_id = %s AND available ORDER BY kind",
                      (build_id,))
        g = db.scalar(conn, "SELECT coalesce(generation, 'reports') FROM rd_build WHERE id = %s", (build_id,))
    final = [r for r in reps if r["final"]] or reps
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for r in final:
            name, data = report_file(actor, r["id"])
            z.writestr(name, data)
    return with_person(f"{g}-reports.zip", actor), buf.getvalue()


# ================================================================== home and history
def home(actor: dict) -> dict:
    runs = list_runs(actor, limit=8)
    with access.tx(actor) as conn:
        pubs = db.all(conn, "SELECT b.id, b.generation, b.published_at, c.code AS client_code, b.run_id FROM rd_build b "
                            "JOIN client c ON c.id = b.client_id WHERE b.state IN ('PUBLISHED','SUPERSEDED') "
                            "ORDER BY b.published_at DESC LIMIT 6")
        clients = access.client_ids(conn, actor)
    attention = [r for r in runs if r["state"]["code"] in ("FILES_NEED_ATTENTION", "NEEDS_REVIEW", "FAILED", "READY")] \
        if access.can(actor, "prepare") else []                    # a viewer is never asked to act
    nxt = None
    if attention and access.can(actor, "prepare"):
        r = attention[0]
        nxt = dict(kind="CONTINUE", run_id=r["id"], text=f"Continue '{r['title']}' - {r['state']['text'].lower()}.")
    elif access.can(actor, "prepare"):
        nxt = dict(kind="NEW_RUN", text="Start a new processing run.")
    else:
        nxt = dict(kind="REPORTS", text="Open Past Reports to download published reports.")
    return dict(next=nxt, runs=runs, attention=attention, published=pubs,
                can_prepare=access.can(actor, "prepare"), clients=len(clients))


def past_reports(actor: dict, client_id: int | None = None) -> list[dict]:
    with access.tx(actor) as conn:
        bs = db.all(conn, "SELECT b.*, c.code AS client_code, c.name AS client_name, r.title AS run_title FROM rd_build b "
                          "JOIN client c ON c.id = b.client_id JOIN rd_run r ON r.id = b.run_id "
                          "WHERE b.state IN ('PUBLISHED','SUPERSEDED','CHECKED','SUBMITTED') AND r.state <> 'DELETED' "
                          "AND (%(c)s::BIGINT IS NULL OR b.client_id = %(c)s) ORDER BY coalesce(b.published_at, "
                          "b.finished_at) DESC LIMIT 100", dict(c=client_id))
        people = {r["id"]: r["name"] for r in db.all(conn, "SELECT id, name FROM app_user")}
        deletable = {rid: _deletable(conn, actor, rid)["can_delete"] for rid in {b["run_id"] for b in bs}}
        out = []
        for b in bs:
            reps = db.all(conn, "SELECT id, kind, final, filename, available, reason, rows FROM rd_report WHERE "
                                "build_id = %s ORDER BY kind", (b["id"],))
            final = [r for r in reps if r["final"]]
            if b["state"] in ("PUBLISHED", "SUPERSEDED"):
                reps = final
            elif not access.can(actor, "draft_download"):
                continue
            else:
                reps = [r for r in reps if not r["final"]]
            out.append(dict(build_id=b["id"], run_id=b["run_id"], run_title=b["run_title"],
                            client_code=b["client_code"], client_name=b["client_name"], version=b["version"],
                            state=b["state"], state_text=_build_state_text(b), generation=b["generation"],
                            published_at=b["published_at"], finished_at=b["finished_at"],
                            published_by=people.get(b["approved_by"]), reports=reps,
                            royalty=(b["summary"] or {}).get("royalty"), can_delete=deletable[b["run_id"]]))
        return out


def diagnostics(actor: dict, limit: int = 200) -> dict:
    access.require(actor, "diagnostics")
    with access.tx(actor) as conn:
        full = access.can(actor, "audit")              # the activity log is the Super Admin's
        events = audit.recent(conn, limit=limit) if full else []
        chain = audit.verify_chain(conn) if full else None
        failed = db.all(conn, "SELECT b.id, b.run_id, b.version, b.error, b.error_detail, b.finished_at, c.code FROM "
                              "rd_build b JOIN client c ON c.id = b.client_id WHERE b.state = 'FAILED' ORDER BY b.id "
                              "DESC LIMIT 20")
        jobs = db.all(conn, "SELECT * FROM rd_job ORDER BY id DESC LIMIT 20")
    return dict(events=events, chain=chain, failed_builds=failed, jobs=jobs, software=SOFTWARE,
                ruleset=rules.RULESET_V1, storage=str(config.settings.storage_root))


def file_details(actor: dict, file_id: int) -> dict:
    with access.tx(actor) as conn:
        f = db.one(conn, "SELECT * FROM rd_file WHERE id = %s", (file_id,))
        if not f:
            raise RDError("This file does not exist, or you do not have access to it.")
        out = dict(id=f["id"], filename=f["filename"], sha256=f["sha256"], bytes=f["bytes"], kind=f["kind"],
                   layout=f["layout"], why=f["why"], meta=f["meta"], meta_source=f["meta_source"],
                   problems=f["problems"], uploaded_at=f["uploaded_at"])
        if access.can(actor, "diagnostics") or access.can(actor, "prepare"):
            out["facts"] = f["facts"]
        return out
