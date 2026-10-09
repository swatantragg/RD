"""H. Output + resolution (pages 9-11): Excel evidence pack, client summary, editable email; approval, sending,
reply tracking. Export is not dispatch: a pack is SENT only when an authorised person records it."""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import zipfile
from collections import defaultdict
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import audit, db, storage
from .compare import ENQUIRY_OUTCOMES, OUTCOME_LABEL, P_STATUS_LABEL
from .normalize import period_label, q2, s
from .security import not_same_person

ROUTE_GREETING = {"IPRS": "Dear IPRS Team,", "PRS": "Dear PRS for Music Team,", "DIRECT": "Dear Publisher Accounts Team,"}
ROUTE_NAME = {"IPRS": "IPRS", "PRS": "PRS for Music", "DIRECT": "the paying publisher"}


class PackError(Exception):
    pass


# ------------------------------------------------------------------ email (page 11 template)
def email(client: dict, run: dict, route: str, groups: list[dict], method: str | None, rule: dict | None,
          approver_name: str) -> tuple[str, str]:
    period = period_label(run["period_from"], run["period_to"])
    member = f"{client['beneficiary_name'] or client['name']} (membership {client['member_id'] or 'not stated'})"
    subject = (f"Request for royalty reconciliation and clarification | {client['beneficiary_name'] or client['name']} "
               f"| {run['platform']} | {period}")
    kinds = defaultdict(int)
    for g in groups:
        kinds[g["outcome"]] += 1
    parts = []
    if kinds.get("NO_MATCHING_ENTRY"):
        parts.append(f"{kinds['NO_MATCHING_ENTRY']} with no matching entry")
    if kinds.get("BELOW_COMPARISON"):
        parts.append(f"{kinds['BELOW_COMPARISON']} below the comparison amount")
    if kinds.get("COVERAGE_ONLY"):
        parts.append(f"{kinds['COVERAGE_ONLY']} usage-only / coverage")
    if kinds.get("STATUS_ENQUIRY"):
        parts.append(f"{kinds['STATUS_ENQUIRY']} registration or distribution-status queries")
    if method == "MRM_BENCHMARK" and rule:
        basis = (f"For the eligible rows marked MRM benchmark, we used {format((Decimal(rule['rate']) * 100).normalize(), 'f')}% of "
                 f"{rule['denominator']} ({(rule['revenue_basis'] or '').lower()} basis) as an internal screening "
                 f"reference, with the share treatment disclosed "
                 f"({'beneficiary-specific' if rule['share_treatment'] == 'BENEFICIARY_SPECIFIC' else 'share of 100% of the work'}). "
                 f"We request confirmation of its relevance to these cases; we have not treated it as an established "
                 f"{ROUTE_NAME[route]} tariff or confirmed amount payable.")
    elif method == "EVIDENCED" and rule:
        basis = (f"For the eligible rows we applied the evidenced method {rule['name']} (reference "
                 f"{rule['evidence_ref']}), effective {rule['effective_from']:%d %b %Y}, with the share treatment "
                 f"disclosed in the annexure.")
    else:
        basis = ("No monetary benchmark applies to this scope; the annexure reports actual income and usage and asks "
                 "coverage and status questions only.")
    body = f"""{ROUTE_GREETING[route]}

On behalf of {member}, MRM has reviewed the attached recording revenue/usage reports against the publishing royalty statements available to us for {run['platform']}, {run['territory']}, {period}, with a statement cutoff of {run['cutoff']:%d %B %Y}.

We reconciled the supplied statements to their source totals, reviewed duplicates and linked the recordings to the relevant musical works using the identifiers and supporting metadata listed in the annexure. We have kept recording and publishing income separate and identified the beneficiary/share and comparison basis for each queried line.

{basis}

The annexure identifies {len(groups)} works requiring clarification, including {', '.join(parts) or 'no flagged items'}. Where only streams are available and no applicable monetary rate is evidenced, we have reported usage without estimating an amount due.

Please confirm for each flagged work:
1. Whether the recording-to-work mapping, registration and beneficiary share are correct.
2. Whether the relevant platform usage and revenue have been received and allocated.
3. The applicable calculation basis, deductions, thresholds and distribution period.
4. Any reason for a missing or different allocation, and the documents or corrective action required.
5. Where an adjustment is warranted, the expected distribution/reference.

Please return your response against the case IDs in the attached schedule so we can update the publisher and close each item accurately.

Regards,
{approver_name}
Music Rights Management (MRM) India"""
    return subject, body


# ------------------------------------------------------------------ workbook
HEAD = PatternFill("solid", fgColor="1F3A5F")
BAND = PatternFill("solid", fgColor="EEF2F7")
HOLDF = PatternFill("solid", fgColor="FFF4E5")
THIN = Side(style="thin", color="C9D3E0")
MONEY = "#,##0.00;[Red]-#,##0.00"
INT = "#,##0"


class Sheet:
    def __init__(self, wb: Workbook, title: str, heading: str, note: str | None = None):
        self.ws = wb.create_sheet(title)
        self.row = 1
        self.ws.cell(1, 1, heading).font = Font(bold=True, size=13, color="1F3A5F")
        self.row = 2
        if note:
            self.ws.cell(2, 1, note).font = Font(italic=True, size=9, color="555555")
            self.row = 3
        self.row += 1
        self.widths: dict[int, int] = {}

    def kv(self, label: str, value, fmt: str | None = None):
        a = self.ws.cell(self.row, 1, label)
        a.font = Font(bold=True)
        b = self._put(self.row, 2, value)
        if fmt:
            b.number_format = fmt
        self._w(1, label)
        self._w(2, value)
        self.row += 1

    def blank(self, n=1):
        self.row += n

    def title(self, text: str):
        self.ws.cell(self.row, 1, text).font = Font(bold=True, size=11, color="1F3A5F")
        self.row += 1

    def table(self, headers: list[str], rows: list[list], formats: dict[int, str] | None = None, hold_col: int | None = None):
        formats = formats or {}
        for j, h in enumerate(headers, 1):
            c = self.ws.cell(self.row, j, h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD
            c.alignment = Alignment(wrap_text=True, vertical="top")
            self._w(j, h, 10)
        start = self.row
        self.row += 1
        for i, r in enumerate(rows):
            for j, v in enumerate(r, 1):
                c = self._put(self.row, j, v)
                c.border = Border(bottom=THIN)
                if j in formats and isinstance(v, (int, float, Decimal)):
                    c.number_format = formats[j]
                if i % 2:
                    c.fill = BAND
                self._w(j, v)
            if hold_col is not None and r[hold_col - 1] in ("Hold", "Status enquiry"):
                for j in range(1, len(r) + 1):
                    self.ws.cell(self.row, j).fill = HOLDF
            self.row += 1
        if rows:
            self.ws.auto_filter.ref = f"A{start}:{get_column_letter(len(headers))}{self.row - 1}"
        self.ws.freeze_panes = self.ws.cell(start + 1, 1)
        self.row += 1

    def _put(self, r: int, c: int, v):
        if isinstance(v, Decimal):
            v = float(q2(v)) if abs(v) < Decimal("1e15") else float(v)
        if isinstance(v, (list, tuple)):
            v = ", ".join(str(x) for x in v)
        if isinstance(v, dict):
            v = str(v)
        cell = self.ws.cell(r, c, v)
        if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
            cell.data_type = "s"                       # text that looks like a formula stays text
        if isinstance(v, (dt.date, dt.datetime)):
            cell.number_format = "dd-mmm-yyyy"
        return cell

    def _w(self, j: int, v, minimum: int = 8):
        n = min(60, max(minimum, len(str(v if v is not None else "")) + 2))
        if n > self.widths.get(j, 0):
            self.widths[j] = n
            self.ws.column_dimensions[get_column_letter(j)].width = n


def _d(v):
    return None if v is None else Decimal(v)


def workbook(conn, pack: dict, run: dict, client: dict, route: str, enquiry: list[dict]) -> bytes:
    groups = db.all(conn, "SELECT g.*, r.rule_key, r.version_no AS rule_version FROM comparison_group g LEFT JOIN "
                          "rule_version r ON r.id = g.rule_version_id WHERE g.run_id = %s ORDER BY g.id", (run["id"],))
    exc = db.all(conn, "SELECT * FROM run_exception WHERE run_id = %s ORDER BY id", (run["id"],))
    iv = db.one(conn, "SELECT * FROM identity_version WHERE id = %s", (run["identity_version_id"],))
    files = db.all(conn, "SELECT f.*, u.email AS uploader FROM source_file f LEFT JOIN app_user u ON u.id = "
                         "f.uploaded_by WHERE f.client_id = %s ORDER BY f.id", (client["id"],))
    rule_hist = db.all(conn, "SELECT r.*, a.email AS approver, c.email AS author FROM rule_version r LEFT JOIN app_user a "
                             "ON a.id = r.approved_by LEFT JOIN app_user c ON c.id = r.created_by WHERE r.client_id IS "
                             "NULL OR r.client_id = %s ORDER BY r.rule_key, r.version_no", (client["id"],))
    fname = {f["id"]: f["filename"] for f in files}
    sm = run["summary"]
    wb = Workbook()
    wb.remove(wb.active)
    period = period_label(run["period_from"], run["period_to"])

    # 1 client summary
    sh = Sheet(wb, "1 Client summary", "MRM Royalty Comparison - client summary",
               "Recording and publishing income are kept separate. A 6% MRM figure is an indicative benchmark and its "
               "variance a difference for clarification - not an amount due.")
    for k, v in [("Publisher / label", client["name"]), ("Publishing beneficiary", client["beneficiary_name"]),
                 ("Member ID", client["member_id"]),
                 ("Mandate", f"{'verified' if client['mandate_verified'] else 'NOT verified'} {client['mandate_ref'] or ''}"
                             f" ({client['mandate_from'] or '?'} to {client['mandate_to'] or 'open'})"),
                 ("Collection route", route), ("Platform", run["platform"]), ("Territory", run["territory"]),
                 ("Use type", run["use_type"].replace("_", " ").lower()), ("Publishing-right scope", run["right_scope"]),
                 ("Usage period", period), ("Statement cutoff", run["cutoff"]),
                 ("Comparison run", f"v{run['version_no']} of {run['created_at']:%d %b %Y %H:%M}"),
                 ("Identity version", f"v{iv['version_no']} approved {iv['approved_at']:%d %b %Y}"),
                 ("Rule", f"{sm['rule']['key']} v{sm['rule']['version']} ({sm['method'].replace('_', ' ').lower()})"
                          if sm.get("rule") else sm["method"].replace("_", " ").lower()),
                 ("Report pack", f"v{pack['version_no']} built {pack['built_at']:%d %b %Y %H:%M}"),
                 ("Works in input", sm.get("works_input")), ("Eligible (comparable) works", sm.get("works_eligible")),
                 ("Comparison groups", sm["groups"]), ("Publishing statement coverage", sm["coverage"].lower())]:
        sh.kv(k, v)
    sh.blank()
    sh.title("Separate totals in scope")
    sh.table(["Measure", "Amount"], [
        ["Recording revenue R in scope", _d(sm["r_in_scope"])], ["  of which mapped to works", _d(sm["r_grouped"])],
        ["  of which unmapped (exceptions)", _d(sm["r_unmapped"])], ["Publishing P in scope", _d(sm["p_in_scope"])],
        ["  of which matched to works", _d(sm["p_grouped"])], ["  of which not in register", _d(sm["p_unregistered"])],
        ["Streams N in scope (usage rows)", sm["n_in_scope"]]], {2: MONEY})
    sh.title("Summary control - per currency and right; only comparable groups are added")
    sh.table(["Currency", "Right", "Comparable groups", "R", "B (indicative benchmark)", "P", "Positive gaps",
              "Above benchmark (separate)", "Below", "Within/above", "No matching entry (provisional)",
              "Provisional B", "Coverage-only"],
             [[t["currency"], t["right_scope"], t["comparable_groups"], _d(t["r"]), _d(t["b"]), _d(t["p"]),
               _d(t["positive_gaps"]), _d(t["above_benchmark"]), t["below_groups"], t["within_groups"],
               t["provisional_groups"], _d(t["provisional_b"]), t["coverage_only"]] for t in sm["by_currency"]],
             {4: MONEY, 5: MONEY, 6: MONEY, 7: MONEY, 8: MONEY, 12: MONEY})
    sh.title("Outcomes")
    sh.table(["Outcome", "Groups"], [[OUTCOME_LABEL[k], v] for k, v in sorted(sm["outcomes"].items())])
    if sm.get("holds"):
        sh.title("Holds and status enquiries by reason (no monetary comparison)")
        sh.table(["Reason", "Groups", "Recording revenue held"],
                 [[k.replace("_", " "), v["groups"], _d(v["r"])] for k, v in sm["holds"].items()], {3: MONEY})
    sh.title("Coverage limits")
    for line in [f"Recording revenue months: {', '.join(m[:7] for m in sm['r_months']) or 'none'}",
                 f"Usage (stream) months: {', '.join(m[:7] for m in sm['n_months']) or 'none'}",
                 f"Statements: {'; '.join(x['filename'] for x in sm['statements']) or 'none received'}",
                 f"Exceptions listed on sheet 5: {sm['exceptions']}"]:
        sh.kv("", line)

    # 2 mapped comparison
    sh = Sheet(wb, "2 Mapped comparison", "One row per comparable group")
    sh.table(["Case ID", "Work title", "Film / album", "Composer", "Lyricist", "Internal work ID", "IPRS work no",
              "Linked ISRCs", "Masters", "Mapping status / version", "Outcome"],
             [[g["case_code"], g["title"], g["film"], g["composer"], g["lyricist"], g["work_code"], g["work_code"],
               g["isrcs"], len(g["isrcs"]), g["mapping_status"], OUTCOME_LABEL[g["outcome"]]] for g in groups],
             hold_col=11)

    # 3 financial + usage detail
    sh = Sheet(wb, "3 Financial + usage", "Recording amount R, streams N and publishing P per group",
               "R is summed across the group's masters once; P is the composition-level amount, counted once.")
    sh.table(["Case ID", "Work", "Platform", "Usage period", "R", "Currency", "Basis", "Share", "N", "N definition",
              "P", "Statement / payment status", "Distribution ref", "Rows (R / N / P)"],
             [[g["case_code"], g["title"], run["platform"], period, _d(g["r_amount"]), g["currency"],
               _basis(files, g), _share(files, g), g["n_streams"], g["n_definition"], _d(g["p_amount"]),
               P_STATUS_LABEL[g["p_status"]], _dist(files, g),
               f"{len(g['evidence']['r_rows'])} / {g['evidence'].get('n_rows_total', 0)} / {len(g['evidence']['p_rows'])}"]
              for g in groups], {5: MONEY, 9: INT, 11: MONEY})
    sh.title("Per master (aggregation control: R per ISRC adds up to the group R)")
    sh.table(["Case ID", "ISRC", "Recording title", "R", "N", "Allocation"],
             [[g["case_code"], m["isrc"], m["title"], _d(m["r"]), int(Decimal(m["n"])) if m["n"] is not None else None,
               _d(m["allocation"])] for g in groups for m in g["recordings"]], {4: MONEY, 5: INT})

    # 4 calculation + question
    sh = Sheet(wb, "4 Calculation + question", "Rule, comparison amount B, variance V = B - P and the exact question")
    sh.table(["Case ID", "Work", "Rule / version", "Method", "Share treatment", "Route", "Formula", "B", "Stream "
              "cross-check (not added)", "P", "Signed variance V", "Positive gap", "Tolerance", "Outcome",
              "Hold / query reason", "Confidence / limitations", "Clarification requested"],
             [[g["case_code"], g["title"], f"{g['rule_key']} v{g['rule_version']}" if g["rule_key"] else "",
               _method(g["method"]), _share_treatment(rule_hist, g["rule_version_id"]), g["route"].lower(),
               g["evidence"].get("formula"), _d(g["b_amount"]), _d(g["b_stream"]), _d(g["p_amount"]),
               _d(g["variance"]), _d(g["positive_gap"]), _d(g["evidence"].get("tolerance")),
               OUTCOME_LABEL[g["outcome"]] + (" (provisional)" if g["provisional"] else ""),
               (g["hold_reason"] or "").replace("_", " ").lower(), g["limitations"], g["question"]] for g in groups],
             {8: MONEY, 9: MONEY, 10: MONEY, 11: MONEY, 12: MONEY, 13: MONEY}, hold_col=14)

    # 5 exceptions + review
    sh = Sheet(wb, "5 Exceptions + review", "Unknown and unmatched items - no silent exclusions")
    held = [g for g in groups if g["outcome"] in ("HOLD", "STATUS_ENQUIRY")]
    sh.table(["Kind", "Subject", "Amount", "Streams", "Detail", "Owner", "Next action", "Deadline", "Resolution evidence"],
             [[e["kind"].replace("_", " ").lower(), e["subject"], _d(e["amount"]), e["streams"], _detail(e["detail"]),
               e["owner"], e["next_action"], None, None] for e in exc]
             + [[f"group {g['outcome'].lower().replace('_', ' ')}: {(g['hold_reason'] or '').replace('_', ' ').lower()}",
                 f"{g['case_code']} {g['title']}", _d(g["r_amount"]), g["n_streams"], g["evidence"].get("hold_detail"),
                 _owner(g["hold_reason"]), "rerun the affected gate after correction", None, None] for g in held],
             {3: MONEY, 4: INT})

    # 6 source + reconciliation
    sh = Sheet(wb, "6 Source + reconciliation", "Original files, versions, totals, validation results and history")
    used = set(run["file_ids"])
    sh.table(["File ID", "File", "SHA-256", "Type", "Status", "Used in run", "Period", "Platform", "Currency",
              "Printed total", "Lines total", "Alternative control", "Control reference", "Rows", "Invalid rows",
              "Supersedes", "Uploaded by", "Uploaded at", "Checks"],
             [[f["id"], f["filename"], f["sha256"], f["income_type"], f["status"], "yes" if f["id"] in used else "no",
               f"{f['meta'].get('period_from', '')} to {f['meta'].get('period_to', '')}", f["meta"].get("platform"),
               f["meta"].get("currency"), _d(f["printed_total"]), _d(f["lines_total"]), _d(f["control_total"]),
               f["control_ref"], f["rows_total"], f["rows_invalid"], f["supersedes_id"], f["uploader"],
               f["uploaded_at"].replace(tzinfo=None) if f["uploaded_at"] else None,
               "; ".join(f"{c.get('gate', '')} {c['check']}: {c['result']}" for c in f["checks"] or [])]
              for f in files], {10: MONEY, 11: MONEY, 12: MONEY})
    sh.title("Run controls")
    sh.table(["Control", "Result", "Detail"], [[c["check"], c["result"], c["detail"]] for c in run["checks"]])
    sh.title("Rule history")
    sh.table(["Rule", "Version", "Status", "Method", "k", "q", "Denominator", "Basis", "Share treatment", "Effective",
              "Tolerance", "Drafted by", "Approved by", "Approved at", "Reason"],
             [[r["rule_key"], r["version_no"], r["status"], r["method"], _d(r["rate"]), _d(r["stream_rate"]),
               r["denominator"], r["revenue_basis"], r["share_treatment"],
               f"{r['effective_from'] or ''} to {r['effective_to'] or 'open'}", _d(r["tolerance"]), r["author"],
               r["approver"], r["approved_at"].replace(tzinfo=None) if r["approved_at"] else None, r["reason"]]
              for r in rule_hist])
    sh.title("Source rows behind every figure (file, sheet, row)")
    trace = []
    for g in groups:
        for x in g["evidence"]["r_rows"]:
            trace.append([g["case_code"], "R", fname.get(x["file_id"]), x["sheet"], x["row"], x["isrc"],
                          (x["month"] or "")[:7], _d(x["amount"]), _d(x["allocation"])])
        for x in g["evidence"]["p_rows"]:
            trace.append([g["case_code"], "P", fname.get(x["file_id"]), x["sheet"], x["row"], x["work_code"],
                          x["source"], _d(x["amount"]), None])
        for x in g["evidence"]["n_rows"]:
            trace.append([g["case_code"], "N", fname.get(x["file_id"]), x["sheet"], x["row"], x["isrc"],
                          (x["month"] or "")[:7], x["streams"], _d(x["allocation"])])
    sh.table(["Case ID", "Figure", "File", "Sheet", "Row", "Identifier", "Month / source", "Amount / streams",
              "Allocation"], trace, {8: MONEY})
    sh.kv("Export version", f"pack v{pack['version_no']} of run v{run['version_no']}")

    # 7 enquiry schedule
    sh = Sheet(wb, f"7 {route} enquiry schedule", f"{ROUTE_NAME[route]} enquiry schedule - {client['beneficiary_name'] or client['name']}",
               "Only relevant, reviewed cases for the named client/beneficiary. Please complete the reply columns.")
    sh.table(["Case ID", "Work title", "IPRS work no", "ISRCs", "Platform", "Territory", "Usage period", "Query",
              "R", "N", "B (indicative benchmark)", "P", "Difference for clarification", "Statement status",
              "Evidence reference", "Clarification requested", "IPRS status", "Amount / period", "Reason",
              "Correction requested", "Expected distribution"],
             [[g["case_code"], g["title"], g["work_code"], g["isrcs"], run["platform"], run["territory"], period,
               OUTCOME_LABEL[g["outcome"]], _d(g["r_amount"]), g["n_streams"], _d(g["b_amount"]), _d(g["p_amount"]),
               _d(g["variance"]), P_STATUS_LABEL[g["p_status"]],
               f"sheet 6 rows for {g['case_code']}", g["question"], None, None, None, None, None] for g in enquiry],
             {9: MONEY, 10: INT, 11: MONEY, 12: MONEY, 13: MONEY})

    wb.properties.creator = "MRM Royalty Comparison"
    fixed = pack["built_at"].replace(tzinfo=None, microsecond=0)
    wb.properties.created = wb.properties.modified = fixed
    buf = io.BytesIO()
    wb.save(buf)
    return _stable_zip(buf.getvalue(), fixed)


def _stable_zip(data: bytes, when: dt.datetime) -> bytes:
    """Same pack content -> same bytes (zip entry times fixed)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(when.year, when.month, when.day, when.hour, when.minute,
                                                           when.second))
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, src.read(info.filename))
    return out.getvalue()


def _basis(files, g):
    ids = {x["file_id"] for x in g["evidence"]["r_rows"]}
    return ", ".join(sorted({(f["meta"].get("basis") or "?").lower() for f in files if f["id"] in ids})) or None


def _share(files, g):
    ids = {x["file_id"] for x in g["evidence"]["r_rows"]}
    return ", ".join(sorted({str(f["meta"].get("share") or "?") for f in files if f["id"] in ids})) or None


def _dist(files, g):
    ids = {x["file_id"] for x in g["evidence"]["p_rows"]}
    return ", ".join(sorted({f["meta"].get("distribution_ref") or f["filename"][:40] for f in files if f["id"] in ids})) or None


def _method(m):
    return {"MRM_BENCHMARK": "6% MRM benchmark (indicative)", "EVIDENCED": "evidenced method",
            "NO_BENCHMARK": "no monetary benchmark", "RULE_HOLD": "rule hold",
            "CONFIGURATION_HOLD": "configuration hold"}.get(m, m)


def _share_treatment(rules_, rid):
    r = next((x for x in rules_ if x["id"] == rid), None)
    if not r:
        return None
    return "beneficiary-specific" if r["share_treatment"] == "BENEFICIARY_SPECIFIC" else \
        f"100% of work x controlled share {r['controlled_share']}"


def _detail(d):
    if not d:
        return None
    keep = {k: v for k, v in d.items() if k not in ("candidates", "works", "isrcs_a", "isrcs_b")}
    return "; ".join(f"{k}: {v}" for k, v in keep.items())[:300]


def _owner(reason):
    return {"MAPPING": "rights_reviewer", "RIGHTS": "rights_reviewer", "BASIS": "accounts", "DATA_GAP": "importer",
            "PENDING_DISTRIBUTION": "accounts"}.get(reason or "", "importer")


# ------------------------------------------------------------------ H1-H4
def build(conn, actor: dict, run_id: int, note: str | None = None) -> dict:
    run = db.one(conn, "SELECT * FROM comparison_run WHERE id = %s", (run_id,))
    if not run or run["status"] != "COMPLETE":
        raise PackError("the comparison run is not complete")
    client = db.one(conn, "SELECT * FROM client WHERE id = %s", (run["client_id"],))
    route = client["collection_route"]                     # H2: the evidenced route, not assumed IPRS
    groups = db.all(conn, "SELECT * FROM comparison_group WHERE run_id = %s ORDER BY id", (run_id,))
    enquiry = [g for g in groups if g["outcome"] in ENQUIRY_OUTCOMES]
    rule = db.one(conn, "SELECT * FROM rule_version WHERE id = %s", ((run["rule_version_ids"] or [0])[0],))
    approver = db.one(conn, "SELECT u.name FROM app_user u JOIN user_role r ON r.user_id = u.id WHERE r.role = "
                            "'approver' AND u.active ORDER BY u.id LIMIT 1")
    subject, body = email(client, run, route, enquiry, run["summary"].get("method"), rule,
                          approver["name"] if approver else "MRM approver")
    prev = db.one(conn, "SELECT * FROM report_pack WHERE run_id = %s ORDER BY version_no DESC LIMIT 1", (run_id,))
    if prev and prev["status"] in ("DRAFT", "APPROVED"):
        raise PackError(f"pack v{prev['version_no']} for this run is {prev['status'].lower()} - request a revision "
                        f"before building a new version")
    no = (prev["version_no"] if prev else 0) + 1
    pid = db.insert(conn, "report_pack", dict(
        run_id=run_id, client_id=client["id"], version_no=no, route=route,
        recipient={"IPRS": "IPRS distribution team", "PRS": "PRS for Music", "DIRECT": "publisher accounts"}[route],
        status="DRAFT", email_subject=subject, email_body=body, case_codes=[g["case_code"] for g in enquiry],
        built_by=actor["id"], supersedes_id=prev["id"] if prev else None, review_note=note))
    pack = db.one(conn, "SELECT * FROM report_pack WHERE id = %s", (pid,))
    data = workbook(conn, pack, run, client, route, enquiry)
    sha = hashlib.sha256(data).hexdigest()
    key = f"{client['code']}/run{run['version_no']}-pack{no}-{sha[:12]}.xlsx"
    storage.put("packs", key, data)
    conn.execute("UPDATE report_pack SET object_key = %s, sha256 = %s WHERE id = %s", (key, sha, pid))
    for g in enquiry:
        _attach_case(conn, actor, client["id"], g, pid, run)
    audit.record(conn, "pack.built", actor, "report_pack", pid,
                 dict(run=run["version_no"], version=no, route=route, cases=len(enquiry), sha256=sha))
    return db.one(conn, "SELECT * FROM report_pack WHERE id = %s", (pid,))


def _attach_case(conn, actor, client_id, g, pack_id, run):
    c = db.one(conn, "SELECT * FROM enquiry_case WHERE client_id = %s AND group_key = %s FOR UPDATE",
               (client_id, g["group_key"]))
    if not c:
        cid = db.insert(conn, "enquiry_case", dict(client_id=client_id, case_code=g["case_code"], group_key=g["group_key"],
                                                   title=g["title"], work_code=g["work_code"], latest_group_id=g["id"],
                                                   pack_id=pack_id, status="OPEN",
                                                   next_action="approve and send the enquiry pack"))
        event(conn, cid, actor, "opened", dict(run=run["version_no"], outcome=g["outcome"], r=g["r_amount"],
                                              b=g["b_amount"], p=g["p_amount"], v=g["variance"]))
        return
    upd = dict(latest_group_id=g["id"], pack_id=pack_id)
    if c["status"] in ("CLOSED_EXPLAINED", "CLOSED_CORRECTED"):
        upd.update(status="OPEN", next_action="still flagged in a later run - review")
        event(conn, c["id"], actor, "reopened", dict(run=run["version_no"], outcome=g["outcome"],
                                                    previous=c["status"]))
    else:
        event(conn, c["id"], actor, "included_in_pack", dict(run=run["version_no"], outcome=g["outcome"],
                                                            r=g["r_amount"], b=g["b_amount"], p=g["p_amount"],
                                                            v=g["variance"]))
    db.update(conn, "enquiry_case", c["id"], dict(upd, updated_at=dt.datetime.now(dt.timezone.utc)))


def event(conn, case_id: int, actor: dict | None, kind: str, detail: dict) -> None:
    db.insert(conn, "case_event", dict(case_id=case_id, actor_id=(actor or {}).get("id"), kind=kind, detail=detail),
              returning=None)


def _pack(conn, pack_id: int) -> dict:
    p = db.one(conn, "SELECT * FROM report_pack WHERE id = %s FOR UPDATE", (pack_id,))
    if not p:
        raise PackError("unknown pack")
    return p


def approve(conn, actor: dict, pack_id: int, note: str | None = None) -> dict:
    p = _pack(conn, pack_id)
    if p["status"] != "DRAFT":
        raise PackError(f"pack is {p['status'].lower()}, not a draft")
    not_same_person(actor, p["built_by"], "a report pack")
    conn.execute("UPDATE report_pack SET status = 'APPROVED', approved_by = %s, approved_at = now(), review_note = "
                 "coalesce(%s, review_note) WHERE id = %s", (actor["id"], note, pack_id))
    audit.record(conn, "pack.approved", actor, "report_pack", pack_id, dict(note=note, sha256=p["sha256"]))
    return db.one(conn, "SELECT * FROM report_pack WHERE id = %s", (pack_id,))


def request_revision(conn, actor: dict, pack_id: int, note: str) -> dict:
    p = _pack(conn, pack_id)
    if p["status"] not in ("DRAFT", "APPROVED"):
        raise PackError(f"pack is {p['status'].lower()}")
    if not s(note):
        raise PackError("say what must be corrected (scope, totals, sources, confidential data or wording)")
    conn.execute("UPDATE report_pack SET status = 'REVISION_REQUESTED', review_note = %s WHERE id = %s",
                 (note, pack_id))
    audit.record(conn, "pack.revision_requested", actor, "report_pack", pack_id, dict(note=note))
    return db.one(conn, "SELECT * FROM report_pack WHERE id = %s", (pack_id,))


def mark_sent(conn, actor: dict, pack_id: int, reference: str, follow_up, sent_on=None) -> dict:
    p = _pack(conn, pack_id)
    if p["status"] != "APPROVED":
        raise PackError("only an approved pack can be recorded as sent")
    if not s(reference):
        raise PackError("record the dispatch reference (email subject / message ID / letter number)")
    fu = follow_up if isinstance(follow_up, dt.date) else (dt.date.fromisoformat(follow_up) if s(follow_up) else None)
    if not fu:
        raise PackError("set the follow-up date")
    sent_at = dt.datetime.fromisoformat(sent_on) if s(sent_on) else dt.datetime.now(dt.timezone.utc)
    conn.execute("UPDATE report_pack SET status = 'SENT', sent_by = %s, sent_at = %s, sent_reference = %s, "
                 "follow_up_date = %s WHERE id = %s", (actor["id"], sent_at, reference, fu, pack_id))
    for c in db.all(conn, "SELECT * FROM enquiry_case WHERE pack_id = %s", (pack_id,)):
        if c["status"] in ("OPEN", "AWAITING_REPLY"):
            db.update(conn, "enquiry_case", c["id"], dict(status="AWAITING_REPLY", follow_up_date=fu,
                                                          next_action=f"await reply; follow up on {fu:%d %b %Y}",
                                                          updated_at=dt.datetime.now(dt.timezone.utc)))
        event(conn, c["id"], actor, "sent", dict(pack=p["version_no"], reference=reference, follow_up=fu))
    audit.record(conn, "pack.sent", actor, "report_pack", pack_id, dict(reference=reference, follow_up=fu))
    return db.one(conn, "SELECT * FROM report_pack WHERE id = %s", (pack_id,))
