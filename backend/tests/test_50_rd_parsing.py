"""Reading the supplied files: S-1 and S-16 in detail, every reference statement, file-name metadata, file kinds,
zero / blank / negative amounts and the review-status vocabulary."""
from __future__ import annotations

import glob
import io
from decimal import Decimal

import openpyxl
import pytest

from mrm.rd import detect, names, parsers, rules
from mrm.rd.model import (PLATFORM_REVENUE, PLATFORM_USAGE, PREVIOUS_OUTPUT, REVIEW_LIST, SONG_LIST, STATEMENT,
                          UNKNOWN)
from rdflow import L6M, MRM, REF, REVIEW, S1, S16, S46, SKV8, SPOT_IPRS

pytestmark = pytest.mark.skipif(not S1.exists(), reason="reference files not present")
SVF = dict(code="SVF", name="SVF Entertainment Private Limited", holders=[dict(member_no="5154290")],
           settings=dict(distributors=[dict(name="MRM", currency="INR", basis="GROSS", platforms=["Spotify"])]))


def ident(path, name=None, forced=None):
    p = detect.identify(name or path.name, path.read_bytes(), forced)
    ev = detect.evaluate(p.kind, p.summary.get("facts", {}), name or path.name, SVF, None, p.summary.get("file_date"))
    return p, ev


def test_s1_mechanical_statement():
    p, ev = ident(S1)
    d = p.data
    assert p.kind == STATEMENT and d.layout == "IPRS_STANDARD" and d.sheet == "overseasMemberDetail"
    assert (d.holder_member, d.holder_name, d.ipi_name_no) == ("5154290", "SVF ENTERTAINMENT PRIVATE LIMITED",
                                                               "00871526032")
    assert len(d.lines) == 1686 and len({x.work_no for x in d.lines}) == 573
    assert d.printed_total == d.lines_total == d.subtotal_total == Decimal("82769.6412")
    assert {x.role for x in d.lines} == {"E"}
    pool = next(f for f in d.footers if f.name == "POOL")
    lang = next(f for f in d.footers if f.name == "SUMMARY OF ROYALTY")
    assert pool.rows_total == pool.grand_total == lang.grand_total == Decimal("82769.6412")
    assert len(pool.rows) == 4 and "MUSERK JULY" in pool.rows[0]["label"]
    # carried forward: the money lines of a block belong to the work on its first line
    aaj = [x for x in d.lines if x.work_no == "27855040"]
    assert [x.amount for x in aaj] == [Decimal("15.5356"), Decimal("26.8148"), Decimal("3.3018"), Decimal("28.9255")]
    assert aaj[0].title == "AAJ JYOTSNA RAATE" and aaj[0].source == "MUSERK JANUARY"
    m = ev["meta"]
    assert ev["status"] == "READY"
    assert (m["section"], m["dist_no"], m["period_label"]) == ("Mechanical (MUSERK)", "M2506",
                                                               "Apr 2025 - Mar 2026 (FY 2025-26)")


def test_s16_overseas_statement():
    p, ev = ident(S16)
    d = p.data
    assert d.layout == "IPRS_OVERSEAS" and d.sheet == "MemberDetail"
    assert d.printed_total == d.lines_total == d.subtotal_total == Decimal("10.2588")
    works = {x.work_no for x in d.lines}
    assert works == {"15662148", "25809104", "25809105", "25809107", "25809110"}
    zero_named = [w for w in works if all(x.amount == 0 for x in d.lines if x.work_no == w)]
    assert len(zero_named) == 4                                        # named, 0.00 paid - real lines, not blanks
    svf = [x for x in d.lines if x.work_no == "15662148" and x.role == "E"]
    assert [x.amount for x in svf] == [Decimal("0.1116"), Decimal("10.1472")]
    assert svf[1].usage["DEMAND"] == Decimal("9.664") and svf[1].usage["OTHERS"] == Decimal("0.4832")
    assert ev["meta"]["section"] == "Overseas - CASH (Hong Kong)" and ev["meta"]["dist_no"] == "P2555"
    assert ev["meta"]["period_label"] == "Apr 2024 - Mar 2025"


def test_every_reference_statement_reconciles():
    paths = sorted(glob.glob(str(REF / "batch-1" / "(S-*)Royalty*.xlsx")))
    assert len(paths) == 45
    printed = Decimal("0")
    layouts = {"IPRS_STANDARD": 0, "IPRS_OVERSEAS": 0}
    for path in paths:
        p = detect.identify(path.rsplit("/", 1)[1], open(path, "rb").read())
        d = p.data
        layouts[d.layout] += 1
        assert abs(d.lines_total - d.printed_total) <= Decimal("0.05"), path
        assert abs(d.subtotal_total - d.printed_total) <= Decimal("0.05"), path
        printed += d.printed_total
    assert layouts == {"IPRS_STANDARD": 27, "IPRS_OVERSEAS": 18}
    assert names.q2(printed) == Decimal("3209945.73")                  # = SVF-RD-SKV4's total


@pytest.mark.parametrize("name,cat,dist,period", [
    ("(S-2)Royalty Distribution - P2101 to P2540 - P2107A008~5154290.xlsx", "Other / Unclassified", "P2107A008", None),
    ("(S-44)Royalty Redistribution - for P2101 to P2550 - SVF ENTERTAINMENT PRIVATE LIMITED.xlsx", "Redistribution",
     "P2101 to P2550", None),
    ("(S-24)Royalty Distribution - P2567 Spotify - for the Period April 2025 to Septmber 2025.xlsx", "Spotify", "P2567",
     ("2025-04-01", "2025-09-30")),
    ("(S-9)Royalty Distribution - P2540 - Radio - for the period April 2023 to March 2025.xlsx", "Radio", "P2540",
     ("2023-04-01", "2025-03-31")),
    ("Spotify - for the Period October 2025 to March 2026(from IPRS).xlsx", "Spotify", None,
     ("2025-10-01", "2026-03-31")),
    ("(S-36)Royalty Distribution - P2580 Overseas Societies - April 2025 to December 2025 - 008 APRA.xlsx", "Overseas",
     "P2580", ("2025-04-01", "2025-12-31")),
])
def test_file_name_metadata(name, cat, dist, period):
    m = names.statement_meta_from_name(name)
    assert m.get("category") == cat and m.get("dist_no") == dist
    if period:
        assert (m["period_start"].isoformat(), m["period_end"].isoformat()) == period
    else:
        assert "period_start" not in m


def test_renamed_statement_gets_its_category_from_the_content():
    path = REF / "batch-1" / "(S-24)Royalty Distribution - P2567 Spotify - for the Period April 2025 to Septmber 2025.xlsx"
    p, ev = ident(path, "statement.xlsx")
    assert ev["meta"]["category"] == "Spotify" and ev["meta_source"]["category"] == "file content"
    assert [q["field"] for q in ev["problems"] if q.get("field")] == ["period", "dist_no"]   # asks only for what's missing
    ev2 = detect.evaluate(p.kind, p.summary["facts"], "statement.xlsx", SVF,
                          dict(period="Apr 2025 - Sep 2025", dist_no="P2567"), None)
    assert ev2["meta"]["period_label"] == "Apr 2025 - Sep 2025" and ev2["meta_source"]["dist_no"] == "you entered"
    assert ev2["status"] == "READY"


def test_file_kinds_are_detected_from_content():
    assert ident(S46)[0].kind == SONG_LIST
    assert ident(MRM)[0].kind == PLATFORM_REVENUE
    assert ident(L6M)[0].kind == PLATFORM_USAGE
    assert ident(SPOT_IPRS)[0].kind == STATEMENT
    p, ev = ident(REVIEW)
    assert p.kind == REVIEW_LIST and len(p.data.entries) == 93
    assert p.summary["facts"]["statuses"]["NOT_OURS"] == 34
    p, ev = ident(SKV8)
    assert p.kind == PREVIOUS_OUTPUT and len(p.data.rows) == 3144 and not p.data.integrity
    assert p.data.grand_total == Decimal("3323794.31")
    assert sum(1 for c in p.data.columns if c.kind == "ROYALTY") == 46


def test_mrm_report_metadata_comes_from_the_client_settings():
    p, ev = ident(MRM)
    m, src = ev["meta"], ev["meta_source"]
    assert (m["platform"], m["distributor"], m["currency"], m["basis"]) == ("Spotify", "MRM", "INR", "GROSS")
    assert src["currency"] == "client setting" and src["platform"] == "file content"
    assert Decimal(p.summary["facts"]["total"]).quantize(Decimal("0.01")) == Decimal("7889496.76")


def _xlsx(rows) -> bytes:
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _statement(lines, total, member=5154290):
    rows = [["INTERNAL NO", member, "IPI NAME NO", "00871526032"], ["NAME", "SVF ENTERTAINMENT PRIVATE LIMITED, "], [],
            ["WORK INT NO", "TITLE", "AV", "LANGUAGE", "NAME", "ROLE", "SOCIETY", "OWN", "COLL", "POOL", "SOURCE",
             "ROYALTY AMT"]]
    for work, title, amt in lines:
        rows.append([work, title, None, "BENGALI", "COMPOSER, A", "C", "036", 25, 25])
        rows.append([None, None, None, None, "SVF ENTERTAINMENT PRIVATE LIMITED, ", "E", "036", 50, 50, "INTERNET",
                     "SPOTIFY AB", amt])
        rows.append([None] * 11 + [amt])
    rows.append([None] * 7 + ["TOTAL ROYALTIES", None, None, None, total])
    return _xlsx(rows)


def test_blank_zero_and_negative_amounts():
    data = _statement([(111, "SONG A", 10.5), (222, "SONG B", 0), (333, "SONG C", -2.25), (444, "SONG D", None)], 8.25)
    p = detect.identify("(S-90)Royalty Distribution - P2999 Spotify - for the Period April 2025 to June 2025.xlsx",
                        data)
    d = p.data
    amounts = {x.work_no: x.amount for x in d.lines}
    assert amounts == {"111": Decimal("10.5"), "222": Decimal("0"), "333": Decimal("-2.25")}    # blank is not zero
    assert len(d.blank_amount_rows) == 1
    ev = detect.evaluate(p.kind, p.summary["facts"], "(S-90)Royalty Distribution - P2999 Spotify - for the Period "
                         "April 2025 to June 2025.xlsx", SVF, None, None)
    codes = {x["code"] for x in ev["problems"]}
    assert ev["status"] == "READY" and {"BLANK_AMOUNTS", "NEGATIVE_LINES"} <= codes


def test_a_statement_that_does_not_reconcile_is_left_out_with_the_difference_shown():
    data = _statement([(111, "SONG A", 10.5)], 99)
    name = "(S-91)Royalty Distribution - P2998 Spotify - for the Period April 2025 to June 2025.xlsx"
    p = detect.identify(name, data)
    ev = detect.evaluate(p.kind, p.summary["facts"], name, SVF, None, None)
    assert ev["status"] == "ERROR"
    msg = next(x for x in ev["problems"] if x["code"] == "RECON_MISMATCH")["message"]
    assert "₹10.50" in msg and "₹99.00" in msg and "-₹88.50" in msg


def test_macro_workbook_is_rejected_and_unknown_file_asks_one_question():
    p = detect.identify("x.xlsx", b"PK\x03\x04 not really a zip")
    assert p.kind == UNKNOWN and p.problems[0].severity == "ERROR"
    p = detect.identify("mystery.xlsx", _xlsx([["a", "b"], [1, 2]]))
    assert p.kind == UNKNOWN and p.problems[0].field == "kind"


@pytest.mark.parametrize("text,status", [
    ("Not SVF work", rules.NOT_OURS), ("already linked", rules.ALREADY_IN_LIST),
    ("Updated in SVF Report", rules.ALREADY_IN_LIST), ("need to update ISRC", rules.IDENTIFIER_TO_UPDATE),
    ("INS2X2500159 ISRC to update", rules.IDENTIFIER_TO_UPDATE),
    ("Need to update song in report", rules.IDENTIFIER_TO_UPDATE), ("Need to register", rules.CONFIRMED_MISSING),
    ("Need Hangover Songs", rules.CONFIRMED_MISSING),
    ("Need to register. Ask for all the Hooligaanism songs", rules.CONFIRMED_MISSING),
    ("Please confirm if these work belongs to svf or not. Internal no is registered under saregama", rules.NEEDS_CHECK),
    ("something else entirely", rules.NEEDS_CHECK), ("", None)])
def test_review_status_mapping(text, status):
    assert rules.map_status(text) == status


def test_previous_report_edited_after_export_is_detected():
    wb = openpyxl.load_workbook(SKV8)
    ws = wb.active
    ws["D3"] = 1.0                                                     # someone typed over a total
    buf = io.BytesIO()
    wb.save(buf)
    p = detect.identify("SVF-RD-SKV8.xlsx", buf.getvalue())
    assert p.kind == PREVIOUS_OUTPUT and p.data.integrity
    ev = detect.evaluate(p.kind, p.summary["facts"], "SVF-RD-SKV8.xlsx", SVF, None, None)
    assert ev["status"] == "NEEDS_REVIEW" and any(x["field"] == "accept_edited" for x in ev["problems"])
    parsers.looks_like_skv  # noqa: B018 - imported for the reader under test
