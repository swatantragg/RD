"""Matching, money and the three workbooks: engine rules on small examples, then the supplied reference files.

Expected values come from the reference workbooks (SVF-RD-SKV8, Non-SVF-Works-Revenue-From-IPRS-And-Spotify) and the
approved business decisions (S-32 excluded; title matches never move money; combined total kept in companions)."""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
from decimal import Decimal

import openpyxl
import pytest

from mrm.rd import assemble, detect, engine, excel, rules, verify
from mrm.rd.model import (PLATFORM_REVENUE, PLATFORM_USAGE, PREVIOUS_OUTPUT, REVIEW_LIST, SONG_LIST, STATEMENT,
                          PlatformLine, RoyaltyLine, SongEntry)
from rdflow import L6M, MRM, NON_SVF_REF, REF, REVIEW, S1, S46, SKV8, SPOT_IPRS

D = Decimal
SVF = dict(code="SVF", name="SVF Entertainment Private Limited", holders=[dict(member_no="5154290")],
           settings=dict(distributors=[dict(name="MRM", currency="INR", basis="GROSS", platforms=["Spotify"])]))


# ------------------------------------------------------------------ small, readable examples
def stmt(fid, lines, section="Spotify", period=("2025-04-01", "2025-09-30"), dist="P2567", total=None):
    ls = [RoyaltyLine(line_no=i + 1, row_no=i + 5, work_no=w, raw_work=w or "?", title=t, amount=D(str(a)))
          for i, (w, t, a) in enumerate(lines)]
    a, b = dt.date.fromisoformat(period[0]), dt.date.fromisoformat(period[1])
    from mrm.rd import names
    return engine.SrcStatement(file_id=fid, label=f"S-{fid}", filename=f"S-{fid}.xlsx", section=section, dist_no=dist,
                               period_label=names.period_label(a, b), period_start=a, period_end=b,
                               date=dt.date(2025, 11, 1), date_estimated=True, holder_member="5154290",
                               printed_total=total if total is not None else sum((x.amount for x in ls), D(0)),
                               lines=ls, identity_key=f"id{fid}")


def plat(fid, lines):
    return engine.SrcPlatform(file_id=fid, label=f"P-{fid}", filename="p.xlsx", platform="Spotify", distributor="MRM",
                              basis="GROSS", currency="INR",
                              lines=[PlatformLine(line_no=i + 1, row_no=i + 2, month=dt.date.fromisoformat(m), isrc=i_,
                                                  title=t, amount=D(str(a))) for i, (m, i_, t, a) in enumerate(lines)])


def song_list(entries):
    return engine.SrcList(file_id=99, label="S-46", entries=[
        SongEntry(entry_no=n + 1, row_no=n + 2, title=t, work_nos=w, isrcs=i) for n, (t, w, i) in enumerate(entries)])


def run(**kw):
    inp = engine.RunInput(client_code="SVF", client_name="SVF", ruleset=rules.RULESET_V1, **kw)
    res = engine.build(inp)
    assert all(c["ok"] for c in res.checks + verify.master(res)), [c for c in res.checks if not c["ok"]]
    return res


def row(res, title):
    return next(r for r in res.rows if r.title == title)


def test_a_work_on_several_list_rows_is_booked_once_on_the_title_match():
    res = run(statements=[stmt(1, [("100", "TOMAKE CHAI", 50)])],
              song_list=song_list([("Tomake Chai - Cover", ["100"], ["INS2X1600999"]),
                                   ("Tomake Chai", ["100"], ["INS2X1600134"])]))
    owner, sib = row(res, "Tomake Chai"), row(res, "Tomake Chai - Cover")
    assert owner.royalty == D("50.00") and sib.royalty == D("0")
    assert sib.booked_on == ["IPRS work 100 - royalty booked once, on 'Tomake Chai' (INS2X1600134)"]
    assert "separate recording of an IPRS work paid on another row" in sib.status
    assert res.totals["royalty"] == D("50.00")                        # never double counted


def test_a_recording_claimed_by_two_different_songs_is_held_until_assigned():
    lst = song_list([("Mon Bojhe Naa", ["1"], ["INS231421096"]), ("Eka Ekela Mon", ["2"], ["INS231421096"])])
    p = plat(5, [("2025-10-01", "INS231421096", "Eka Ekela Mon", "100.00")])
    res = run(song_list=lst, platforms=[p])
    held = next(r for r in res.rows if r.origin == "hold")
    assert held.platform and sum(held.platform.values()) == D("100.00")
    issue = next(i for i in res.issues if i["kind"] == "SHARED_RECORDING")
    assert len(issue["subject"]["candidates"]) == 2
    assert res.totals["held_platform"] == D("100.00") and sum(res.totals["platform"].values()) == D("100.00")
    dec = engine.Decision(id=7, kind="ASSIGN_HELD", subject=dict(code_kind="ISRC", code="INS231421096"),
                          value=dict(target=dict(code_kind="WORK", code="2", title="Eka Ekela Mon")))
    res2 = run(song_list=lst, platforms=[p], decisions=[dec])
    assert not [r for r in res2.rows if r.origin == "hold"]
    assert sum(row(res2, "Eka Ekela Mon").platform.values()) == D("100.00")


def test_a_title_never_moves_money_without_an_approved_decision():
    lst = song_list([("Amare Tumi", ["10"], ["INS2X0900001"])])
    p = plat(5, [("2025-10-01", "INS2X0900077", "Amare Tumi", "40.00")])
    res = run(song_list=lst, platforms=[p])
    paid = next(r for r in res.rows if r.origin == "recording")
    assert sum(paid.platform.values()) == D("40.00")
    assert sum(row_by_origin(res, "list").platform.values()) == D("0")          # the list row got nothing
    assert res.findings == [paid] and paid.proposals[0]["title"] == "Amare Tumi"
    assert any(i["kind"] == "PROPOSAL" for i in res.issues)
    link = engine.Decision(id=3, kind="LINK", subject=dict(code_kind="ISRC", code="INS2X0900077"),
                           value=dict(target=dict(code_kind="ISRC", code="INS2X0900001", title="Amare Tumi")))
    res2 = run(song_list=lst, platforms=[p], decisions=[link])
    base = row_by_origin(res2, "list")
    assert sum(base.platform.values()) == D("40.00") and base.all_isrcs == ["INS2X0900001", "INS2X0900077"]
    assert not res2.findings
    keep = engine.Decision(id=4, kind="NOT_SAME", subject=dict(code_kind="ISRC", code="INS2X0900077"),
                           value=dict(target=dict(code_kind="ISRC", code="INS2X0900001")))
    res3 = run(song_list=lst, platforms=[p], decisions=[keep])
    assert res3.findings[0].proposals == []                           # a 'keep separate' answer is never asked again


def row_by_origin(res, origin):
    return next(r for r in res.rows if r.origin == origin)


def test_already_in_the_previous_report():
    prior = assemble.prior_src(1, "SVF-RD-SKV8.xlsx", detect.identify("SVF-RD-SKV8.xlsx", SKV8.read_bytes()).data, {})
    p = detect.identify(S1.name, S1.read_bytes())
    ev = detect.evaluate(p.kind, p.summary["facts"], S1.name, SVF, None, p.summary.get("file_date"))
    s1 = assemble.statement_src(2, S1.name, p.data, ev["meta"])
    res = run(prior=prior, statements=[s1])
    assert any("already in SVF-RD-SKV8" in n for n in res.notices)
    assert res.totals["royalty"] == D("3323794.31")                   # S-1 not added twice
    changed = assemble.statement_src(2, S1.name, p.data, ev["meta"])
    changed.lines = changed.lines[:-1]
    changed.printed_total = sum((x.amount for x in changed.lines), D(0))
    res2 = run(prior=prior, statements=[changed])
    assert any(i["kind"] == "ALREADY_IN_BASE" and i["severity"] == "BLOCKING" for i in res2.issues)
    assert res2.totals["royalty"] == D("3323794.31")


def test_two_statements_with_the_same_identity_fail_the_build():
    with pytest.raises(engine.BuildFailed):
        engine.build(engine.RunInput(client_code="SVF", client_name="SVF", ruleset=rules.RULESET_V1,
                                     statements=[stmt(1, [("1", "A", 1)]), stmt(1, [("1", "A", 2)])]))


def test_financial_year_split_and_period_not_stated():
    radio = stmt(1, [("1", "SONG", "100.00")], section="Radio", period=("2023-04-01", "2025-03-31"), dist="P2540")
    other = stmt(2, [("1", "SONG", "10.00")], section="Other / Unclassified", period=("2023-04-01", "2023-04-30"),
                 dist="P2107A008")
    other.period_start = other.period_end = None
    other.period_label = "Not specified"
    res = run(statements=[radio, other])
    r = res.rows[0]
    assert r.fy == {2023: D("50.00"), 2024: D("50.00"), "NA": D("10.00")} and r.royalty == D("110.00")


def test_printed_cells_add_up_to_the_printed_totals():
    lines = [(str(i), f"S{i}", "0.333") for i in range(1, 10)]
    res = run(statements=[stmt(1, lines, total=D("2.997"))])
    col = res.bands[0].cols[0]
    assert col.shown_total == D("3.00") == sum((r.cells[col.key].shown for r in res.rows), D(0))


def test_money_kinds_are_never_added_in_the_master():
    res = run(statements=[stmt(1, [("1", "A", "5")])], platforms=[plat(2, [("2025-10-01", "INS2X0000001", "A", "7")])],
              song_list=song_list([("A", ["1"], ["INS2X0000001"])]))
    r = row(res, "A")
    assert r.royalty == D("5.00") and sum(r.platform.values()) == D("7.00")
    data, _ = excel.master(res, client_code="SVF", generation="SVF-RD-SKV9", state="DRAFT", build_ref="t",
                           ruleset=rules.RULESET_V1, software="t", inputs=[], left_out=[], decisions=[])
    ws = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
    assert ws["D3"].value == 5 and ws["E3"].value == 7
    total_rev = next(c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value == "Total Revenue")
    assert ws.cell(3, total_rev).value == 5                           # 'Total Revenue' is society royalty only


# ------------------------------------------------------------------ the supplied reference files
def _load(paths):
    inp = engine.RunInput(client_code="SVF", client_name="SVF Entertainment Private Limited", ruleset=rules.RULESET_V1,
                          generation="SVF-RD-SKV9")
    seen = {}
    for n, p in enumerate(paths, 1):
        par = detect.identify(p.name, p.read_bytes())
        ev = detect.evaluate(par.kind, par.summary["facts"], p.name, SVF, None, par.summary.get("file_date"))
        if ev["status"] != "READY":
            continue
        m = ev["meta"]
        if par.kind == STATEMENT:
            if m["identity_key"] in seen:                             # S-44 / S-45: the user answers 'use both'
                m["identity_key"] += f"#supplementary{n}"
            seen[m["identity_key"]] = n
            inp.statements.append(assemble.statement_src(n, p.name, par.data, m))
        elif par.kind == PLATFORM_REVENUE:
            inp.platforms.append(assemble.platform_src(n, p.name, par.data, m))
        elif par.kind == PLATFORM_USAGE:
            inp.usages.append(assemble.usage_src(n, p.name, par.data, m))
        elif par.kind == SONG_LIST:
            inp.song_list = assemble.list_src(n, p.name, par.data, m)
        elif par.kind == REVIEW_LIST:
            inp.reviews += assemble.review_marks(n, "review", par.data)
        elif par.kind == PREVIOUS_OUTPUT:
            inp.prior = assemble.prior_src(n, p.name, par.data, m)
    return inp


@pytest.fixture(scope="module")
def full():
    if not S1.exists():
        pytest.skip("reference files not present")
    paths = sorted((REF / "batch-1").glob("*.xlsx")) + [SPOT_IPRS, MRM, L6M, REVIEW]
    res = engine.build(_load(paths))
    data, meta = excel.master(res, client_code="SVF", generation="SVF-RD-SKV9", state="DRAFT", build_ref="t",
                              ruleset=rules.RULESET_V1, software="t", inputs=[], left_out=[], decisions=[])
    return res, data, meta


def test_full_reference_build_totals(full):
    res, _, _ = full
    assert res.ok and all(c["ok"] for c in verify.master(res))
    # SKV8 total 33,23,794.31 less the other member's statement (S-32, 53.47) - business decision Q7
    assert res.totals["royalty"] == D("3323794.31") - D("53.47")
    assert res.totals["platform"] == {"PLATFORM|Spotify|MRM|GROSS": D("7889496.76")}
    assert sum(1 for r in res.rows if r.origin == "list") == 2762
    assert res.totals["distributions"] == 45 and res.totals["platform_months"] == 6


def test_master_keeps_the_skv8_structure(full):
    _, data, _ = full
    ref = openpyxl.load_workbook(SKV8).worksheets[0]
    gen = openpyxl.load_workbook(io.BytesIO(data)).worksheets[0]
    strip = lambda v: re.sub(r"\s{2,}\(.*$", "", v or "")                       # noqa: E731
    ref_bands = [strip(c.value) for c in ref[1] if c.value]
    gen_bands = [strip(c.value) for c in gen[1] if c.value]
    assert gen_bands == ref_bands                                      # same sections, same order
    assert [c.value for c in gen[2]][:5] == [c.value for c in ref[2]][:5]
    assert (gen.freeze_panes, gen.row_dimensions[1].height, gen.row_dimensions[2].height) == ("F3", 22, 30)
    for col in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "BN", "BO", "BP", "BQ"):
        assert gen.column_dimensions[col].width == ref.column_dimensions[col].width, col
    ref_fy = [c.value for c in ref[2] if c.value and str(c.value).startswith(("FY", "Period Not", "Total Rev"))]
    gen_fy = [c.value for c in gen[2] if c.value and str(c.value).startswith(("FY", "Period Not", "Total Rev"))]
    assert gen_fy == ref_fy
    heads = [c.value for c in gen[2] if c.value]
    assert heads[-4:] == ["Catalogue Status", "Royalty Booked On", "Spotify (MRM) Match Basis", "Held / Flags"]
    assert gen["F3"].number_format == "DD-MMM-YYYY" and gen["D3"].number_format == "#,##0.00"
    assert gen["F1"].fill.fgColor.rgb == ref["F1"].fill.fgColor.rgb    # measured colour system


def test_master_distribution_totals_match_skv8(full):
    _, data, _ = full
    def totals(ws):
        tr = next(r for r in range(3, ws.max_row + 1) if ws.cell(r, 2).value == "TOTAL")
        bands = [(c.column, c.value) for c in ws[1] if c.value]
        out = {}
        for c in range(1, ws.max_column + 1):
            if ws.cell(2, c).value == "Amount":
                band = re.sub(r"\s{2,}\(.*$", "", max((b for b in bands if b[0] <= c), key=lambda b: b[0])[1])
                out.setdefault((band, ws.cell(tr, c + 2).value), []).append(D(str(ws.cell(tr, c).value)))
        return out
    ref = totals(openpyxl.load_workbook(SKV8).worksheets[0])
    gen = totals(openpyxl.load_workbook(io.BytesIO(data)).worksheets[0])
    ref[("YOUTUBE POST-CLAIMS", "P2577")].remove(D("53.47"))             # S-32 belongs to member 8875288
    for k, v in ref.items():
        if k[0].startswith("SPOTIFY - GROSS"):
            continue                                                   # month columns: checked below
        assert sorted(gen[k]) == sorted(v), k
    ref_months = sum((sum(v) for k, v in ref.items() if k[0].startswith("SPOTIFY - GROSS")), D(0))
    gen_months = sum((sum(v) for k, v in gen.items() if k[0].startswith("SPOTIFY - GROSS")), D(0))
    assert ref_months == gen_months == D("7889496.76")


def test_reports_are_reproducible_byte_for_byte(full):
    res, data, _ = full
    again, _ = excel.master(res, client_code="SVF", generation="SVF-RD-SKV9", state="DRAFT", build_ref="t",
                            ruleset=rules.RULESET_V1, software="t", inputs=[], left_out=[], decisions=[])
    assert hashlib.sha256(again).digest() == hashlib.sha256(data).digest()


def test_previous_report_as_base_carries_every_rupee(full):
    inp = _load([SKV8])
    res = engine.build(inp)
    assert res.totals["royalty"] == D("3323794.31") and res.totals["rows"] == 3144
    assert sum(1 for i in res.issues if i["kind"] == "TITLE_PLACED_IN_BASE") == 81


def test_non_svf_report_reproduces_the_reference():
    """Spotify sources only, as the reference was built: the songs SVF marked 'Not SVF work' with the same money."""
    if not S1.exists():
        pytest.skip("reference files not present")
    res = engine.build(_load([S46, SPOT_IPRS, MRM, REVIEW]))
    groups = excel.companion_groups(res, rules.NOT_OURS)
    assert all(c["ok"] for c in verify.companion(res, groups))
    data = excel.companion(res, groups, kind="NON_CLIENT", client_code="SVF", client_name="SVF Entertainment",
                           generation="SVF-RD-SKV9", state="DRAFT", combined_total=True)
    ws_ref = openpyxl.load_workbook(NON_SVF_REF, read_only=True).worksheets[0]
    ws_gen = openpyxl.load_workbook(io.BytesIO(data), read_only=True).worksheets[0]
    def songs(ws):
        out = {}
        for r in ws.iter_rows(min_row=3, values_only=True):
            if r and r[1] and not str(r[1]).startswith("TOTAL") and r[0]:
                out[str(r[0])] = (D(str(r[3] or 0)), D(str(r[4] or 0)))
        return out
    ref, gen = songs(ws_ref), songs(ws_gen)
    assert set(gen) == set(ref) and len(gen) == 34
    for isrc, (iprs, spot) in ref.items():
        assert gen[isrc][0] == iprs and abs(gen[isrc][1] - spot) <= D("0.01"), isrc
    assert ws_gen.title == "Non-SVF Works Revenue"
    hdr = [c for c in next(ws_gen.iter_rows(min_row=2, max_row=2, values_only=True))]
    assert hdr[:5] == ["ISRC", "Song Name", "Internal No", "Total Royalty From IPRS", "Total Revenue From Spotify (MRM)"]
    assert "Total (Both Sources)" in hdr                               # business decision Q1


def test_missing_songs_rule_lists_only_codes_outside_the_song_list():
    if not S1.exists():
        pytest.skip("reference files not present")
    res = engine.build(_load([S46, SPOT_IPRS, MRM, REVIEW]))
    groups = excel.companion_groups(res)
    checks = verify.companion(res, groups)
    assert all(c["ok"] for c in checks), checks
    base = {c for r in res.rows if r.in_base for c in r.all_isrcs + r.all_works}
    assert not any(c in base for g in groups for c in g["isrcs"] + g["works"])
    statuses = [g["status"] for g in groups]
    first_resolved = next(i for i, s in enumerate(statuses) if s not in rules.UNRESOLVED)
    assert all(s not in rules.UNRESOLVED for s in statuses[first_resolved:])   # unreviewed songs listed first


def test_missing_songs_report_can_be_reviewed_in_excel_and_uploaded_back():
    """Download the Missing Songs report, type statuses in Excel, upload it: the statuses come back by code."""
    if not S1.exists():
        pytest.skip("reference files not present")
    res = engine.build(_load([S46, SPOT_IPRS, MRM]))
    groups = excel.companion_groups(res)
    data = excel.companion(res, groups, kind="MISSING", client_code="SVF", client_name="SVF", generation="SVF-RD-SKV9",
                           state="DRAFT", combined_total=True)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    col = next(c for c in range(1, ws.max_column + 1) if ws.cell(2, c).value == "Status Given In Review")
    assert ws.cell(3, col).value == "Not reviewed yet"
    ws.cell(3, col).value = "Not SVF work"
    ws.cell(4, col).value = "Need to register"
    buf = io.BytesIO()
    wb.save(buf)
    p = detect.identify("Missing-Songs-reviewed.xlsx", buf.getvalue())
    assert p.kind == REVIEW_LIST
    marks = assemble.review_marks(1, "x", p.data)
    assert [m.status for m in marks] == [rules.NOT_OURS, rules.CONFIRMED_MISSING]   # 'Not reviewed yet' is no answer
    first = groups[0]
    assert set(marks[0].isrcs + marks[0].work_nos) == set(first["isrcs"] + first["works"])
