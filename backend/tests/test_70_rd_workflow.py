"""The whole workflow through the HTTP API with the supplied files: select client, upload, check, process, review,
submit, approve (by someone else), publish, download - plus client isolation and secure downloads."""
from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
import psycopg
import pytest

from conftest import ADMIN, DBNAME
from mrm import db
from mrm.rd import jobs
from rdflow import (MRM, REVIEW, S1, S16, S32, S44, S45, S46, SPOT_IPRS, login, process, report, svf_id, upload)

pytestmark = pytest.mark.skipif(not S1.exists(), reason="reference files not present")


def _wb(resp) -> openpyxl.Workbook:
    assert resp.status_code == 200, resp.text
    assert resp.headers["content-disposition"].startswith("attachment")
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def test_full_workflow_publish_and_download():
    sw, sh, viewer = login("swatantra@test.local"), login("sherley@test.local"), login("viewer@test.local")
    cid = svf_id(sw)
    run = sw.post("/api/runs", json=dict(client_id=cid)).json()
    rid = run["id"]
    out = upload(sw, rid, S1, S16, SPOT_IPRS, MRM, S46, REVIEW, S32)
    assert all("error" not in x for x in out), out
    v = sw.get(f"/api/runs/{rid}").json()
    st = {f["filename"][:6]: f["status"] for f in v["files"]}
    assert st["(S-1)R"] == "READY" and st["(S-16)"] == "READY" and st["(S-32)"] == "LEFT_OUT"
    s32 = next(f for f in v["files"] if f["filename"].startswith("(S-32)"))
    assert "member 8875288" in s32["problems"][0]["message"]           # plain-language reason
    assert v["can_process"], v["blocked"]
    # the same file again is recognised, not counted twice
    again = upload(sw, rid, S1)[0]
    assert "already" in again["note"]
    b = process(sw, rid)
    assert b["state"] == "CHECKED", (b["error"], b["stages"])
    assert [s["state"] for s in b["stages"]] == ["DONE"] * 5
    assert Decimal(b["summary"]["royalty"]) == Decimal("82769.64") + Decimal("10.26") + Decimal("113848.57")
    assert b["summary"]["statements"] == 3 and b["summary"]["platform_reports"] == 1
    assert all(c["ok"] for c in b["checks"] if c.get("blocking", True))
    kinds = {r["kind"]: r for r in b["reports"]}
    assert set(kinds) == {"MASTER", "MISSING", "NON_CLIENT"} and all(r["available"] for r in kinds.values())
    assert kinds["MASTER"]["state"]["code"] == "VALIDATED"
    # draft downloads: preparer yes, viewer no
    wb = _wb(sw.get(f"/api/reports/{kinds['MASTER']['id']}/download"))
    assert wb.sheetnames[:2] == ["SVF-RD-SKV9", "Run Summary"] and "_rd_meta" in wb.sheetnames
    assert viewer.get(f"/api/reports/{kinds['MASTER']['id']}/download").status_code == 403
    # submit; the preparer cannot approve their own work; the approver can
    assert sw.post(f"/api/builds/{b['id']}/submit").status_code == 200
    assert sw.post(f"/api/builds/{b['id']}/approve").status_code == 403
    b2 = sh.get(f"/api/builds/{b['id']}").json()
    assert b2["state"] == "SUBMITTED" and b2["can_approve"]
    assert sh.post(f"/api/builds/{b['id']}/approve").status_code == 200
    jobs.run_pending()
    b3 = sh.get(f"/api/builds/{b['id']}").json()
    assert b3["state"] == "PUBLISHED", b3["returned_reason"]
    assert b3["generation"] == "SVF-RD-SKV9"
    final = {r["kind"]: r for r in b3["reports"]}
    assert all(r["final"] for r in final.values())
    wbf = _wb(viewer.get(f"/api/reports/{final['MASTER']['id']}/download"))      # published: any granted person
    assert wbf.sheetnames[0] == "SVF-RD-SKV9"
    nonsvf = _wb(viewer.get(f"/api/reports/{final['NON_CLIENT']['id']}/download"))
    assert nonsvf.active.title == "Non-SVF Works Revenue"
    # a published report cannot be changed, even by the application's database login
    with db.tx() as conn:
        conn.execute("SELECT set_config('app.client_ids', %s, true)", (str(cid),))
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            conn.execute("UPDATE rd_build SET summary = '{}' WHERE id = %s", (b["id"],))
    # next publication gets the next generation; the earlier one is kept and marked superseded
    b4 = process(sw, rid)
    sw.post(f"/api/builds/{b4['id']}/submit")
    sh.post(f"/api/builds/{b4['id']}/approve")
    jobs.run_pending()
    assert sh.get(f"/api/builds/{b4['id']}").json()["generation"] == "SVF-RD-SKV10"
    assert sh.get(f"/api/builds/{b['id']}").json()["state"] == "SUPERSEDED"
    hist = viewer.get("/api/past-reports").json()
    assert [h["generation"] for h in hist if h["state"] in ("PUBLISHED", "SUPERSEDED")][:2] == \
        ["SVF-RD-SKV10", "SVF-RD-SKV9"]


def test_same_distribution_twice_needs_a_decision():
    sw = login("swatantra@test.local")
    rid = sw.post("/api/runs", json=dict(client_id=svf_id(sw), title="redistribution check")).json()["id"]
    upload(sw, rid, S44, S45)
    v = sw.get(f"/api/runs/{rid}").json()
    assert not v["can_process"] and v["issues"][0]["kind"] == "SAME_DISTRIBUTION"
    act = next(a for a in v["issues"][0]["actions"] if a["action"] == "SUPPLEMENTARY")
    v = sw.post(f"/api/runs/{rid}/files/{act['file_id']}/choice", json=dict(choice="SUPPLEMENTARY")).json()
    assert v["can_process"], v["blocked"]
    b = process(sw, rid)
    assert b["state"] == "CHECKED"
    assert Decimal(b["summary"]["royalty"]) == Decimal("60.73") + Decimal("2435.92")


def test_unreadable_and_unknown_files_are_explained():
    sw = login("swatantra@test.local")
    rid = sw.post("/api/runs", json=dict(client_id=svf_id(sw), title="bad files")).json()["id"]
    r = sw.post(f"/api/runs/{rid}/files", files=[("files", ("notes.xlsx", b"this is not a workbook", "text/plain"))])
    f = r.json()[0]["file"]
    assert f["status"] in ("UNSUPPORTED", "ERROR") and "Excel" in f["problems"][0]["message"]
    wb = openpyxl.Workbook()
    wb.active.append(["Hello", "World"])
    wb.active.append([1, 2])
    buf = io.BytesIO()
    wb.save(buf)
    f = sw.post(f"/api/runs/{rid}/files", files=[("files", ("mystery.xlsx", buf.getvalue(), "x"))]).json()[0]["file"]
    assert f["status"] == "NEEDS_REVIEW" and f["questions"][0]["field"] == "kind"
    v = sw.get(f"/api/runs/{rid}").json()
    assert not v["can_process"] and any("answer" in x for x in v["blocked"])
    sw.delete(f"/api/runs/{rid}/files/{f['id']}")


def test_client_isolation_and_download_authorisation(client_id):
    from mrm.seed import create_user
    with db.tx() as conn:
        uid = create_user(conn, "other.preparer@test.local", "Other Preparer", ["importer"], "test-password-1")
        conn.execute("INSERT INTO client_grant (user_id, client_id) VALUES (%s, %s)", (uid, client_id))
    other = login("other.preparer@test.local")
    assert [c["id"] for c in other.get("/api/clients").json()] == [client_id]
    sw = login("swatantra@test.local")
    runs = sw.get("/api/runs").json()
    svf_run = next(r for r in runs if r["client_code"] == "SVF")
    assert other.get(f"/api/runs/{svf_run['id']}").status_code in (400, 403, 404)
    assert other.post("/api/runs", json=dict(client_id=svf_id(sw))).status_code == 403
    with psycopg.connect(ADMIN.rsplit("/", 1)[0] + f"/{DBNAME}") as su:
        rid = su.execute("SELECT id FROM rd_report LIMIT 1").fetchone()[0]
    assert other.get(f"/api/reports/{rid}/download").status_code == 403
    # the database itself refuses rows of another client when the scope is set to this one
    with db.tx() as conn:
        conn.execute("SELECT set_config('app.client_ids', %s, true)", (str(client_id),))
        assert db.scalar(conn, "SELECT count(*) FROM rd_file") == 0
        conn.execute("SELECT set_config('app.client_ids', '', true)")
        assert db.scalar(conn, "SELECT count(*) FROM rd_run") == 0


def test_review_decisions_need_a_second_person_and_carry_into_the_next_run():
    sw, sh = login("swatantra@test.local"), login("sherley@test.local")
    rid = sw.post("/api/runs", json=dict(client_id=svf_id(sw), title="review decisions")).json()["id"]
    upload(sw, rid, SPOT_IPRS, MRM, S46)
    b = process(sw, rid)
    prop = next(i for i in b["issues"] if i["kind"] == "PROPOSAL")
    before = b["summary"]["missing"]
    r = sw.post(f"/api/builds/{b['id']}/issues/{prop['no']}/act", json=dict(action="CONFIRM_MATCH"))
    assert r.status_code == 200 and "approver" in r.json()["message"]
    did = r.json()["decision_id"]
    assert sw.post(f"/api/decisions/{did}/approve").status_code == 403            # not the maker, not a preparer
    waiting = sh.get("/api/approvals").json()["decisions"]
    assert any(d["id"] == did and d["can_approve"] for d in waiting)
    assert sh.post(f"/api/decisions/{did}/approve", json=dict(note="checked the ISRC")).status_code == 200
    # a paid song marked 'Not ours' lands in the out-of-scope report; the status moves no money
    f = sw.get(f"/api/builds/{b['id']}/findings").json()["items"][0]
    assert sw.post(f"/api/builds/{b['id']}/findings/status", json=dict(findings=[f["no"]], status="NOT_OURS")).status_code == 200
    assert sw.get(f"/api/runs/{rid}").json()["changed_since_build"]
    b2 = process(sw, rid)
    assert b2["summary"]["missing"] == before - 1                                  # the confirmed match left the list
    assert Decimal(b2["summary"]["royalty"]) == Decimal(b["summary"]["royalty"])  # no money created or lost
    non = report(b2, "NON_CLIENT")
    assert non["available"] and non["rows"] >= 1
    # the decisions are keyed by codes: a new run with the same files gets the same answers without asking
    rid2 = sw.post("/api/runs", json=dict(client_id=svf_id(sw), title="next month")).json()["id"]
    upload(sw, rid2, SPOT_IPRS, MRM, S46)
    b3 = process(sw, rid2)
    assert b3["summary"]["missing"] == b2["summary"]["missing"]
    assert b3["summary"]["findings_by_status"].get("NOT_OURS", 0) >= 1
    assert sw.post(f"/api/runs/{rid2}/close").status_code == 200
    assert rid2 not in [r["id"] for r in sw.get("/api/runs").json()]


def test_viewer_cannot_prepare_and_drafts_are_hidden_from_viewers():
    v = login("viewer@test.local")
    assert v.post("/api/runs", json=dict(client_id=svf_id(v))).status_code == 403
    assert all(e["state"] in ("PUBLISHED", "SUPERSEDED") for e in v.get("/api/past-reports").json())
