"""Workbook -> grid, with the dimension repair (architecture §5.1, §22).

Several statements carry a 1x1 dimension record, and openpyxl's read-only reader stops at
the declared dimension. `reset_dimensions()` drops the record so every row is read.
"""
from __future__ import annotations

import io
import warnings

import openpyxl

warnings.filterwarnings("ignore", module="openpyxl")


def read_grid(data: bytes, max_rows: int | None = None) -> tuple[str, list[list]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = None
        for cand in wb.worksheets:            # first worksheet that holds anything
            cand.reset_dimensions()
            ws = cand
            first = next(cand.iter_rows(max_row=1, values_only=True), None)
            if first and any(v is not None for v in first):
                break
        if ws is None:
            return "", []
        ws.reset_dimensions()
        rows: list[list] = []
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))
            if max_rows and len(rows) >= max_rows:
                break
        title = ws.title
    finally:
        wb.close()
    while rows and all(v is None or (isinstance(v, str) and not v.strip()) for v in rows[-1]):
        rows.pop()
    width = max((len(r) for r in rows), default=0)
    for r in rows:
        if len(r) < width:
            r.extend([None] * (width - len(r)))
    return title, rows


def preview(grid: list[list], n: int = 8) -> list[list]:
    """First rows as JSON-safe values (stored on source_file for quarantined files)."""
    out = []
    for r in grid[:n]:
        out.append([None if v is None else (v if isinstance(v, (int, float, str)) else str(v))
                    for v in r[:20]])
    return out
