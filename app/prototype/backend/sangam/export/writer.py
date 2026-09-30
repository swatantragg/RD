"""Band-table writer shared by every SKV-styled workbook, plus the determinism pass.

Column roles are declared up front and number formats are applied BY ROLE, never by
inspecting the Python type of a value (defect D5), and TOTAL rows are computed from the
column roles, never from arithmetic on band offsets (defect D4).
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border
from openpyxl.utils import get_column_letter

from . import style as S

FIXED_TS = "2000-01-01T00:00:00Z"
ZIP_TS = (1980, 1, 1, 0, 0, 0)


@dataclass
class Col:
    header: str
    width: float
    role: str                                        # text | money | date | int | pct
    get: Callable[[Any], Any]
    bold: bool = False
    total: bool = False                              # summed on the TOTAL row


@dataclass
class Band:
    title: str | None
    fill: str | None                                 # band / header colour (None = spacer)
    tint: str | None                                 # data-row tint (None = identity stripe)
    cols: list[Col] = field(default_factory=list)
    header_fill: str | None = None
    header_fills: list[str] | None = None            # per-column header fill override


def _fmt(cell, role):
    if role == "money":
        cell.number_format = S.MONEY
    elif role == "date":
        cell.number_format = S.DATE
    elif role == "pct":
        cell.number_format = "0.00%"


def write_bands(ws, bands: list[Band], rows: list, *, first_row: int = 3, freeze: str | None = "F3",
                total_label: str | None = "TOTAL", total_label_col: int = 2, total_gap: int = 1,
                total_values: dict[int, Any] | None = None, formulas: bool = False) -> dict:
    """Row 1 = merged section bands, row 2 = headers, rows 3.. = data, then the TOTAL row."""
    col = 1
    spans = []
    for b in bands:
        start, end = col, col + len(b.cols) - 1
        spans.append((b, start, end))
        col = end + 1
    ncol = col - 1
    # ---- row 1 bands
    for b, start, end in spans:
        if b.title:
            c = ws.cell(row=1, column=start, value=b.title)
            c.font = S.BAND_FONT
            c.alignment = S.CENTER
            for cc in range(start, end + 1):
                ws.cell(row=1, column=cc).fill = S.fill(b.fill)
            ws.cell(row=1, column=start).border = Border(left=S.MED, top=S.MED, bottom=S.THIN)
            if end > start:
                ws.cell(row=1, column=end).border = Border(right=S.MED, top=S.MED, bottom=S.THIN)
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
    # ---- row 2 headers
    for b, start, end in spans:
        if b.fill is None:                               # spacer band: narrow, unstyled
            for j, cdef in enumerate(b.cols):
                ws.column_dimensions[get_column_letter(start + j)].width = cdef.width
            continue
        for j, cdef in enumerate(b.cols):
            c = ws.cell(row=2, column=start + j, value=cdef.header)
            c.font = S.HEAD_FONT
            hf = (b.header_fills[j] if b.header_fills else None) or b.header_fill or b.fill
            c.fill = S.fill(hf)
            c.alignment = S.WRAP_CENTER
            c.border = Border(top=S.THIN, bottom=S.MED, left=S.MED if (j == 0 and b.tint) else S.THIN,
                              right=S.THIN)
            ws.column_dimensions[get_column_letter(start + j)].width = cdef.width
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 30
    # ---- data rows
    r = first_row
    sums: dict[int, float] = {}
    stripe = S.fill(S.IDENTITY[1])
    left_med = Border(left=S.MED)
    for n, item in enumerate(rows):
        for b, start, end in spans:
            if b.fill is None:
                continue
            tint = S.fill(b.tint) if b.tint else (stripe if n % 2 == 0 else None)
            for j, cdef in enumerate(b.cols):
                v = cdef.get(item)
                c = ws.cell(row=r, column=start + j, value=v)
                if cdef.role != "text":
                    _fmt(c, cdef.role)
                if cdef.bold:
                    c.font = S.BOLD
                if tint is not None:
                    c.fill = tint
                if cdef.total and isinstance(v, (int, float)):
                    sums[start + j] = sums.get(start + j, 0.0) + v
            if b.tint:
                ws.cell(row=r, column=start).border = left_med
        r += 1
    last = r - 1
    total_row = None
    if total_label:
        total_row = last + 1 + total_gap
        ws.cell(row=total_row, column=total_label_col, value=total_label).font = S.BOLD
        for b, start, end in spans:
            if b.fill is None:
                continue
            for j, cdef in enumerate(b.cols):
                cc = start + j
                c = ws.cell(row=total_row, column=cc)
                c.fill = S.fill(b.header_fill or b.fill)
                c.border = Border(top=S.MED, bottom=S.MED, left=S.MED if (j == 0 and b.tint) else None)
                if cdef.total:
                    if formulas:
                        L = get_column_letter(cc)
                        c.value = f"=SUM({L}{first_row}:{L}{last})" if last >= first_row else 0
                    else:
                        c.value = round(sums.get(cc, 0.0) + 0.0, 2)
                    _fmt(c, cdef.role)
                    c.font = S.BOLD
        for cc, v in (total_values or {}).items():
            ws.cell(row=total_row, column=cc, value=v).font = S.BOLD
    if freeze:
        ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A2:{get_column_letter(ncol)}{max(last, 2)}"
    return dict(ncol=ncol, last=last, total_row=total_row, spans=[(b.title, s, e) for b, s, e in spans],
                sums=sums)


def note(ws, row: int, text: str, ncol: int, height: float | None = None, italic: bool = False,
         start_col: int = 1) -> None:
    c = ws.cell(row=row, column=start_col, value=text)
    c.font = S.NOTE_ITALIC if italic else S.NOTE_FONT
    c.alignment = S.WRAP_TOP
    if ncol > start_col:
        ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=ncol)
    if height:
        ws.row_dimensions[row].height = height


def new_workbook(title: str):
    wb = Workbook()
    wb.properties.creator = "Sangam"
    wb.properties.created = dt.datetime(2000, 1, 1)
    ws = wb.active
    ws.title = title[:31]
    return wb, ws


def save_deterministic(wb, path: Path) -> tuple[int, str]:
    """Save, then rewrite the zip with fixed timestamps and fixed document dates, so the same
    content always gives the same bytes (I25)."""
    buf = io.BytesIO()
    wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                txt = data.decode("utf-8")
                txt = re.sub(r"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)", rf"\g<1>{FIXED_TS}\g<2>", txt)
                txt = re.sub(r"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rf"\g<1>{FIXED_TS}\g<2>", txt)
                data = txt.encode("utf-8")
            zi = zipfile.ZipInfo(info.filename, date_time=ZIP_TS)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16
            z.writestr(zi, data)
    raw = out.getvalue()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(raw)
    return len(raw), hashlib.sha256(raw).hexdigest()


def as_date(v):
    if isinstance(v, str) and v:
        return dt.datetime.fromisoformat(v)
    if isinstance(v, dt.date) and not isinstance(v, dt.datetime):
        return dt.datetime(v.year, v.month, v.day)
    return v


def center(ws, cells):
    for c in cells:
        c.alignment = Alignment(horizontal="center")
