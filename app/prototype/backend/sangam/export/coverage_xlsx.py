"""§11.5 / §11.6 coverage deliverables: the 3-column finding list and the SKV-styled revenue
workbook. The master build supplies layout and colour only - every figure is read from the
coverage run, which recomputed it from the source lines (defect D7)."""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Border, PatternFill

from ..core import db
from ..domain.filename import SECTION_ORDER, dist_display
from ..domain.normalize import JOIN, fy_label, month_end_of, month_label
from . import style as S
from .writer import Band, Col, as_date, new_workbook, note, save_deterministic, write_bands


def load_run(conn, run_id: int) -> dict:
    run = db.row(conn, "SELECT * FROM coverage_run WHERE id = ?", (run_id,))
    run["stats"] = db.loads(run["stats_json"], {})
    run["sources"] = db.loads(run["sources_json"], [])
    run["sheet"] = db.row(conn, "SELECT * FROM user_sheet WHERE id = ?", (run["sheet_id"],))
    rows = db.rows(conn, "SELECT * FROM coverage_row WHERE run_id = ? ORDER BY ordinal", (run_id,))
    for r in rows:
        r["detail"] = db.loads(r.pop("detail_json"), {})
        r["isrcs"] = [x["isrc"] for x in db.rows(conn, "SELECT isrc FROM coverage_row_isrc WHERE row_id = ? "
                                                       "ORDER BY ordinal", (r["id"],))]
        r["work_nos"] = [x["work_no"] for x in db.rows(conn, "SELECT work_no FROM coverage_row_work WHERE "
                                                              "row_id = ? ORDER BY ordinal", (r["id"],))]
    run["rows"] = rows
    sids = run["stats"].get("statement_ids", [])
    run["statements"] = db.rows(conn, f"SELECT * FROM statement WHERE id IN ({','.join('?' * len(sids))}) "
                                      f"ORDER BY s_no", sids) if sids else []
    rids = run["stats"].get("report_ids", [])
    run["reports"] = db.rows(conn, f"SELECT * FROM platform_report WHERE id IN ({','.join('?' * len(rids))})",
                             rids) if rids else []
    return run


def banner(run: dict) -> str:
    sh = run["sheet"]
    per = (f"{run['period_from'] or '...'} to {run['period_to'] or '...'}"
           if (run["period_from"] or run["period_to"]) else "all periods")
    return (f"SONGS NOT AVAILABLE IN THE SONG SHEET  -  {sh['label']}  |  sources: {', '.join(run['sources'])}  |  "
            f"period: {per}  |  mode: {run['mode']}  |  every row verified absent from the sheet on BOTH internal "
            f"number and ISRC")


def write_list(run: dict, path) -> dict:
    wb, ws = new_workbook("Songs Not In SVF List")
    ws.merge_cells("A1:C1")
    c = ws["A1"]
    c.value = banner(run)
    c.font = S.BAND_FONT
    c.fill = S.fill(S.SECTION_PALETTE["Spotify MRM"][0])
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 48
    for i, h in enumerate(["ISRC", "Song Name", "Internal No"], 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = S.HEAD_FONT
        c.fill = S.fill(S.IDENTITY[0])
        c.alignment = S.WRAP_CENTER
        c.border = S.BOX
    ws.row_dimensions[2].height = 30
    for n, r in enumerate(run["rows"]):
        for i, v in enumerate([JOIN.join(r["isrcs"]) or None, r["song_name"], JOIN.join(r["work_nos"]) or None], 1):
            c = ws.cell(row=3 + n, column=i, value=v)
            if n % 2 == 0:
                c.fill = S.fill(S.IDENTITY[1])
    last = 2 + len(run["rows"])
    for i in range(1, 4):
        cc = ws.cell(row=last + 2, column=i)
        cc.fill = S.fill(S.IDENTITY[0])
        cc.border = Border(top=S.MED)
        cc.font = S.BOLD
    ws.cell(row=last + 2, column=2, value="TOTAL")
    ws.cell(row=last + 2, column=3, value=f"{len(run['rows'])} songs")
    for col, w in zip("ABC", (34, 46, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:C{max(last, 2)}"
    size, sha = save_deterministic(wb, path)
    return dict(size=size, sha256=sha, rows=len(run["rows"]))


def write_revenue(run: dict, path, client: str = "Spotify") -> dict:
    sts = run["statements"]
    months = run["stats"].get("months", [])
    sections = []
    for sec in sorted({s["section"] for s in sts}, key=lambda x: (0, SECTION_ORDER.index(x)) if x in SECTION_ORDER
                      else (1, x)):
        sections.append((sec, sorted([s for s in sts if s["section"] == sec], key=lambda s: (s["stmt_date"] or "",
                                                                                              s["s_no"]))))
    pal = S.section_palettes([s for s, _ in sections])
    bands = [Band(None, S.IDENTITY[0], None, [
        Col("ISRC", 34, "text", lambda r: JOIN.join(r["isrcs"]) or None),
        Col("Song Name", 42, "text", lambda r: r["song_name"]),
        Col("Internal No", 13, "text", lambda r: JOIN.join(r["work_nos"]) or None),
        Col("Total Amount", 15, "money", lambda r: r["royalty"], bold=True, total=True),
        Col(f"Total {client} Revenue (MRM)", 21, "money", lambda r: r["platform_rev"], bold=True, total=True),
    ])]
    for sec, group in sections:
        cols = []
        for st in group:
            key = str(st["id"])
            # §11.6 rule 3: the statement carries no distribution date - blank is a fact, zero is a claim
            cols += [Col("Date", 13, "date", lambda r: None),
                     Col("Amount", 12, "money", lambda r, k=key: r["detail"]["statements"].get(k), total=True),
                     Col("Period", 22, "text", lambda r, k=key, p=st["period"]: p if k in r["detail"]["statements"] else None),
                     Col("Distribution Number", 17, "text",
                         lambda r, k=key, d=dist_display(st["dist_no"]): d if k in r["detail"]["statements"] else None)]
        n = len(group)
        base = sec.upper() + (" - ROYALTY DISTRIBUTED BY IPRS" if sec == "Spotify" else "")
        title = f"{base}  ({n} distribution{'s' if n > 1 else ''}: {' and '.join(s['period'] for s in group)})"
        bands.append(Band(title, pal[sec][0], pal[sec][1], cols))
    if months:
        mcols = []
        dist = f"MRM {month_label(months[0])} - {month_label(months[-1])}"
        for mk in months:
            mcols += [Col("Date", 13, "date", lambda r, mk=mk: as_date(month_end_of(mk)) if r["detail"]["months"].get(mk) else None),
                      Col("Amount", 12, "money", lambda r, mk=mk: r["detail"]["months"].get(mk) or None, total=True),
                      Col("Period", 13, "text", lambda r, mk=mk: month_label(mk) if r["detail"]["months"].get(mk) else None),
                      Col("Distribution Number", 24, "text", lambda r, mk=mk: dist if r["detail"]["months"].get(mk) else None)]
        bands.append(Band(f"{client.upper()} - GROSS REVENUE REPORTED BY {client.upper()} (MRM report)  "
                          f"({month_label(months[0])} - {month_label(months[-1])}, {len(months)} monthly figures)",
                          S.SECTION_PALETTE["Spotify MRM"][0], S.SECTION_PALETTE["Spotify MRM"][1], mcols))
    fy_keys = sorted({k for r in run["rows"] for k in r["detail"].get("fy", {}) if k != "NA"})
    if any("NA" in r["detail"].get("fy", {}) for r in run["rows"]):
        fy_keys.append("NA")
    fcols = [Col(fy_label(k), 16, "money", lambda r, k=k: r["detail"].get("fy", {}).get(k, 0.0), total=True) for k in fy_keys]
    fcols.append(Col("Total Revenue", 16, "money", lambda r: r["royalty"], bold=True, total=True))
    bands.append(Band("TOTAL REVENUE  (Indian Financial Year, 1 Apr - 31 Mar)", S.TOTAL_REVENUE["band"],
                      S.TOTAL_REVENUE["tint"], fcols,
                      header_fills=[S.TOTAL_REVENUE["header"]] * len(fy_keys) + [S.TOTAL_REVENUE["band"]]))
    present = {"statement": "Source 1 (IPRS statement)", "platform": f"Source 2 ({client} MRM)",
               "both": f"Source 1 (IPRS statement) + Source 2 ({client} MRM)"}
    bands.append(Band("AUDIT / TRACEABILITY", S.AUDIT[0], S.AUDIT[1], [
        Col("Present In", 40, "text", lambda r: present[r["present_in"]]),
        Col("Song Name Exists In SVF List", 24, "text",
            lambda r: "Yes - identifiers differ" if r["evidence"].startswith("NAME PRESENT") else "No"),
    ]))
    wb, ws = new_workbook("Missing Songs Revenue")
    info = write_bands(ws, bands, run["rows"], total_gap=0, formulas=True)
    tr = info["total_row"]
    note(ws, tr + 2, (
        f"Songs earning money that are NOT in the song sheet '{run['sheet']['label']}' (mode {run['mode']}).  "
        "Layout and colours follow the SKV master workbook; no figure is taken from it - every amount is recomputed "
        "from the source statement lines and the platform report.  Total Amount = royalty actually distributed by "
        f"IPRS (statement money).  Total {client} Revenue (MRM) = gross revenue the platform reported for the ISRC - "
        "it is gross, not a distribution, so it is excluded from Total Amount and from Total Revenue.  A blank "
        "Internal No means the song reached us only through the platform report (ISRC only); a blank ISRC means it "
        "reached us only through a statement (internal number only).  Blank is a fact; zero is a claim.  "
        "'Song Name Exists In SVF List = Yes - identifiers differ' flags a title that does appear in the sheet while "
        "its internal number and ISRC do not - typically a cover, a male/female version or a re-recording that still "
        "needs registering.  The TOTAL row is a live SUM over the rows above."), info["ncol"], height=110)
    size, sha = save_deterministic(wb, path)
    return dict(size=size, sha256=sha, rows=len(run["rows"]), columns=info["ncol"])


def check_population_identity(list_path, revenue_path) -> dict:
    """I16 - the finding list and the revenue workbook hold the same songs in the same order."""
    def read(p):
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        rows = [list(x) for x in wb.worksheets[0].iter_rows(min_row=3, values_only=True)]
        wb.close()
        out = []
        for r in rows:
            if r[1] in (None, "") or str(r[1]).upper() == "TOTAL":
                if r[1] and str(r[1]).upper() == "TOTAL":
                    break
                continue
            out.append((r[0] or "", r[1], r[2] or ""))
        return out
    a, b = read(list_path), read(revenue_path)
    return dict(list_rows=len(a), revenue_rows=len(b), identical=a == b)
