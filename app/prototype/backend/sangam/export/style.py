"""The SKV visual system (architecture §19.1) - every fill, font, border, width and number
format lives here as a named token. A new section gets a token, not a hard-coded hex.
The same tokens are served to the UI through the build layout, so the screen and the
workbook use identical colours."""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

IDENTITY = ("D9D9D9", "F2F2F2")
SECTION_PALETTE = {
    "YouTube Pre-Claims": ("E7C3D8", "F9E9F2"),
    "YouTube Post-Claims": ("CBC1E0", "EEEAF7"),
    "Facebook / Meta": ("B7CBE4", "E6EDF7"),
    "Spotify MRM": ("B7CBE4", "E6EDF7"),
    "Spotify": ("BFD9BF", "E8F2E8"),
    "Apple Music": ("E0B4B4", "F7E7E7"),
    "Radio": ("E6CBA6", "F9F0E2"),
    "Zee TV Broadcast": ("ADCECA", "E6F1EF"),
    "Mechanical (MUSERK)": ("DAD7A0", "F3F2DF"),
    "Other / Unclassified": ("CFCFCF", "EFEFEF"),
    "Redistribution": ("D5BFA7", "F2E9DE"),
}
# overseas societies, alphabetical, cycle through eight tokens:
# APRA/SACEM, BMI/SOCAN, BUMA/SUISA, CASH, COMPASS, IMRO, MACP, MCT
OVERSEAS_CYCLE = [("B9C4D6", "E9EDF4"), ("C2CFB8", "ECF1E8"), ("DCC0C0", "F5EAEA"),
                  ("C9C2D9", "EFECF5"), ("DBCEB1", "F5F0E5"), ("B3CDD1", "E7F0F2"),
                  ("D2BFCE", "F2E9F0"), ("C9CDA8", "F0F1E3")]
TOTAL_REVENUE = dict(band="C9B99B", header="D8CBB3", tint="E8F2E8")
AUDIT = ("D8CBB7", "F0E9DC")
FY_TINTS = ["E8F2E8", "E6EDF7", "F9E9F2", "F5F0E5", "EEEAF7", "E7F0F2"]
STATUS_COLOURS = {       # the eight catalogue-status values -> the validated categorical slots 1..8
    "IN_RECEIVED": "2A78D6", "IN_NONE": "EB6834", "NOT_PAID": "1BAF7A", "IN_SIBLING": "EDA100",
    "NOT_PLATFORM": "E87BA4", "IN_ZERO": "008300", "NOT_ZERO": "4A3AA7", "NOT_USAGE": "E34948",
}

MONEY = "#,##0.00"
DATE = "DD-MMM-YYYY"
BAND_FONT = Font(bold=True, size=12, color="1F1F1F")
HEAD_FONT = Font(bold=True)
BOLD = Font(bold=True)
NOTE_FONT = Font(size=10, color="404040")
NOTE_ITALIC = Font(italic=True, size=10, color="404040")
THIN = Side(style="thin", color="BFBFBF")
MED = Side(style="medium", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_BORDER = Border(top=THIN, bottom=MED, left=THIN, right=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
WIDTHS = dict(isrc=34, name=42, internal=13, total=15, total_mrm=21, date=13, amount=12,
              period=24, dist=17)

_fills: dict[str, PatternFill] = {}


def fill(hex_: str) -> PatternFill:
    f = _fills.get(hex_)
    if f is None:
        f = _fills[hex_] = PatternFill("solid", fgColor=hex_)
    return f


def section_palettes(sections: list[str]) -> dict[str, tuple[str, str]]:
    out, k = {}, 0
    for sec in sections:
        if sec in SECTION_PALETTE:
            out[sec] = SECTION_PALETTE[sec]
        else:
            out[sec] = OVERSEAS_CYCLE[k % len(OVERSEAS_CYCLE)]
            k += 1
    return out
