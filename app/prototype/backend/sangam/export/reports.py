"""The other deliverables: the merge list (§12.5), the mismatch report V1/V2 (§13.5), the
statement register and the gap lists (§14)."""
from __future__ import annotations

import datetime as dt
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..build.engine import build_summary, load_rows
from ..build.model import BASIS
from ..core import db
from ..domain.filename import dist_display
from ..domain.normalize import JOIN, fy_label, fy_weights, month_label, penny_fix, r2
from . import style as S
from .writer import Band, Col, as_date, new_workbook, note, save_deterministic, write_bands


def _plain_table(ws, headers, widths, rows, *, header_row=1, fills=None, money_cols=(), pct_cols=(), date_cols=(),
                 wrap_cols=(), head_fill="D9D9D9"):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(header_row, c, value=h)
        cell.font = S.HEAD_FONT
        cell.fill = S.fill(head_fill)
        cell.border = S.BOX
        cell.alignment = S.WRAP_CENTER
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    ws.row_dimensions[header_row].height = 32
    r = header_row + 1
    for n, vals in enumerate(rows):
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, value=v)
            cell.border = S.BOX
            if c in money_cols:
                cell.number_format = S.MONEY
            elif c in pct_cols:
                cell.number_format = "0.00%"
            elif c in date_cols:
                cell.number_format = S.DATE
            if c in wrap_cols:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fills and n % 2 == 0:
                cell.fill = S.fill(fills)
        r += 1
    return r


# ------------------------------------------------------------------ §12.5 merge list
def write_merge_list(conn, build_id: int, path) -> dict:
    cands = db.rows(conn, "SELECT * FROM merge_candidate WHERE build_id = ? ORDER BY CASE kind WHEN "
                          "'VERSION_VARIANT' THEN 1 ELSE 0 END, confidence DESC, name_key", (build_id,))
    root = db.scalar(conn, "SELECT root_build_id FROM build_run WHERE id = ?", (build_id,))
    decisions = {(d["kind"], d["name_key"]): d for d in db.rows(
        conn, "SELECT * FROM merge_decision WHERE root_build_id = ? AND active = 1 ORDER BY id", (root,))}
    rows_out = []
    for g, c in enumerate(cands, 1):
        mem = db.rows(conn, "SELECT r.* FROM merge_candidate_member m JOIN rd_row r ON r.id = m.rd_row_id WHERE "
                            "m.candidate_id = ? ORDER BY r.ordinal", (c["id"],))
        isr, nos, names, reg = [], [], [], []
        for m in mem:
            if m["name"] not in names:
                names.append(m["name"])
            for x in db.rows(conn, "SELECT isrc, registered FROM rd_row_isrc WHERE row_id = ? ORDER BY ordinal", (m["id"],)):
                if x["isrc"] not in isr:
                    isr.append(x["isrc"])
                    reg.append(f"{x['isrc']} {'registered' if x['registered'] else 'NOT registered'}")
            for x in db.rows(conn, "SELECT work_no, registered FROM rd_row_work WHERE row_id = ? ORDER BY ordinal", (m["id"],)):
                if x["work_no"] not in nos:
                    nos.append(x["work_no"])
                    reg.append(f"{x['work_no']} {'registered' if x['registered'] else 'NOT registered'}")
        sig = db.loads(c["signals_json"], [])
        d = decisions.get((c["kind"], c["name_key"]))
        verdict = ("MERGED" if d["verdict"] == "MERGE" else "NOT THE SAME") + f" - {d['decided_by']}, {d['decided_at']}" \
            if d else "PENDING"
        rows_out.append([g, c["display_name"], JOIN.join(names), JOIN.join(isr) or None, JOIN.join(nos) or None,
                         c["kind"], f"{c['confidence']:.2f} ({c['band']}): " + "; ".join(
                             f"{s['weight']:+.2f} {s['label']}" for s in sig),
                         verdict, JOIN.join(f"{m['total_amount']:,.2f}" for m in mem),
                         r2(sum(m["total_amount"] for m in mem)), JOIN.join(f"{m['total_mrm']:,.2f}" for m in mem),
                         r2(sum(m["total_mrm"] for m in mem)), "; ".join(reg)])
    wb, ws = new_workbook("Same Name Diff ISRC")
    hdr = ["Group", "Song Name", "Name As Each Source Spells It", "ISRCs In This Group", "Internal Numbers In This Group",
           "Kind", "Confidence", "Decision", "Royalty Before", "Royalty After", "Spotify Revenue Before",
           "Spotify Revenue After", "In The User's Sheet?"]
    _plain_table(ws, hdr, [7, 32, 40, 44, 26, 26, 60, 34, 22, 14, 24, 14, 60], rows_out, fills="F2F2F2",
                 money_cols=(10, 12), wrap_cols=(3, 4, 7, 13))
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:M{len(rows_out) + 1}"
    size, sha = save_deterministic(wb, path)
    return dict(size=size, sha256=sha, rows=len(rows_out))


# ------------------------------------------------------------------ §13.5 mismatch V1 / V2
def load_mismatch(conn, run_id: int) -> dict:
    run = db.row(conn, "SELECT * FROM mismatch_run WHERE id = ?", (run_id,))
    run["counts"] = db.loads(run["counts_json"], {})
    run["totals"] = db.loads(run["totals_json"], {})
    rows = db.rows(conn, "SELECT * FROM mismatch WHERE run_id = ? ORDER BY ordinal", (run_id,))
    by = {r["id"]: r for r in rows}
    for r in rows:
        r.update(isrcs=[], values=[], sources=[])
    for x in conn.execute("SELECT i.* FROM mismatch_main_isrc i JOIN mismatch m ON m.id = i.mismatch_id WHERE "
                          "m.run_id = ? ORDER BY i.mismatch_id, i.ordinal", (run_id,)):
        by[x["mismatch_id"]]["isrcs"].append(x["isrc"])
    for x in conn.execute("SELECT v.* FROM mismatch_value v JOIN mismatch m ON m.id = v.mismatch_id WHERE "
                          "m.run_id = ? ORDER BY v.mismatch_id, v.value_text", (run_id,)):
        by[x["mismatch_id"]]["values"].append(x["value_text"])
    for x in conn.execute("SELECT s.* FROM mismatch_source s JOIN mismatch m ON m.id = s.mismatch_id WHERE "
                          "m.run_id = ? ORDER BY s.mismatch_id, s.source", (run_id,)):
        by[x["mismatch_id"]]["sources"].append(x["source"])
    run["rows"] = rows
    return run


def write_mismatch(run: dict, path, revenue: bool) -> dict:
    labels = run["counts"].get("labels", {})
    hdr = ["S. No", "ISRC (as per SVF Song List)", "Song Name (as per SVF Song List)", "Internal No. (as per SVF Song List)",
           "Type of Mismatch", "Conflicting Value(s) Found", "Sheet(s) Where the Conflict Was Found"]
    wid = [7, 34, 34, 17, 24, 46, 42]
    rhdr = ["Royalty Paid on This Song", "Gross Revenue on This Song", "Total (Both Files)"]
    hfill, sfill, tfill = "DCE9F7", "FCE4D6", "E2EFDA"
    thin = Side(style="thin", color="B7C9DC")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb, ws = new_workbook("Song ID Mismatches")
    top = 2 if revenue else 1
    nc = len(hdr) + (3 if revenue else 0)
    t = run["totals"]
    if revenue:
        c = ws.cell(1, 1, value=f"SONG IDENTITY MISMATCHES BETWEEN THE SONG SHEET, THE IPRS STATEMENT(S) AND THE "
                                f"PLATFORM'S OWN REPORT   -   {len(run['rows']):,} rows")
        c.font = Font(bold=True, size=11, color="1F3864")
        c.fill = S.fill(hfill)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(hdr))
        secs = [f"ROYALTY DISTRIBUTED BY IPRS\n{labels.get('S3', '')}\nTotal royalty in the selection : {t['iprs_file']:,.2f}",
                f"GROSS REVENUE REPORTED BY THE PLATFORM\n{labels.get('S4', '')}\nTotal revenue in the selection : "
                f"{t['mrm_file']:,.2f}", "BOTH SOURCES TOGETHER\nbooked once per IPRS work / per ISRC"]
        for j, lab in enumerate(secs):
            c = ws.cell(1, len(hdr) + 1 + j, value=lab)
            c.font = Font(bold=True, size=9, color="833C0C")
            c.fill = S.fill(sfill)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 58
        for cc in range(1, nc + 1):
            ws.cell(1, cc).border = bd
    for c, h in enumerate(hdr + (rhdr if revenue else []), 1):
        cell = ws.cell(top, c, value=h)
        cell.font = Font(bold=True, size=11, color="1F3864")
        cell.fill = S.fill(hfill)
        cell.border = bd
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = (wid + [18, 19, 17])[c - 1]
    ws.row_dimensions[top].height = 34
    lab = {k: labels.get(k, k) for k in ("S2", "S3", "S4")}
    for i, d in enumerate(run["rows"], 1):
        r = top + i
        vals = [i, JOIN.join(d["isrcs"]) or None, d["main_name"] or None, d["main_no"] or None, d["kind"],
                JOIN.join(d["values"]), JOIN.join(lab[s] for s in d["sources"])]
        if revenue:
            vals += [d["royalty"], d["gross"], r2(d["royalty"] + d["gross"])]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, value=v)
            cell.border = bd
            cell.alignment = Alignment(vertical="top", horizontal="center" if c in (1, 4) else "left",
                                       wrap_text=c in (2, 6, 7))
            if c > len(hdr):
                cell.number_format = S.MONEY
                cell.alignment = Alignment(vertical="top", horizontal="right")
    if revenue:
        r = top + len(run["rows"]) + 1
        c = ws.cell(r, 1, value="TOTAL  (booked once per IPRS work / per ISRC)")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(hdr))
        c.font = Font(bold=True, size=11, color="1F3864")
        c.alignment = Alignment(horizontal="right", vertical="center")
        for j, v in enumerate([t["royalty"], t["gross"], t["both"]]):
            cell = ws.cell(r, len(hdr) + 1 + j, value=v)
            cell.font = S.BOLD
            cell.number_format = S.MONEY
        for cc in range(1, nc + 1):
            ws.cell(r, cc).fill = S.fill(tfill)
            ws.cell(r, cc).border = bd
    ws.freeze_panes = f"A{top + 1}"
    ws.auto_filter.ref = f"A{top}:{get_column_letter(nc)}{top + max(1, len(run['rows']))}"
    size, sha = save_deterministic(wb, path)
    return dict(size=size, sha256=sha, rows=len(run["rows"]))


def check_v1_v2(v1_path, v2_path) -> dict:
    """I12 - V1 and V2 carry identical A-G content and row count."""
    def read(p, top):
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        rows = [tuple(x[:7]) for x in wb.worksheets[0].iter_rows(min_row=top + 1, values_only=True)
                if isinstance(x[0], int)]
        wb.close()
        return rows
    a, b = read(v1_path, 1), read(v2_path, 2)
    return dict(v1_rows=len(a), v2_rows=len(b), identical=a == b)


# ------------------------------------------------------------------ §14 statement register + gap lists
def write_register(conn, build_id: int, path) -> dict:
    b = build_summary(conn, build_id)
    sts = db.rows(conn, "SELECT st.*, sf.filename FROM statement st JOIN source_file sf ON sf.id = st.source_file_id "
                        "JOIN build_statement bs ON bs.statement_id = st.id WHERE bs.build_id = ? ORDER BY st.s_no",
                  (build_id,))
    grand = r2(sum(r2(s["extracted_total"]) for s in sts))
    rows = []
    for s in sts:
        w = fy_weights(dt.date.fromisoformat(s["p_start"]) if s["p_start"] else None,
                       dt.date.fromisoformat(s["p_end"]) if s["p_end"] else None)
        rows.append([f"S-{s['s_no']}", s["filename"], dist_display(s["dist_no"]), s["section"], s["period"],
                     ", ".join(fy_label(y) for y in w) if w else "Not specified", as_date(s["stmt_date"]),
                     s["work_count"], s["line_count"], r2(s["extracted_total"]),
                     r2(s["extracted_total"]) / grand if grand else 0, "PASS" if s["reconciled"] else "FAIL"])
    wb, ws = new_workbook("Statement Register")
    t = ws.cell(1, 1, value=f"STATEMENT REGISTER - {b['label']} - one row per royalty distribution statement "
                            f"({len(sts)} statements)")
    t.font = Font(bold=True, size=13)
    t.fill = S.fill(S.IDENTITY[0])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=12)
    hdr = ["S. No", "Statement File", "Distribution Number", "Source / Category", "Period Covered",
           "Financial Year(s) Covered", "Statement Date", "Works", "Royalty Lines", "Total Amount", "% of Grand Total",
           "Reconciled"]
    r = _plain_table(ws, hdr, [8, 70, 18, 30, 34, 28, 14, 9, 13, 16, 14, 12], rows, header_row=2, fills="F2F2F2",
                     money_cols=(10,), pct_cols=(11,), date_cols=(7,))
    for j, v in enumerate([f"TOTAL ({len(sts)} statements)", None, None, None, None, None, None,
                           sum(s["work_count"] for s in sts), sum(s["line_count"] for s in sts), grand, 1.0, None], 1):
        c = ws.cell(r, j, value=v)
        c.font = S.BOLD
        c.fill = S.fill(S.TOTAL_REVENUE["band"])
        c.border = Border(top=S.MED, bottom=S.MED)
        if j == 10:
            c.number_format = S.MONEY
        if j == 11:
            c.number_format = "0.00%"
    r += 3
    by_sec = defaultdict(lambda: [0, 0.0])
    for s in sts:
        by_sec[s["section"]][0] += 1
        by_sec[s["section"]][1] += r2(s["extracted_total"])
    ws.cell(r, 1, value="BY SOURCE / CATEGORY").font = Font(bold=True, size=12)
    r = _plain_table(ws, ["Source / Category", "Statements", "Total Amount", "% of Grand Total"], [8, 70, 18, 30],
                     [[k, v[0], r2(v[1]), v[1] / grand if grand else 0]
                      for k, v in sorted(by_sec.items(), key=lambda kv: -kv[1][1])],
                     header_row=r + 1, money_cols=(3,), pct_cols=(4,))
    r += 2
    ws.cell(r, 1, value="BY EARNING PERIOD (Indian FY, month-weighted)").font = Font(bold=True, size=12)
    fy = b["layout"]["fy_totals"]
    _plain_table(ws, ["Financial Year", "Total Amount", "% of Grand Total"], [8, 70, 18],
                 [[fy_label(k), v, v / grand if grand else 0] for k, v in fy.items()], header_row=r + 1,
                 money_cols=(2,), pct_cols=(3,))
    ws.column_dimensions["A"].width = 30
    ws.freeze_panes = "C3"
    size, sha = save_deterministic(wb, path)
    return dict(size=size, sha256=sha, rows=len(rows))


def write_gap(conn, build_id: int, which: str, path) -> dict:
    """Views of the build with a fixed configuration (kept for continuity; §14)."""
    b = build_summary(conn, build_id)
    L = b["layout"]
    rows = [r for r in load_rows(conn, build_id)
            if r.origin == {"works": "statement", "platform": "platform", "usage": "usage"}[which]]
    stmts = {st["id"]: sec["section"] for sec in L["sections"] if sec["kind"] == "statement" for st in sec["statements"]}
    sections = [sec["section"] for sec in L["sections"] if sec["kind"] == "statement"]
    months = [m["month"] for sec in L["sections"] if sec["kind"] == "platform" for m in sec["months"]]
    ident = [Col("Internal No", 14, "text", lambda r: JOIN.join(r.ordered_works()) or None),
             Col("Song Name", 44, "text", lambda r: r.name),
             Col("ISRC", 34, "text", lambda r: JOIN.join(r.ordered_isrcs()) or None)]
    if which == "works":
        title = "Missing-From-SVF-List"
        bands = [Band(None, S.IDENTITY[0], None, ident + [
            Col("ISRC Provenance", 30, "text", lambda r: " ; ".join(BASIS[b_] for b_ in r.basis) or None),
            Col("Total Royalty", 15, "money", lambda r: r.total_amount, bold=True, total=True),
            Col(f"{L['client']} Revenue (MRM)", 18, "money", lambda r: r.total_mrm, total=True)])]
        bands.append(Band("ROYALTY BY SOURCE / CATEGORY", "D8CBB3", "F5F0E5", [
            Col(sec, 16, "money", lambda r, sec=sec: r2(sum(a for sid, a in r.amounts.items() if stmts.get(sid) == sec))
                or None, total=True) for sec in sections]))
    elif which == "platform":
        title = "Spotify-Revenue-Only"
        bands = [Band(None, S.IDENTITY[0], None, ident[1:] + [
            Col("Album", 34, "text", lambda r: next((d.get("album") for d in r.mrm_isrcs.values() if d.get("album")), None)),
            Col("Total Revenue", 15, "money", lambda r: r.total_mrm, bold=True, total=True),
            Col("Streams", 12, "int", lambda r: int(r.usage), total=True),
            Col("IPRS Royalty", 13, "money", lambda r: r.total_amount, total=True),
            Col("Match Basis", 40, "text", lambda r: " ; ".join(BASIS[b_] for b_ in r.basis) or None)])]
        bands.append(Band("MONTHLY GROSS REVENUE", *S.SECTION_PALETTE["Spotify MRM"],
                          [Col(month_label(m), 13, "money", lambda r, m=m: r.mrm.get(m), total=True) for m in months]))
    else:
        title = "Spotify-Streams-Only"
        bands = [Band(None, S.IDENTITY[0], None, ident[1:] + [
            Col("Streams", 14, "int", lambda r: int(r.usage), bold=True, total=True),
            Col("Revenue", 13, "money", lambda r: r.total_mrm, total=True),
            Col("Royalty", 13, "money", lambda r: r.total_amount, total=True)])]
    wb, ws = new_workbook(title)
    info = write_bands(ws, bands, rows, freeze="C3", total_label_col=1)
    size, sha = save_deterministic(wb, path)
    return dict(size=size, sha256=sha, rows=len(rows))
